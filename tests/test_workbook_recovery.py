"""Manual checkpoints and offline recovery use only synthetic output pairs."""

import unittest
from argparse import Namespace
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook
from pydantic import ValidationError

from scdp_automation import extrator, workbook_recovery, xlsx_output
from scdp_automation.output_history import OutputHistory
from scdp_automation.relatorio import load_trips, save_json
from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import final_template_fixture, make_trip


class WorkbookRecoveryTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template = self.root / "template.xlsx"
        final_template_fixture(self.template)
        self.checkpoint = self.root / "viagens_scdp_2026.json"
        self.workbook = self.root / "gastos_scdp_2026.xlsx"
        profile = self.root / ".scdp-browser"
        (profile / "Default").mkdir(parents=True)
        (profile / "Local State").write_text("{}")
        self.trips = [make_trip(), make_trip("123456/26-3C")]
        for trip in self.trips:
            trip.descricao_do_motivo_da_viagem = "Descrição persistida"
            trip.data_da_ultima_verificacao = date(2026, 1, 1)
        save_json(self.checkpoint, self.trips)
        xlsx_output.publish_workbook(
            self.trips, self.workbook, template_path=self.template
        )
        default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.template)
        default.start()
        self.addCleanup(default.stop)

    def edit_manual(self, code="AGRONOMIA", decision="Sim"):
        book = load_workbook(self.workbook)
        base = book["BASE VIAGENS"]
        base["Q2"] = code
        base["R2"] = decision
        # Reordering travel rows cannot change the full-identity association.
        first = [cell.value for cell in base[2]]
        second = [cell.value for cell in base[3]]
        for column, (a, b) in enumerate(zip(first, second, strict=True), 1):
            base.cell(2, column, b)
            base.cell(3, column, a)
        book.save(self.workbook)
        book.close()

    def snapshot(self):
        return {
            path.relative_to(self.root): path.read_bytes()
            for path in self.root.rglob("*")
            if path.is_file()
        }

    async def run_extraction(self, listed, *, fail_browser=False):
        page = MagicMock()
        page.url = "https://www2.scdp.gov.br/novoscdp/home.xhtml"
        page.bring_to_front = AsyncMock()
        browser = MagicMock()
        browser.close = AsyncMock()
        manager = MagicMock()
        manager.__aenter__ = AsyncMock(return_value=object())
        manager.__aexit__ = AsyncMock(return_value=False)
        checkpoints = []
        original_save = extrator.save_json

        def observe(path, trips):
            original_save(path, trips)
            checkpoints.append(load_trips(path))

        async def reason(*_):
            return "Descrição atualizada"

        with (
            patch.object(extrator, "configure_logging"),
            patch.object(
                extrator,
                "parse_args",
                return_value=Namespace(
                    output=self.checkpoint, workbook=self.workbook, ano=2026, limite=0
                ),
            ),
            patch.object(
                extrator,
                "async_playwright",
                side_effect=AssertionError("browser started") if fail_browser else None,
                return_value=manager,
            ),
            patch.object(
                extrator, "selected_profile_directory", return_value="Default"
            ),
            patch.object(extrator, "REPO_ROOT", self.root),
            patch.object(
                extrator, "connect_visible_chrome", new=AsyncMock(return_value=browser)
            ),
            patch.object(
                extrator, "select_scdp_page", new=AsyncMock(return_value=page)
            ),
            patch.object(extrator, "start_login_if_needed", new=AsyncMock()),
            patch.object(extrator, "wait_for_login", new=AsyncMock()),
            patch.object(extrator, "open_annual_cch_report", new=AsyncMock()),
            patch.object(
                extrator, "collect_listing", new=AsyncMock(return_value=listed)
            ),
            patch.object(
                extrator, "consult_trip_reason", new=AsyncMock(side_effect=reason)
            ),
            patch.object(extrator, "current_date", return_value=date(2026, 2, 1)),
            patch.object(extrator, "save_json", side_effect=observe),
        ):
            await extrator.run([])
        return checkpoints

    async def test_extraction_checkpoints_preserve_manual_values(self):
        self.edit_manual()
        originals = (self.checkpoint.read_bytes(), self.workbook.read_bytes())
        listed = [make_trip("123456/26-3C"), make_trip()]
        listed[1].total_da_viagem_r += 100
        checkpoints = await self.run_extraction(listed)
        for records in checkpoints:
            trip = next(
                t
                for t in records
                if t.numero_da_solicitacao == self.trips[0].numero_da_solicitacao
            )
            self.assertEqual(
                (trip.codigo_de_debito, trip.descontar_do_curso), ("AGRONOMIA", "Sim")
            )
            self.assertIsNotNone(trip.descricao_do_motivo_da_viagem)
            self.assertIsNotNone(trip.data_da_ultima_verificacao)
        final = load_trips(self.checkpoint)[1]
        self.assertEqual(final.total_da_viagem_r, listed[1].total_da_viagem_r)
        self.assertEqual(final.descricao_do_motivo_da_viagem, "Descrição atualizada")
        self.assertEqual(final.data_da_ultima_verificacao, date(2026, 2, 1))
        self.assertEqual(len(list((self.root / "backup").glob("*.history.*"))), 2)
        self.assertEqual(
            next((self.root / "backup").glob("*.json")).read_bytes(), originals[0]
        )
        self.assertEqual(
            next((self.root / "backup").glob("*.xlsx")).read_bytes(), originals[1]
        )

    async def test_cleared_choices_stay_cleared_after_second_publish(self):
        self.trips[0].codigo_de_debito = "AGRONOMIA"
        self.trips[0].descontar_do_curso = "Sim"
        save_json(self.checkpoint, self.trips)
        self.edit_manual(None, None)
        checkpoints = await self.run_extraction(
            [make_trip(), make_trip("123456/26-3C")]
        )
        self.assertGreaterEqual(len(checkpoints), 4)
        for records in checkpoints:
            self.assertIsNone(records[0].codigo_de_debito)
            self.assertIsNone(records[0].descontar_do_curso)
        book = load_workbook(self.workbook)
        self.addCleanup(book.close)
        self.assertIsNone(book["BASE VIAGENS"]["Q2"].value)
        self.assertIsNone(book["BASE VIAGENS"]["R2"].value)

    def test_retry_after_publication_failure_keeps_manual_fields_and_original_backup(
        self,
    ):
        self.edit_manual()
        original = self.snapshot()
        trips = workbook_recovery.prepare_checkpoint_trips(
            self.checkpoint, self.workbook
        )
        trips[0].total_da_viagem_r += 100
        replace = xlsx_output.os.replace

        def fail_workbook_replace(source, destination):
            if Path(destination) == self.workbook:
                raise PermissionError("locked")
            return replace(source, destination)

        with (
            OutputHistory(self.checkpoint, self.workbook) as history,
            patch.object(xlsx_output.os, "replace", side_effect=fail_workbook_replace),
            self.assertRaises(xlsx_output.WorkbookPublishError),
        ):
            extrator.save_checkpoint_and_publish(
                trips, self.checkpoint, self.workbook, history=history
            )
        self.assertEqual(load_trips(self.checkpoint), trips)
        self.assertEqual(self.workbook.read_bytes(), original[Path(self.workbook.name)])
        recoveries = list((self.root / "backup").glob("recovery.*.xlsx"))
        self.assertEqual(len(recoveries), 1)
        original_json = next((self.root / "backup").glob("*.history.*.json"))
        original_xlsx = next((self.root / "backup").glob("*.history.*.xlsx"))
        workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
        self.assertEqual(load_trips(self.checkpoint), trips)
        self.assertEqual(
            original_json.read_bytes(), original[Path(self.checkpoint.name)]
        )
        self.assertEqual(original_xlsx.read_bytes(), original[Path(self.workbook.name)])
        self.assertFalse(recoveries[0].exists())

    def test_json_only_recovery_preserves_persisted_base_choices(self):
        self.workbook.unlink()
        self.trips[0].codigo_de_debito = "AGRONOMIA"
        self.trips[0].descontar_do_curso = "Não"
        save_json(self.checkpoint, self.trips)
        workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
        self.assertEqual(load_trips(self.checkpoint), self.trips)
        book = load_workbook(self.workbook)
        self.addCleanup(book.close)
        self.assertEqual(book["BASE VIAGENS"]["Q2"].value, "AGRONOMIA")
        self.assertEqual(book["BASE VIAGENS"]["R2"].value, "Não")
        template = load_workbook(self.template)
        self.addCleanup(template.close)
        self.assertEqual(book["APOIO"]["D2"].value, template["APOIO"]["D2"].value)

    def test_prepare_is_read_only_and_returns_copies(self):
        self.edit_manual()
        before = self.snapshot()
        prepared = workbook_recovery.prepare_checkpoint_trips(
            self.checkpoint, self.workbook
        )
        self.assertEqual(prepared[0].codigo_de_debito, "AGRONOMIA")
        self.assertEqual(self.snapshot(), before)
        self.assertIsNone(load_trips(self.checkpoint)[0].codigo_de_debito)

    async def test_invalid_sources_rejected_without_writes_or_browser(self):
        original_json = self.checkpoint.read_bytes()
        original_workbook = self.workbook.read_bytes()
        for failure in (
            "corrupt",
            "orphan",
            "duplicate JSON",
            "duplicate XLSX",
            "ambiguous",
            "missing PCDP",
            "incompatible APOIO",
        ):
            with self.subTest(failure=failure):
                self.checkpoint.write_bytes(original_json)
                self.workbook.write_bytes(original_workbook)
                if failure == "corrupt":
                    self.checkpoint.write_text("not JSON")
                elif failure == "orphan":
                    self.checkpoint.unlink()
                elif failure == "duplicate JSON":
                    save_json(self.checkpoint, [self.trips[0], self.trips[0]])
                else:
                    book = load_workbook(self.workbook)
                    if failure == "duplicate XLSX":
                        book["BASE VIAGENS"]["A3"] = book["BASE VIAGENS"]["A2"].value
                    elif failure == "ambiguous":
                        book["BASE VIAGENS"]["S1"] = "Código de débito"
                    elif failure == "missing PCDP":
                        book["BASE VIAGENS"]["A2"] = "999999/26"
                    else:
                        book["APOIO"].insert_cols(8)
                        book["APOIO"]["H1"] = "Transportes pago (R$)"
                    book.save(self.workbook)
                    book.close()
                before = self.snapshot()
                with self.assertRaises((WorkbookValidationError, ValidationError)):
                    await self.run_extraction(self.trips, fail_browser=True)
                self.assertEqual(self.snapshot(), before)
                with self.assertRaises((WorkbookValidationError, ValidationError)):
                    workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
                self.assertEqual(self.snapshot(), before)

    def test_unknown_json_only_code_fails_before_writes(self):
        self.workbook.unlink()
        self.trips[0].codigo_de_debito = "UNKNOWN"
        save_json(self.checkpoint, self.trips)
        before = self.snapshot()
        with self.assertRaises(WorkbookValidationError):
            workbook_recovery.prepare_checkpoint_trips(self.checkpoint, self.workbook)
        self.assertEqual(self.snapshot(), before)
        with self.assertRaises(WorkbookValidationError):
            workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
        self.assertEqual(self.snapshot(), before)

    def test_recovery_never_opens_browser_or_selects_profile(self):
        with (
            patch.object(
                extrator, "async_playwright", side_effect=AssertionError("browser")
            ),
            patch.object(
                extrator,
                "selected_profile_directory",
                side_effect=AssertionError("profile"),
            ),
            patch.object(
                extrator, "collect_listing", side_effect=AssertionError("listing")
            ),
        ):
            workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
        self.assertEqual(load_trips(self.checkpoint), self.trips)

    def test_missing_checkpoint_never_creates_empty_workbook(self):
        self.checkpoint.unlink()
        self.workbook.unlink()
        before = self.snapshot()
        with self.assertRaises(WorkbookValidationError):
            workbook_recovery.recreate_workbook(self.checkpoint, self.workbook)
        self.assertEqual(self.snapshot(), before)

    async def test_missing_listed_pcdp_reconciles_before_overwrite(self):
        self.edit_manual()
        original_json = self.checkpoint.read_bytes()
        with self.assertRaises(WorkbookValidationError) as raised:
            await self.run_extraction([make_trip()])
        self.assertEqual(raised.exception.missing_pcdps, ("123456/26-3C",))
        records = load_trips(self.checkpoint)
        self.assertEqual(len(records), 2)
        self.assertEqual(records[0].codigo_de_debito, "AGRONOMIA")
        self.assertEqual(
            next((self.root / "backup").glob("*.json")).read_bytes(), original_json
        )
