"""Read the original workbook's inputs without saving or changing it."""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from scdp_automation.relatorio import PCDP_PATTERN
from scdp_automation.xlsx_models import DEBIT_CATEGORIES, TripSummary


@dataclass(frozen=True)
class BudgetInputs:
    total: float | None = None
    daily: float | None = None
    transport: float | None = None
    scheduled: float | None = None
    paid: float | None = None


@dataclass(frozen=True)
class ReferenceData:
    trips: tuple[TripSummary, ...]
    codes: dict[str, str]
    inputs: dict[str, BudgetInputs]
    labels: dict[str, str]
    occurrences: int
    rateio: tuple[float, float, float] | None
    decisions: dict[str, str] = field(default_factory=dict)


def number(value: object, location: str) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
    ):
        raise ValueError(f"Valor numérico ausente ou inválido em {location}.")
    return float(value)


def canonical_code(value: object) -> str | None:
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        raise TypeError("Código de débito inválido na referência.")
    code = value.strip()
    if code == "LS Enf em Oncologia":
        code = "Lato Oncologia"
    if code not in {category.code for category in DEBIT_CATEGORIES} | {"PPGH", "PPGDH"}:
        raise ValueError("Código de débito desconhecido na referência.")
    return code


def _read_number(
    formulas: Workbook, cached: Workbook, sheet: str, coordinate: str
) -> float:
    value = cached[sheet][coordinate].value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return number(value, f"{sheet}!{coordinate}")
    formula = formulas[sheet][coordinate].value
    if isinstance(formula, str) and re.fullmatch(r"=[A-Z]+[0-9]+", formula):
        return number(cached[sheet][formula[1:]].value, f"{sheet}!{formula[1:]}")
    return number(value, f"{sheet}!{coordinate}")


def read_reference(path: Path) -> ReferenceData:
    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    try:
        return _read_reference(formulas, cached)
    finally:
        formulas.close()
        cached.close()


def _read_reference(formulas: Workbook, cached: Workbook) -> ReferenceData:
    required = {"BD D&P", "APOIO", "RESUMO GASTOS"}
    if not required.issubset(cached.sheetnames):
        raise ValueError("A referência não contém BD D&P, APOIO e RESUMO GASTOS.")
    base = cached["BD D&P"]
    trips: dict[str, TripSummary] = {}
    codes: dict[str, str] = {}
    current: tuple[str, str, str, str | None] | None = None
    subtotal: tuple[float, float, float] | None = None
    occurrences = 0
    starts: list[date] = []
    ends: list[date] = []
    for row in range(5, base.max_row + 1):
        pcdp = base.cell(row, 4).value
        if isinstance(pcdp, str) and re.fullmatch(PCDP_PATTERN, pcdp):
            if current is not None:
                raise ValueError("Viagem sem rodapé completo na referência.")
            proposed = base.cell(row, 5).value
            status = base.cell(row, 9).value
            if not isinstance(proposed, str) or not isinstance(status, str):
                raise ValueError("Identificação da viagem incompleta na referência.")
            code = canonical_code(base.cell(row, 3).value)
            if code:
                expected_segment = next(
                    c.segment
                    for c in DEBIT_CATEGORIES
                    if c.code == ("PPGH/PPGDH" if code in ("PPGH", "PPGDH") else code)
                )
                if base.cell(row, 2).value != expected_segment:
                    raise ValueError(
                        "Segmento incompatível com o código de débito na referência."
                    )
            current = pcdp, proposed, status, code
            subtotal = None
            starts, ends = [], []
        if current is None:
            continue
        for column, dates in ((11, starts), (12, ends)):
            value = base.cell(row, column).value
            if isinstance(value, datetime):
                dates.append(value.date())
            elif isinstance(value, date):
                dates.append(value)
            elif isinstance(value, str) and re.fullmatch(
                r"\d{2}/\d{2}/\d{4}", value.strip()
            ):
                day, month, year = map(int, value.strip().split("/"))
                dates.append(date(year, month, day))
        if (
            base.cell(row, 15).value == "Sub-Total"
            or base.cell(row, 4).value == "Sub-Total"
        ):
            subtotal = (
                number(base.cell(row, 16).value, f"BD D&P linha {row}"),
                number(base.cell(row, 17).value, f"BD D&P linha {row}"),
                number(base.cell(row, 18).value, f"BD D&P linha {row}"),
            )
        if base.cell(row, 16).value != "Total da Viagem (R$)":
            continue
        if subtotal is None and current[0] in trips:
            previous = trips[current[0]]
            unlabeled = (
                number(base.cell(row - 1, 16).value, f"BD D&P linha {row - 1}"),
                number(base.cell(row - 1, 17).value, f"BD D&P linha {row - 1}"),
                number(base.cell(row - 1, 18).value, f"BD D&P linha {row - 1}"),
            )
            if unlabeled == (
                previous.daily_count,
                previous.daily_amount,
                previous.ticket_amount,
            ):
                subtotal = unlabeled
        if subtotal is None:
            raise ValueError("Viagem sem subtotal na referência.")
        pcdp, proposed, status, code = current
        additional, discount, restitution, reimbursement, total = (
            number(base.cell(row, col).value, f"BD D&P linha {row}")
            for col in (6, 9, 12, 15, 19)
        )
        trip = TripSummary(
            pcdp,
            proposed,
            status,
            *subtotal,
            additional,
            discount,
            restitution,
            reimbursement,
            total,
            min(starts) if starts else None,
            max(ends) if ends else None,
        )
        previous = trips.get(pcdp)
        if previous is not None and previous != trip:
            raise ValueError(
                "PCDP repetida com dados financeiros conflitantes na referência."
            )
        trips[pcdp] = trip
        if code:
            if pcdp in codes and codes[pcdp] != code:
                raise ValueError(
                    "PCDP repetida com classificações conflitantes na referência."
                )
            codes[pcdp] = code
        occurrences += 1
        current = None
    if current is not None or not trips:
        raise ValueError("Base de viagens vazia ou incompleta na referência.")

    support = cached["APOIO"]
    summary = cached["RESUMO GASTOS"]
    labels: dict[str, str] = {}
    inputs: dict[str, BudgetInputs] = {}
    for row in range(7, 35):
        label = support.cell(row, 2).value
        raw_code = support.cell(row, 3).value
        if not label or not raw_code:
            continue
        code = canonical_code(raw_code)
        assert code is not None
        if not isinstance(label, str):
            raise TypeError("Rótulo de orçamento inválido na referência.")
        labels[code] = label
        summary_rows = [
            r
            for r in (*range(7, 22), *range(26, 36), 41, 43)
            if summary.cell(r, 2).value == label
        ]
        if len(summary_rows) > 1:
            raise ValueError("Orçamento duplicado na referência.")
        total = daily = transport = None
        if summary_rows:
            summary_row = summary_rows[0]
            total = number(
                summary.cell(summary_row, 3).value, f"RESUMO GASTOS!C{summary_row}"
            )
            daily = number(
                summary.cell(summary_row, 5).value, f"RESUMO GASTOS!E{summary_row}"
            )
            transport = total - daily
        inputs[code] = BudgetInputs(
            total,
            daily,
            transport,
            _read_number(formulas, cached, "APOIO", f"M{row}"),
            _read_number(formulas, cached, "APOIO", f"P{row}"),
        )
    rateio = None
    split_total = sum(
        trip.trip_total for trip in trips.values() if codes.get(trip.pcdp) == "PPGEL +"
    )
    if split_total:
        if tuple(support.cell(row, 9).value for row in (58, 59, 60)) != (
            "PPGE",
            "PPGEL",
            "PPGH",
        ):
            raise ValueError(
                "Destinos do rateio PPGEL + não reconhecidos na referência."
            )
        amounts = tuple(
            _read_number(formulas, cached, "APOIO", f"J{row}") for row in (58, 59, 60)
        )
        if not math.isclose(sum(amounts), split_total, abs_tol=0.005):
            raise ValueError("O rateio PPGEL + não reconcilia com o total das viagens.")
        rateio = (
            amounts[0] / split_total,
            amounts[1] / split_total,
            amounts[2] / split_total,
        )
    decisions = infer_decisions(tuple(trips.values()), codes, cached)
    history_inputs = [inputs[c] for c in ("PPGH", "PPGDH") if c in inputs]
    if history_inputs:
        combined = {}
        for key in ("total", "daily", "transport", "scheduled", "paid"):
            numbers = [
                getattr(v, key) for v in history_inputs if getattr(v, key) is not None
            ]
            combined[key] = sum(numbers) if numbers else None
        inputs["PPGH/PPGDH"] = BudgetInputs(**combined)
        labels["PPGH/PPGDH"] = "Mestrado e Doutorado em História"
    codes = {
        pcdp: "PPGH/PPGDH" if code in ("PPGH", "PPGDH") else code
        for pcdp, code in codes.items()
    }
    return ReferenceData(
        tuple(trips.values()), codes, inputs, labels, occurrences, rateio, decisions
    )


def infer_decisions(
    trips: tuple[TripSummary, ...], codes: dict[str, str], cached: Workbook
) -> dict[str, str]:
    """Require reconciliation with original utilized amounts, never status alone."""
    utilized = {}
    support = cached["APOIO"]
    for row in range(41, support.max_row + 1):
        raw_code = support.cell(row, 2).value
        amount = support.cell(row, 3).value
        if not isinstance(raw_code, str) or not isinstance(amount, (int, float)):
            continue
        try:
            code = canonical_code(raw_code)
        except TypeError, ValueError:
            continue
        if code:
            utilized[code] = float(amount)
    decisions = {}
    for code, amount in utilized.items():
        members = [t for t in trips if codes.get(t.pcdp) == code]
        cancelled = [t for t in members if t.status == "Cancelada"]
        if not members or any(t.trip_total < 0 for t in cancelled):
            continue
        included = sum(t.trip_total for t in members)
        excluded = sum(t.trip_total for t in members if t.status != "Cancelada")
        all_in = math.isclose(amount, included, abs_tol=0.005)
        all_out = math.isclose(amount, excluded, abs_tol=0.005)
        for trip in members:
            if trip.trip_total == 0:
                continue
            if trip.status != "Cancelada" and (all_in or all_out):
                decisions[trip.pcdp] = "Sim"
            elif trip.status == "Cancelada" and all_in != all_out:
                decisions[trip.pcdp] = "Sim" if all_in else "Não"
    return decisions
