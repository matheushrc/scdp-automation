"""Semantic manual imports must survive reordered legacy columns."""

import importlib.util
import unittest
from typing import Literal, cast

from openpyxl import Workbook

from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import make_trip

if importlib.util.find_spec("scdp_automation.xlsx_manual") is not None:
    from scdp_automation.xlsx_manual import (
        ManualValues,
        apply_base_manual_values,
        read_base_manual_values,
    )


class ManualImportTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(
            importlib.util.find_spec("scdp_automation.xlsx_manual"),
            "semantic manual importer is not implemented",
        )
        self.workbook = Workbook()
        self.addCleanup(self.workbook.close)
        self.base = self.workbook.active
        self.base.title = "BASE VIAGENS"

    def rows(self, headers, rows):
        self.base.append(headers)
        for row in rows:
            self.base.append(row)

    def test_reordered_and_aliased_headers_keep_full_pcdp_choices(self):
        for headers in (
            ["Descontar do curso?", "Código de débito", "PCDP", "Segmento"],
            [
                " Descontar do curso ",
                "Codigo de debito",
                "Número da Solicitação",
                "Outro",
            ],
        ):
            with self.subTest(headers=headers):
                self.base.delete_rows(1, self.base.max_row)
                self.rows(
                    headers,
                    [
                        ["Não", "PPGE", "123456/26-2B", "=1"],
                        ["Sim", "AGRONOMIA", "123456/26-1A", "old"],
                    ],
                )
                self.assertEqual(
                    read_base_manual_values(self.base),
                    {
                        "123456/26-2B": ManualValues("PPGE", "Não"),
                        "123456/26-1A": ManualValues("AGRONOMIA", "Sim"),
                    },
                )

    def test_blank_cells_clear_json_values(self):
        self.rows(
            ["PCDP", "Código de débito", "Descontar do curso?"],
            [["123456/26-1A", None, ""]],
        )
        first = make_trip("123456/26-1A")
        first.codigo_de_debito, first.descontar_do_curso = "AGRONOMIA", "Sim"
        second = make_trip("123456/26-2B")
        second.codigo_de_debito, second.descontar_do_curso = "PPGE", "Não"
        trips = [second, first]
        before = [trip.model_dump() for trip in trips]
        result = apply_base_manual_values(trips, read_base_manual_values(self.base))
        self.assertEqual(
            [(trip.codigo_de_debito, trip.descontar_do_curso) for trip in result],
            [("PPGE", "Não"), (None, None)],
        )
        self.assertEqual([trip.model_dump() for trip in trips], before)
        self.assertIsNot(result[0], second)
        self.assertIsNot(result[1], first)

    def test_invalid_headers_are_rejected(self):
        for headers in (
            ["PCDP", "PCDP", "Código de débito", "Descontar do curso?"],
            [
                "PCDP",
                "Número da Solicitação",
                "Código de débito",
                "Descontar do curso?",
            ],
            ["PCDP", "Código de débito", "Codigo de debito", "Descontar do curso?"],
            ["PCDP", "Código de débito", "Descontar do curso?", "Descontar do curso"],
            ["Código de débito", "Descontar do curso?"],
            ["PCDP", "Descontar do curso?"],
            ["PCDP", "Código de débito"],
        ):
            with self.subTest(headers=headers):
                self.base.delete_rows(1, self.base.max_row)
                self.rows(headers, [])
                with self.assertRaisesRegex(WorkbookValidationError, "BASE VIAGENS"):
                    read_base_manual_values(self.base)

    def test_invalid_rows_are_rejected_without_changes(self):
        for rows in (
            [["123456/26", "PPGE", "Sim"], ["123456/26", "PPGE", "Não"]],
            [[None, "PPGE", None]],
            [["123456/26", 12, None]],
            [["123456/26", "=1", None]],
            [["123456/26", None, "=1"]],
            [["123456/26", None, "sim"]],
            [["bad", None, None]],
            [[123456, None, None]],
        ):
            with self.subTest(rows=rows):
                self.base.delete_rows(1, self.base.max_row)
                self.rows(["PCDP", "Código de débito", "Descontar do curso?"], rows)
                before = list(self.base.values)
                with self.assertRaisesRegex(
                    WorkbookValidationError, "BASE VIAGENS.*[234]"
                ):
                    read_base_manual_values(self.base)
                self.assertEqual(list(self.base.values), before)

    def test_apply_rejects_missing_pcdps_and_keeps_inputs(self):
        trips = [make_trip()]
        values = {"123456/26-1A": ManualValues("PPGE", "Sim")}
        before = trips[0].model_dump()
        with self.assertRaises(WorkbookValidationError) as caught:
            apply_base_manual_values(trips, values)
        self.assertEqual(caught.exception.missing_pcdps, ("123456/26-1A",))
        self.assertEqual(trips[0].model_dump(), before)
        self.assertEqual(values["123456/26-1A"].codigo_de_debito, "PPGE")

    def test_apply_validates_pcdps_fields_and_duplicates(self):
        trip = make_trip()
        for trips, values in (
            ([trip, trip], {}),
            ([trip], {"bad": ManualValues(None, None)}),
            ([trip], {trip.numero_da_solicitacao: ManualValues("=1", None)}),
            ([trip], {trip.numero_da_solicitacao: ManualValues(cast(str, 12), None)}),
            (
                [trip],
                {
                    trip.numero_da_solicitacao: ManualValues(
                        None, cast(Literal["Sim", "Não"], "sim")
                    )
                },
            ),
            ([trip.model_copy(update={"numero_da_solicitacao": "bad"})], {}),
        ):
            with (
                self.subTest(values=values),
                self.assertRaises(WorkbookValidationError),
            ):
                apply_base_manual_values(trips, values)

    def test_blank_rows_and_unknown_columns_are_ignored(self):
        self.rows(
            ["PCDP", "Código de débito", "Descontar do curso?", "Segmento"],
            [[None, None, None, "=1"], ["123456/26", " PPGE ", " Não ", "=2"]],
        )
        self.assertEqual(
            read_base_manual_values(self.base),
            {"123456/26": ManualValues("PPGE", "Não")},
        )

    def test_apply_rejects_invalid_retained_manual_value(self):
        trip = make_trip()
        trip.codigo_de_debito = "=1"
        with self.assertRaises(WorkbookValidationError):
            apply_base_manual_values([trip], {})
