"""Offline CLI integration with synthetic annual files."""

import unittest
from contextlib import ExitStack, redirect_stderr, redirect_stdout
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from scdp_automation import cli, config, extrator, xlsx_output
from scdp_automation.relatorio import load_trips, save_json
from scdp_automation.xlsx_validation import BASE_HEADERS
from tests.support.workbooks import final_template_fixture, make_trip


class CliRecoveryTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template = self.root / "template.xlsx"
        final_template_fixture(self.template)
        self.checkpoint = self.root / "output" / "viagens_scdp_2026.json"
        self.workbook = self.root / "output" / "gastos_scdp_2026.xlsx"
        self.configuration = self.root / ".scdp-config.toml"
        self.configuration.write_text("ano = 2026\n")
        stack = self.enterContext(ExitStack())
        stack.enter_context(patch.object(extrator, "REPO_ROOT", self.root))
        stack.enter_context(
            patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.template)
        )
        # Configuration boundary points at a real synthetic local config.
        stack.enter_context(
            patch.object(
                extrator,
                "sync_extraction_config",
                side_effect=lambda: config.sync_extraction_config(
                    self.configuration, current_year=2026
                ),
            )
        )
        for name in ("prepare_chrome_profile", "open_browser", "login_only", "run"):
            stack.enter_context(
                patch.object(cli, name, side_effect=AssertionError(name))
            )
        stack.enter_context(
            patch.object(cli.asyncio, "run", side_effect=AssertionError("asyncio"))
        )

    def snapshot(self):
        return {
            p.relative_to(self.root): p.read_bytes()
            for p in self.root.rglob("*")
            if p.is_file()
        }

    def test_recreate_workbook_does_not_prepare_profile_or_start_browser(self):
        save_json(self.checkpoint, [make_trip()])
        xlsx_output.publish_workbook(load_trips(self.checkpoint), self.workbook)
        book = load_workbook(self.workbook)
        base = book["BASE VIAGENS"]
        base["Q2"] = "AGRONOMIA"
        base["R2"] = "Não"
        base["A1"] = "Número da Solicitação"
        base["Q1"] = "Codigo de debito"
        book["APOIO"]["D2"] = 987
        book.save(self.workbook)
        book.close()
        cli.main(["--recriar-planilha"])
        book = load_workbook(self.workbook)
        self.addCleanup(book.close)
        self.assertEqual(
            [cell.value for cell in book["BASE VIAGENS"][1]], list(BASE_HEADERS)
        )
        self.assertEqual(book["APOIO"]["D2"].value, 987)
        self.assertEqual(load_trips(self.checkpoint)[0].codigo_de_debito, "AGRONOMIA")
        self.assertEqual(book["BASE VIAGENS"]["R2"].value, "Não")

    def test_recreate_requires_existing_valid_checkpoint(self):
        for content in (None, "not JSON", '[{"numero_da_solicitacao": "invalid"}]'):
            with self.subTest(content=content):
                if content is not None:
                    self.checkpoint.parent.mkdir(exist_ok=True)
                    self.checkpoint.write_text(content)
                before = self.snapshot()
                error = StringIO()
                with redirect_stderr(error), self.assertRaises(SystemExit) as raised:
                    cli.main(["--recriar-planilha"])
                self.assertEqual(raised.exception.code, 1)
                self.assertIn(str(self.checkpoint), error.getvalue())
                self.assertIn("JSON existente", error.getvalue())
                self.assertNotIn("Traceback", error.getvalue())
                self.assertEqual(self.snapshot(), before)

    def test_recreate_rejects_conflicting_browser_actions(self):
        for flag in ("--login", "--abrir-navegador", "--selecionar-perfil-chrome"):
            with self.subTest(flag=flag):
                with (
                    redirect_stderr(StringIO()),
                    self.assertRaises(SystemExit) as raised,
                ):
                    cli.main(["--recriar-planilha", flag])
                self.assertEqual(raised.exception.code, 2)
                self.assertFalse(self.workbook.exists())

    def test_help_includes_recovery_without_profile(self):
        output = StringIO()
        with redirect_stdout(output), self.assertRaises(SystemExit) as raised:
            cli.main(["--help"])
        self.assertEqual(raised.exception.code, 0)
        self.assertIn("--recriar-planilha", output.getvalue())

    def test_unexpected_internal_error_remains_visible(self):
        save_json(self.checkpoint, [make_trip()])
        with (
            patch.object(
                cli,
                "recreate_workbook",
                create=True,
                side_effect=RuntimeError("internal bug"),
            ),
            self.assertRaisesRegex(RuntimeError, "internal bug"),
        ):
            cli.main(["--recriar-planilha"])

    def test_incompatible_manual_sheet_reports_reconciliation(self):
        save_json(self.checkpoint, [make_trip()])
        xlsx_output.publish_workbook(load_trips(self.checkpoint), self.workbook)
        book = load_workbook(self.workbook)
        book["APOIO"].insert_cols(8)
        book["APOIO"]["H1"] = "Transportes pago (R$)"
        book.save(self.workbook)
        book.close()
        before = self.snapshot()
        error = StringIO()
        with redirect_stderr(error), self.assertRaises(SystemExit) as raised:
            cli.main(["--recriar-planilha"])
        self.assertEqual(raised.exception.code, 1)
        self.assertIn(str(self.workbook), error.getvalue())
        self.assertIn("reconcil", error.getvalue())
        self.assertNotIn("começar novamente", error.getvalue())
        self.assertEqual(self.snapshot(), before)
