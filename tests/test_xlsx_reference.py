import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scdp_automation.xlsx_reference import read_reference
from tests.xlsx_fixtures import reference_fixture


class ReferenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "original.xlsx"
        reference_fixture(self.path)

    def test_deduplicates_and_recovers_classification_without_writing_source(self):
        before = self.path.read_bytes()
        data = read_reference(self.path)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(len(data.trips), 2)
        self.assertEqual(data.occurrences, 3)
        self.assertEqual(
            data.codes, {"999001/26": "AGRONOMIA", "999002/26-1C": "PPGEL +"}
        )
        first = data.trips[0]
        self.assertEqual(
            (
                first.daily_count,
                first.daily_amount,
                first.ticket_amount,
                first.trip_total,
            ),
            (2.5, 100, 50, 156),
        )
        self.assertEqual((first.discount_amount, first.restitution_amount), (5, -2))

    def test_imports_summary_budgets_and_scheduled_and_paid_transport(self):
        data = read_reference(self.path)
        inputs = data.inputs["AGRONOMIA"]
        self.assertEqual(
            (inputs.total, inputs.daily, inputs.transport), (1000, 400, 600)
        )
        self.assertEqual((inputs.scheduled, inputs.paid), (200, 150))
        self.assertEqual(data.rateio, (1 / 3, 1 / 3, 1 / 3))

    def test_conflicting_financial_duplicate_is_rejected(self):
        workbook = load_workbook(self.path)
        workbook["BD D&P"]["S15"] = 200
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "financeiros conflitantes"):
            read_reference(self.path)

    def test_conflicting_populated_classification_is_rejected(self):
        workbook = load_workbook(self.path)
        workbook["BD D&P"]["B11"] = "SEG 1 GRADUAÇÃO"
        workbook["BD D&P"]["C11"] = "MEDICINA"
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "classificações conflitantes"):
            read_reference(self.path)

    def test_missing_amount_is_not_zero(self):
        workbook = load_workbook(self.path)
        workbook["BD D&P"]["Q8"] = None
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "numérico"):
            read_reference(self.path)

    def test_subtotal_label_can_be_in_the_description_column(self):
        workbook = load_workbook(self.path)
        workbook["BD D&P"]["O8"] = None
        workbook["BD D&P"]["D8"] = "Sub-Total"
        workbook.save(self.path)
        workbook.close()
        self.assertEqual(read_reference(self.path).trips[0].daily_amount, 100)

    def test_duplicate_without_subtotal_label_must_match_explicit_original(self):
        workbook = load_workbook(self.path)
        workbook["BD D&P"]["O14"] = None
        workbook.save(self.path)
        workbook.close()
        self.assertEqual(len(read_reference(self.path).trips), 2)

    def test_rateio_must_reconcile_with_the_classified_trip_total(self):
        workbook = load_workbook(self.path)
        workbook["APOIO"]["J58"] = 53
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "rateio.*não reconcilia"):
            read_reference(self.path)
