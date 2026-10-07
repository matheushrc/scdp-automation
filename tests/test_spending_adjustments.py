import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import AsyncMock, patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from scdp_automation import extrator, xlsx_output
from scdp_automation.xlsx_models import summarize_trips
from tests.support.workbooks import (
    final_template_fixture,
    make_trip,
    write_existing_workbook,
)


class DateTests(unittest.IsolatedAsyncioTestCase):
    def test_dates_and_future_selection(self):
        trip = make_trip()
        trip.trechos = [trip.trechos[0]]
        trip.trechos[0].inicio = "06/10/2026"
        trip.trechos[0].termino = "08/10/2026"
        trip.data_da_ultima_verificacao = date(2026, 10, 4)
        trip.descricao_do_motivo_da_viagem = "Anterior"
        summary = summarize_trips([trip])[0]
        self.assertEqual(summary.start_date, date(2026, 10, 6))
        self.assertEqual(summary.end_date, date(2026, 10, 8))
        self.assertEqual(summary.verified_date, date(2026, 10, 4))
        self.assertTrue(extrator.needs_verification(trip, date(2026, 10, 5)))
        trip.trechos[0].inicio = "05/10/2026"
        self.assertFalse(extrator.needs_verification(trip, date(2026, 10, 5)))

    async def test_success_records_date_failure_keeps_old_date(self):
        trip = make_trip()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "checkpoint.json"
            with patch.object(
                extrator, "_consult_with_safe_failure", AsyncMock(return_value="Atual")
            ):
                await extrator.collect_pending_descriptions(
                    AsyncMock(), [trip], [trip], output
                )
            self.assertEqual(
                trip.data_da_ultima_verificacao,
                datetime.now(ZoneInfo("America/Sao_Paulo")).date(),
            )
            old_date = date(2026, 1, 1)
            trip.data_da_ultima_verificacao = old_date
            with (
                patch.object(
                    extrator,
                    "_consult_with_safe_failure",
                    AsyncMock(side_effect=RuntimeError("falha")),
                ),
                self.assertRaises(RuntimeError),
            ):
                await extrator.collect_pending_descriptions(
                    AsyncMock(), [trip], [trip], output
                )
            self.assertEqual(trip.data_da_ultima_verificacao, old_date)


class WorkbookAdjustmentsTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.reference = self.root / "reference.xlsx"
        final_template_fixture(self.reference)
        default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.reference)
        default.start()
        self.addCleanup(default.stop)

    def test_rateio_required_for_classified_expense_without_budget(self):
        from scdp_automation.xlsx_validation import (
            WorkbookValidationError,
            validate_workbook,
        )

        path = self.root / "final.xlsx"
        final_template_fixture(path)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        base, support = workbook["BASE VIAGENS"], workbook["APOIO"]
        rateio_row = next(
            r
            for r in range(2, support.max_row + 1)
            if support.cell(r, 1).value == "PPGEL +"
        )
        base["Q2"] = "PPGEL +"
        base["K2"] = 100
        for column in (4, 5, 7, 8, 9, 10):
            support.cell(rateio_row, column).value = None
        with self.assertRaisesRegex(WorkbookValidationError, "100%"):
            validate_workbook(workbook)
        for weights, valid in (
            ((0.5, 0.25, 0.25), True),
            ((0.5, 0.25, 0.20), False),
            ((0.5, 0.25, 0.2500000005), False),
            ((float("nan"), 0.25, 0.25), False),
            ((True, 0, 0), False),
        ):
            with self.subTest(weights=weights):
                for column, weight in zip((8, 9, 10), weights, strict=True):
                    support.cell(rateio_row, column).value = weight
                if valid:
                    validate_workbook(workbook)
                else:
                    with self.assertRaisesRegex(WorkbookValidationError, "100%"):
                        validate_workbook(workbook)
        base["K2"] = None
        for column in (8, 9, 10):
            support.cell(rateio_row, column).value = None
        validate_workbook(workbook)

    def test_decisions_and_dates_follow_complete_pcdp(self):
        current, candidate = self.root / "current.xlsx", self.root / "candidate.xlsx"
        write_existing_workbook(
            current,
            [make_trip("999001/26"), make_trip("999002/26-1C")],
            {"999001/26": "AGRONOMIA", "999002/26-1C": "PPGE"},
        )
        workbook = load_workbook(current)
        base = workbook["BASE VIAGENS"]
        self.assertEqual(base["R1"].value, "Descontar do curso?")
        base["R2"] = "Não"
        base["R3"] = "Sim"
        workbook.save(current)
        workbook.close()
        trips = [make_trip("999002/26-1C"), make_trip("999001/26")]
        xlsx_output.build_candidate(trips, current, candidate)
        workbook = load_workbook(candidate)
        self.addCleanup(workbook.close)
        base = workbook["BASE VIAGENS"]
        self.assertEqual([base.cell(r, 18).value for r in (2, 3)], ["Sim", "Não"])
        self.assertEqual(base["L2"].value.date(), date(2026, 3, 1))
        self.assertEqual(base["L2"].number_format, "dd/mm/yyyy")
        self.assertTrue(
            any(
                v.formula1 == '"Sim,Não"' and "R1048576" in str(v.sqref)
                for v in base.data_validations.dataValidation
            )
        )
        self.assertTrue(
            any(str(c.sqref) == "A2:R1048576" for c in base.conditional_formatting)
        )

    def test_support_order_history_group_and_manual_summary_survive(self):
        current, candidate = self.root / "current.xlsx", self.root / "candidate.xlsx"
        write_existing_workbook(
            current,
            [make_trip("999001/26"), make_trip("999002/26-1C")],
            {"999001/26": "AGRONOMIA", "999002/26-1C": "PPGE"},
        )
        workbook = load_workbook(current)
        support = workbook["APOIO"]
        self.assertEqual(
            [support.cell(1, c).value for c in (4, 5, 6)],
            [
                "Diárias e passagens distribuído (R$)",
                "Transportes distribuído (R$)",
                "Recurso total (R$)",
            ],
        )
        self.assertEqual(support["E3"].value, 600)
        groups = {
            support.cell(r, 1).value: support.cell(r, 2).value
            for r in range(2, support.max_row + 1)
        }
        self.assertIn("PPGH/PPGDH", groups)
        self.assertNotIn("PPGH", groups)
        self.assertNotIn("PPGDH", groups)
        self.assertNotEqual(groups["HISTÓRIA"], groups["PPGH/PPGDH"])
        support.append(
            [
                "NOVO",
                "Categoria nova",
                "SEG 1 GRADUAÇÃO",
                0,
                0,
                "=D40+E40",
                0,
                0,
                None,
            ]
        )
        workbook["RESUMO GASTOS"]["F8"] = "=SUM(1,2)"
        workbook.save(current)
        workbook.close()
        xlsx_output.build_candidate(
            [make_trip("999001/26"), make_trip("999002/26-1C")], current, candidate
        )
        workbook = load_workbook(candidate)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["RESUMO GASTOS"]["F8"].value, "=SUM(1,2)")
        self.assertEqual(workbook["APOIO"]["E3"].value, 600)
        self.assertEqual(
            workbook["APOIO"].cell(workbook["APOIO"].max_row, 1).value, "NOVO"
        )


class SpendingResultsTests(unittest.TestCase):
    def test_cancellation_decisions_and_restitution_use_collected_totals_once(self):
        import shutil

        from tests.support.recalculation import recalculate_workbook

        if not shutil.which("libreoffice"):
            self.skipTest("LibreOffice necessário")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, current, candidate = (
                root / name
                for name in ("template.xlsx", "current.xlsx", "candidate.xlsx")
            )
            final_template_fixture(template)
            trip = make_trip("111111/26")
            trip.situacao_da_viagem = "Cancelada"
            trip.restituicao_r = -44
            trip.total_da_viagem_r = 156
            write_existing_workbook(current, [trip], {"111111/26": "AGRONOMIA"})
            for decision, expected in ((None, "Pendente"), ("Não", 0), ("Sim", 156)):
                with self.subTest(decision=decision):
                    book = load_workbook(current)
                    book["BASE VIAGENS"]["R2"] = decision
                    book.save(current)
                    book.close()
                    candidate.unlink(missing_ok=True)
                    xlsx_output.build_candidate(
                        [trip], current, candidate, template_path=template
                    )
                    recalculate_workbook(candidate)
                    cached = load_workbook(candidate, data_only=True)
                    try:
                        self.assertEqual(cached["BASE VIAGENS"]["I2"].value, -44)
                        self.assertEqual(cached["BASE VIAGENS"]["K2"].value, 156)
                        self.assertEqual(cached["RESUMO GASTOS"]["F7"].value, expected)
                        self.assertEqual(cached["RESUMO GASTOS"]["F20"].value, expected)
                    finally:
                        cached.close()
