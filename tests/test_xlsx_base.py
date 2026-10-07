"""Recovery uses semantic BASE inputs and preserves the manual workbook container."""

import unittest
from copy import copy
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

from scdp_automation.xlsx_output import build_candidate
from scdp_automation.xlsx_validation import BASE_HEADERS, WorkbookValidationError
from tests.support.fidelity import assert_manual_sheets_preserved
from tests.support.workbooks import (
    final_template_fixture,
    make_trip,
    write_existing_workbook,
)


class BaseRecoveryTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.template = self.root / "template.xlsx"
        self.current = self.root / "current.xlsx"
        self.candidate = self.root / "candidate.xlsx"
        final_template_fixture(self.template)
        self.trips = [make_trip("123456/26-1C"), make_trip("123456/26-2C")]
        write_existing_workbook(
            self.current,
            self.trips,
            {"123456/26-1C": "AGRONOMIA", "123456/26-2C": "PPGE"},
        )

    def reorder(self, book):
        base = book["BASE VIAGENS"]
        rows = [[cell.value for cell in row] for row in base]
        # Keep the PCDP key in A; move Total, decision and code away from positions.
        order = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 17, 11, 12, 13, 14, 15, 10, 16]
        for r, values in enumerate(rows, 1):
            for c, old in enumerate(order, 1):
                base.cell(r, c).value = values[old]
        base["A1"] = "Número da Solicitação"
        base["R1"] = "Codigo de debito"
        base["K1"] = "Descontar do curso"
        base["P2"] = "=1"
        base["B1"].font = Font(name="Old", color="FF0000")
        base.data_validations.dataValidation = []
        book.defined_names["SCDPLayoutVersion"].attr_text = '"8"'
        book.defined_names["ViagensM"].attr_text = "'BASE VIAGENS'!$R$2:$R$999"

    def build(self, trips=None):
        build_candidate(
            trips or list(reversed(self.trips)),
            self.current,
            self.candidate,
            template_path=self.template,
        )

    def test_refresh_rebuilds_base_and_preserves_manual_sheets(self):
        book = load_workbook(self.current)
        book["BASE VIAGENS"]["R2"] = "Não"
        book["BASE VIAGENS"]["R3"] = "Sim"
        book["APOIO"]["D2"] = 12345
        book["APOIO"]["D2"].font = Font(name="Manual", color="AABBCC")
        book["APOIO"].column_dimensions["B"].width = 81
        book["RESUMO GASTOS"]["F7"] = "=SUM(ViagensK)"
        book["RESUMO GASTOS"]["F8"] = "=SUM(tblBaseViagens[Total da viagem (R$)])"
        self.reorder(book)
        book.save(self.current)
        book.close()
        before = self.current.read_bytes(), self.template.read_bytes()
        self.build()
        actual, source, template = (
            load_workbook(path)
            for path in (self.candidate, self.current, self.template)
        )
        for workbook in (actual, source, template):
            self.addCleanup(workbook.close)
        self.assertEqual(
            before, (self.current.read_bytes(), self.template.read_bytes())
        )
        self.assertEqual(
            tuple(cell.value for cell in actual["BASE VIAGENS"][1]), BASE_HEADERS
        )
        self.assertEqual(
            copy(actual["BASE VIAGENS"]["B1"].font),
            copy(template["BASE VIAGENS"]["B1"].font),
        )
        self.assertEqual(
            [actual["BASE VIAGENS"].cell(r, 17).value for r in (2, 3)],
            ["PPGE", "AGRONOMIA"],
        )
        self.assertEqual(
            [actual["BASE VIAGENS"].cell(r, 18).value for r in (2, 3)], ["Sim", "Não"]
        )
        managed = {
            name for name in template.defined_names if name.startswith("Viagens")
        } | {"SCDPLayoutVersion"}
        assert_manual_sheets_preserved(self, actual, source, excluded_names=managed)
        for name in managed:
            self.assertEqual(actual.defined_names[name], template.defined_names[name])

    def test_json_fields_populate_first_publication(self):
        self.current.unlink()
        trip = self.trips[0]
        trip.codigo_de_debito, trip.descontar_do_curso = "AGRONOMIA", "Sim"
        self.build([trip])
        book = load_workbook(self.candidate)
        self.addCleanup(book.close)
        self.assertEqual(book["BASE VIAGENS"]["Q2"].value, "AGRONOMIA")
        self.assertEqual(book["BASE VIAGENS"]["R2"].value, "Sim")
        self.assertIn("VLOOKUP", book["BASE VIAGENS"]["P2"].value)

    def test_cleared_output_inputs_override_json_and_json_only_trip_survives(self):
        self.trips[0].codigo_de_debito = "PPGE"
        self.trips[0].descontar_do_curso = "Sim"
        book = load_workbook(self.current)
        book["BASE VIAGENS"]["Q2"] = None
        book.save(self.current)
        book.close()
        added = make_trip("234567/26")
        added.codigo_de_debito, added.descontar_do_curso = "AGRONOMIA", "Não"
        self.build([self.trips[0], self.trips[1], added])
        book = load_workbook(self.candidate)
        self.addCleanup(book.close)
        self.assertIsNone(book["BASE VIAGENS"]["Q2"].value)
        self.assertIsNone(book["BASE VIAGENS"]["R2"].value)
        self.assertEqual(book["BASE VIAGENS"]["Q4"].value, "AGRONOMIA")
        self.assertEqual(book["BASE VIAGENS"]["R4"].value, "Não")

    def test_base_recovery_rejects_incompatible_manual_sheets(self):
        for kind, message in (
            ("apoio9", "APOIO"),
            ("name", "ApoioJ"),
            ("height", "R2"),
            ("budget", "APOIO"),
            ("decision", "Descontar"),
            ("code", "código"),
            ("finance", "financeiros"),
        ):
            with self.subTest(kind=kind):
                write_existing_workbook(self.current, self.trips)
                book = load_workbook(self.current)
                self.reorder(book)
                if kind == "apoio9":
                    book["APOIO"].insert_cols(8)
                    book["APOIO"]["H1"] = "Transportes pago (R$)"
                elif kind == "name":
                    book.defined_names["ApoioJ"].attr_text = "APOIO!I2:I999"
                elif kind == "height":
                    book["RESUMO GASTOS"]["R2"] = "=1"
                elif kind == "budget":
                    book["APOIO"]["D2"] = "inf"
                elif kind == "decision":
                    book["BASE VIAGENS"]["K2"] = "Talvez"
                elif kind == "code":
                    book["BASE VIAGENS"]["R2"] = "FORA"
                else:
                    book["BASE VIAGENS"]["Q2"] = "=1/0"
                book.save(self.current)
                book.close()
                before = self.current.read_bytes()
                with self.assertRaisesRegex(WorkbookValidationError, message):
                    self.build()
                self.assertEqual(before, self.current.read_bytes())
                self.assertFalse(self.candidate.exists())

    def test_manual_physical_base_reference_needs_reconciliation(self):
        for kind in ("formula", "name", "validation", "conditional"):
            with self.subTest(kind=kind):
                write_existing_workbook(self.current, self.trips)
                book = load_workbook(self.current)
                self.reorder(book)
                formula = "='BASE VIAGENS'!$Q$2"
                if kind == "formula":
                    book["APOIO"]["K12"] = formula
                elif kind == "name":
                    book.defined_names.add(
                        DefinedName("ManualTotal", attr_text=formula[1:])
                    )
                elif kind == "validation":
                    from openpyxl.worksheet.datavalidation import DataValidation

                    rule = DataValidation(type="custom", formula1=formula)
                    rule.add("D2")
                    book["APOIO"].add_data_validation(rule)
                else:
                    from openpyxl.formatting.rule import FormulaRule

                    book["APOIO"].conditional_formatting.add(
                        "D2", FormulaRule(formula=[formula[1:]])
                    )
                book.save(self.current)
                book.close()
                before = self.current.read_bytes()
                with self.assertRaisesRegex(WorkbookValidationError, "reconcil"):
                    self.build()
                self.assertEqual(before, self.current.read_bytes())
                self.assertFalse(self.candidate.exists())

    def test_reordered_total_keeps_old_rateio_trigger(self):
        book = load_workbook(self.current)
        book["BASE VIAGENS"]["Q2"] = "PPGEL +"
        support = book["APOIO"]
        row = next(
            r
            for r in range(2, support.max_row + 1)
            if support.cell(r, 1).value == "PPGEL +"
        )
        for col in (4, 5, 7, 8, 9, 10):
            support.cell(row, col).value = None
        self.reorder(book)
        book.save(self.current)
        book.close()
        # Candidate totals would no longer activate allocation; old source must still fail.
        for trip in self.trips:
            trip.total_da_viagem_r = 0
        with self.assertRaisesRegex(WorkbookValidationError, "100%"):
            self.build()

    def test_base_reconstruction_copies_styles_and_metadata_without_style_indices(self):
        from scdp_automation.xlsx_base import rebuild_base_from_template

        source, template = load_workbook(self.current), load_workbook(self.template)
        self.addCleanup(source.close)
        self.addCleanup(template.close)
        support = source["APOIO"]
        support["D2"].font = Font(name="DifferentRegistry", color="112233")
        base = template["BASE VIAGENS"]
        base["B1"].font = Font(name="TemplateUnique", color="998877")
        base.freeze_panes = "C2"
        base.merge_cells("A5:B5")
        base.print_area = "A1:R5"
        base.print_title_rows = "1:1"
        base.row_dimensions[5].height = 42
        base.column_dimensions["C"].font = copy(base["B1"].font)
        base.defined_names.add(
            DefinedName("LocalBase", attr_text="'BASE VIAGENS'!$A$1", localSheetId=0)
        )
        rebuilt = rebuild_base_from_template(source, template)
        self.assertIs(source["APOIO"], support)
        self.assertEqual(source.sheetnames, ["BASE VIAGENS", "APOIO", "RESUMO GASTOS"])
        self.assertEqual(copy(rebuilt["B1"].font), copy(base["B1"].font))
        self.assertEqual(
            copy(rebuilt.column_dimensions["C"].font),
            copy(base.column_dimensions["C"].font),
        )
        for attribute in (
            "freeze_panes",
            "print_area",
            "print_title_rows",
            "data_validations",
            "sheet_properties",
            "page_setup",
            "page_margins",
        ):
            self.assertEqual(getattr(rebuilt, attribute), getattr(base, attribute))
        self.assertEqual(set(map(str, rebuilt.merged_cells)), {"A5:B5"})
        self.assertEqual(rebuilt.row_dimensions[5].height, 42)
        self.assertEqual(dict(rebuilt.defined_names), dict(base.defined_names))

    def test_importer_accepts_trimmed_pcdp_alias_after_key_column_moves(self):
        book = load_workbook(self.current)
        base = book["BASE VIAGENS"]
        for row in range(1, base.max_row + 1):
            a, b = base.cell(row, 1).value, base.cell(row, 2).value
            base.cell(row, 1).value, base.cell(row, 2).value = b, a
        base["B1"] = " Número da Solicitação "
        book.save(self.current)
        book.close()
        self.build()
        actual = load_workbook(self.candidate)
        self.addCleanup(actual.close)
        self.assertEqual(actual["BASE VIAGENS"]["A2"].value, "123456/26-2C")
        self.assertEqual(actual["BASE VIAGENS"]["Q2"].value, "PPGE")

    def test_indirect_physical_reference_is_not_silently_changed(self):
        book = load_workbook(self.current)
        self.reorder(book)
        book["APOIO"]["K12"] = "=INDIRECT(\"'BASE VIAGENS'!Q2\")"
        book.save(self.current)
        book.close()
        before = self.current.read_bytes()
        with self.assertRaisesRegex(WorkbookValidationError, "reconcil"):
            self.build()
        self.assertEqual(self.current.read_bytes(), before)
        self.assertFalse(self.candidate.exists())

    def test_structured_base_reference_to_renamed_header_needs_reconciliation(self):
        book = load_workbook(self.current)
        self.reorder(book)
        book["APOIO"]["K12"] = "=COUNTA(tblBaseViagens[Codigo de debito])"
        book.save(self.current)
        book.close()
        before = self.current.read_bytes()
        with self.assertRaisesRegex(WorkbookValidationError, "reconcil"):
            self.build()
        self.assertEqual(self.current.read_bytes(), before)
        self.assertFalse(self.candidate.exists())

    def test_compatible_custom_base_local_name_survives_reconstruction(self):
        book = load_workbook(self.current)
        self.reorder(book)
        book["BASE VIAGENS"].defined_names.add(
            DefinedName(
                "ManualKey", attr_text="'BASE VIAGENS'!$A$2:$A$3", localSheetId=0
            )
        )
        book.save(self.current)
        book.close()
        self.build()
        actual = load_workbook(self.candidate)
        self.addCleanup(actual.close)
        self.assertIn("ManualKey", actual["BASE VIAGENS"].defined_names)
        self.assertEqual(
            actual["BASE VIAGENS"].defined_names["ManualKey"].attr_text,
            "'BASE VIAGENS'!$A$2:$A$3",
        )

    def test_incompatible_custom_base_local_name_needs_reconciliation(self):
        book = load_workbook(self.current)
        self.reorder(book)
        book["BASE VIAGENS"].defined_names.add(
            DefinedName(
                "ManualTotal", attr_text="'BASE VIAGENS'!$Q$2:$Q$3", localSheetId=0
            )
        )
        book.save(self.current)
        book.close()
        before = self.current.read_bytes()
        with self.assertRaisesRegex(WorkbookValidationError, "ManualTotal.*reconcil"):
            self.build()
        self.assertEqual(self.current.read_bytes(), before)
        self.assertFalse(self.candidate.exists())

    def test_unqualified_custom_base_local_reference_needs_reconciliation(self):
        book = load_workbook(self.current)
        self.reorder(book)
        book["BASE VIAGENS"].defined_names.add(
            DefinedName("ManualTotal", attr_text="$Q$2:$Q$3", localSheetId=0)
        )
        book.save(self.current)
        book.close()
        before = self.current.read_bytes()
        with self.assertRaisesRegex(WorkbookValidationError, "ManualTotal.*reconcil"):
            self.build()
        self.assertEqual(self.current.read_bytes(), before)
        self.assertFalse(self.candidate.exists())

    def test_candidate_keeps_template_base_presentation(self):
        template = load_workbook(self.template)
        base = template["BASE VIAGENS"]
        base["B1"].font = Font(name="TemplatePresentation", color="123ABC")
        base.column_dimensions["B"].width = 73
        base.freeze_panes = "C2"
        base.row_dimensions[1].height = 82
        template.save(self.template)
        template.close()
        self.build()
        actual = load_workbook(self.candidate)
        self.addCleanup(actual.close)
        base = actual["BASE VIAGENS"]
        self.assertEqual(base["B1"].font.name, "TemplatePresentation")
        self.assertEqual(base.column_dimensions["B"].width, 73)
        self.assertEqual(base.freeze_panes, "C2")
        self.assertEqual(base.row_dimensions[1].height, 82)
        self.assertEqual(base["L2"].number_format, "dd/mm/yyyy")
        self.assertIn('"R$"', base["K2"].number_format)

    def test_unsupported_template_drawing_is_rejected(self):
        from openpyxl.chart import BarChart

        from scdp_automation.xlsx_base import rebuild_base_from_template

        source, template = load_workbook(self.current), load_workbook(self.template)
        self.addCleanup(source.close)
        self.addCleanup(template.close)
        template["BASE VIAGENS"].add_chart(BarChart(), "A10")
        old = source["BASE VIAGENS"]
        with self.assertRaisesRegex(WorkbookValidationError, "suportad"):
            rebuild_base_from_template(source, template)
        self.assertIs(source["BASE VIAGENS"], old)
