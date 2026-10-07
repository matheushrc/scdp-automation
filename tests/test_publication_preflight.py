"""Extraction source failures must precede checkpoint and browser access."""

import unittest
from argparse import Namespace
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from scdp_automation import extrator, xlsx_output
from scdp_automation.output_history import OutputHistory
from scdp_automation.relatorio import load_trips, save_json
from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import final_template_fixture, make_trip


class PublicationPreflightTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template = self.root / "template.xlsx"
        final_template_fixture(self.template)
        self.checkpoint = self.root / "viagens_scdp_2026.json"
        self.workbook = self.root / "gastos_scdp_2026.xlsx"
        self.trips = [make_trip()]
        save_json(self.checkpoint, self.trips)
        xlsx_output.publish_workbook(
            self.trips, self.workbook, template_path=self.template
        )
        default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.template)
        default.start()
        self.addCleanup(default.stop)

    def invalidate(self, failure):
        if failure == "missing template":
            self.template.unlink()
        elif failure == "orphan workbook":
            self.checkpoint.unlink()
        else:
            path = self.template if failure == "old template" else self.workbook
            book = load_workbook(path)
            book.defined_names["SCDPLayoutVersion"].attr_text = '"8"'
            book.save(path)
            book.close()

    def snapshot(self):
        return {
            path.relative_to(self.root): path.read_bytes()
            for path in self.root.rglob("*")
            if path.is_file()
        }

    def test_coordinator_rejects_bad_sources_before_checkpoint_or_history_writes(self):
        for failure in (
            "missing template",
            "old template",
            "old output",
            "orphan workbook",
        ):
            with self.subTest(failure=failure):
                # Restore the synthetic pair between independent failure cases.
                final_template_fixture(self.template)
                save_json(self.checkpoint, self.trips)
                self.workbook.unlink()
                xlsx_output.publish_workbook(self.trips, self.workbook)
                self.invalidate(failure)
                before = self.snapshot()
                with (
                    OutputHistory(self.checkpoint, self.workbook) as history,
                    self.assertRaises(WorkbookValidationError) as raised,
                ):
                    extrator.save_checkpoint_and_publish(
                        [make_trip(), make_trip("000002/26")],
                        self.checkpoint,
                        self.workbook,
                        history=history,
                    )
                self.assertEqual(self.snapshot(), before)
                if failure == "orphan workbook":
                    self.assertIn("reconcil", str(raised.exception).lower())

    async def test_run_rejects_bad_sources_before_loading_checkpoint_or_browser(self):
        for failure in (
            "missing template",
            "old template",
            "old output",
            "orphan workbook",
        ):
            with self.subTest(failure=failure):
                final_template_fixture(self.template)
                save_json(self.checkpoint, self.trips)
                self.workbook.unlink()
                xlsx_output.publish_workbook(self.trips, self.workbook)
                self.invalidate(failure)
                before = self.snapshot()
                args = Namespace(
                    output=self.checkpoint, workbook=self.workbook, ano=2026, limite=0
                )
                with (
                    patch.object(extrator, "configure_logging"),
                    patch.object(extrator, "parse_args", return_value=args),
                    patch.object(
                        extrator,
                        "load_trips",
                        side_effect=AssertionError(
                            "checkpoint accessed before rejection"
                        ),
                    ),
                    patch.object(
                        extrator,
                        "async_playwright",
                        side_effect=AssertionError("browser accessed before rejection"),
                    ),
                    self.assertRaises(WorkbookValidationError),
                ):
                    await extrator.run([])
                self.assertEqual(self.snapshot(), before)

    def test_json_only_interrupted_publication_can_resume(self):
        self.workbook.unlink()
        with OutputHistory(self.checkpoint, self.workbook) as history:
            extrator.save_checkpoint_and_publish(
                self.trips, self.checkpoint, self.workbook, history=history
            )
        self.assertEqual(load_trips(self.checkpoint), self.trips)
        self.assertTrue(self.workbook.is_file())
        self.assertFalse((self.root / "backup").exists())

    def test_new_rateio_expense_error_preserves_collected_checkpoint(self):
        book = load_workbook(self.workbook)
        book["BASE VIAGENS"]["Q2"] = "PPGEL +"
        book["BASE VIAGENS"]["K2"] = 0
        for column in (4, 5, 7):
            book["APOIO"].cell(7, column).value = 0
        for column in (8, 9, 10):
            book["APOIO"].cell(7, column).value = None
        book.save(self.workbook)
        book.close()
        save_json(self.checkpoint, [])
        before = self.workbook.read_bytes()
        with (
            OutputHistory(self.checkpoint, self.workbook) as history,
            self.assertRaises(WorkbookValidationError),
        ):
            extrator.save_checkpoint_and_publish(
                self.trips, self.checkpoint, self.workbook, history=history
            )
        self.assertEqual(load_trips(self.checkpoint), self.trips)
        self.assertEqual(self.workbook.read_bytes(), before)
        self.assertEqual(len(list((self.root / "backup").glob("*.history.*"))), 2)

    def test_missing_pcdp_publication_error_preserves_new_checkpoint(self):
        before = self.workbook.read_bytes()
        with (
            OutputHistory(self.checkpoint, self.workbook) as history,
            self.assertRaises(WorkbookValidationError) as raised,
        ):
            extrator.save_checkpoint_and_publish(
                [], self.checkpoint, self.workbook, history=history
            )
        self.assertEqual(raised.exception.missing_pcdps, ("123456/26-2B",))
        self.assertEqual(load_trips(self.checkpoint), [])
        self.assertEqual(self.workbook.read_bytes(), before)
        self.assertEqual(len(list((self.root / "backup").glob("*.history.*"))), 2)
