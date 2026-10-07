"""Read-only validation of the final editable spending workbook."""

from __future__ import annotations

import math
import re
from collections.abc import Mapping
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from scdp_automation.xlsx_manual import ManualValues

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from openpyxl.workbook.workbook import Workbook

BASE_HEADERS = (
    "PCDP",
    "Proposto",
    "Situação",
    "Quantidade de diárias",
    "Diárias (R$)",
    "Passagens (R$)",
    "Adicional (R$)",
    "Descontos (R$)",
    "Restituição (R$)",
    "Reembolso (R$)",
    "Total da viagem (R$)",
    "Data de início da viagem",
    "Data de término da viagem",
    "Data da última verificação",
    "Descrição do pedido",
    "Segmento",
    "Código de débito",
    "Descontar do curso?",
)

SUPPORT_HEADERS = (
    "Código de débito",
    "Nome por extenso",
    "Segmento",
    "Diárias e passagens distribuído (R$)",
    "Transportes distribuído (R$)",
    "Recurso total (R$)",
    "Transportes agendado (R$)",
    "Rateio para PPGE (%)",
    "Rateio para PPGEL (%)",
    "Rateio para PPGH (%)",
    "Total utilizado por categoria (R$)",
)


class WorkbookValidationError(ValueError):
    """The workbook violates the final template or financial input contract."""

    def __init__(self, message: str, *, missing_pcdps: tuple[str, ...] = ()) -> None:
        super().__init__(message)
        self.missing_pcdps = missing_pcdps


def same_formula(actual: object, expected: str) -> bool:
    """Compare formulas while allowing office applications' optional sheet quotes."""
    if not isinstance(actual, str) or not actual.startswith("="):
        return False

    def tokens(formula: str) -> list[tuple[str, str, str]]:
        return [
            (token.type, token.subtype, token.value.replace("'APOIO'!", "APOIO!"))
            for token in Tokenizer(formula).items
        ]

    return tokens(actual) == tokens(expected)


def segment_formula(row: int) -> str:
    return f'=IF($Q{row}="","",IFERROR(VLOOKUP($Q{row},ApoioCatalogo,3,FALSE),""))'


def _named_range(sheet: str, column: str, height_row: int, width: int = 1) -> str:
    return f"OFFSET('{sheet}'!${column}$1,1,0,MAX(1,'RESUMO GASTOS'!$R${height_row}),{width})"


def validate_workbook(workbook: Workbook) -> None:
    """Reject old layouts and invalid inputs without changing any workbook content."""
    _validate_contract(workbook)


def _validate_contract(
    workbook: Workbook, *, manual_values: Mapping[str, ManualValues] | None = None
) -> None:
    preserved = manual_values is not None
    if workbook.sheetnames != ["BASE VIAGENS", "APOIO", "RESUMO GASTOS"]:
        raise WorkbookValidationError(
            "O workbook não contém as três worksheets esperadas."
        )
    version = workbook.defined_names.get("SCDPLayoutVersion")
    if not preserved and (version is None or version.attr_text != '"10"'):
        raise WorkbookValidationError(
            f"Versão do layout inválida: esperada 10; observada "
            f"{version.attr_text if version is not None else 'ausente (marcador SCDPLayoutVersion)'}. "
            "Use o template final versão 10."
        )
    base, support = workbook["BASE VIAGENS"], workbook["APOIO"]
    for sheet, headers in (
        ((support, SUPPORT_HEADERS),)
        if preserved
        else ((base, BASE_HEADERS), (support, SUPPORT_HEADERS))
    ):
        if tuple(cell.value for cell in sheet[1]) != headers:
            raise WorkbookValidationError(
                f"As colunas de {sheet.title} foram alteradas."
            )
    tables = [(support, "tblApoioDebito", 11)]
    if not preserved:
        tables.append((base, "tblBaseViagens", 18))
    for sheet, name, last_column in tables:
        if name not in sheet.tables:
            raise WorkbookValidationError("Tabela necessária ausente.")
        min_column, min_row, max_column, max_row = range_boundaries(
            sheet.tables[name].ref
        )
        if (min_column, min_row, max_column) != (1, 1, last_column) or max_row < 2:
            raise WorkbookValidationError(f"A tabela {sheet.title} está malformada.")
    # OFFSET names depend on these three dynamic heights. A stale constant can
    # omit valid financial rows even when the names themselves are unchanged.
    for row, sheet, key in (
        (1, "APOIO", "A"),
        (2, "BASE VIAGENS", "A"),
        (3, "RESUMO GASTOS", "B"),
    ):
        expected = f"""=MAX(1,IFERROR(LOOKUP(2,1/('{sheet}'!${key}:${key}<>""),ROW('{sheet}'!${key}:${key}))-1,1))"""
        if not same_formula(workbook["RESUMO GASTOS"].cell(row, 18).value, expected):
            raise WorkbookValidationError(
                f"Altura do intervalo essencial ausente ou inválida: RESUMO GASTOS!R{row}."
            )
    required_names = {
        "CodigosDebito": _named_range("APOIO", "A", 1),
        "ApoioCatalogo": _named_range("APOIO", "A", 1, 3),
    }
    # Names are the final template's formula API, independent of category rows.
    for prefix, sheet, columns, height in (
        (
            "Apoio",
            "APOIO",
            {
                "A": "A",
                "B": "B",
                "C": "C",
                "D": "D",
                "E": "E",
                "F": "F",
                "G": "G",
                "J": "H",
                "K": "I",
                "L": "J",
                "M": "K",
            },
            1,
        ),
        (
            "Viagens",
            "BASE VIAGENS",
            {
                "A": "A",
                "B": "B",
                "C": "C",
                "D": "D",
                "E": "E",
                "F": "F",
                "G": "G",
                "H": "H",
                "I": "I",
                "J": "J",
                "K": "K",
                "L": "P",
                "M": "Q",
                "N": "R",
                "O": "L",
                "P": "M",
                "Q": "N",
            },
            2,
        ),
    ):
        if preserved and prefix == "Viagens":
            continue
        for name, column in columns.items():
            required_names[prefix + name] = _named_range(sheet, column, height)
    for column in "BCDEFGHIJKLMNOP":
        required_names["Resumo" + column] = _named_range("RESUMO GASTOS", column, 3)
    for name, expected in required_names.items():
        actual = workbook.defined_names.get(name)
        if actual is None or not same_formula("=" + actual.attr_text, "=" + expected):
            raise WorkbookValidationError(
                f"Intervalo nomeado essencial ausente ou inválido: {name}."
            )
    if not preserved:
        last_base_row = max(
            2, base.max_row, range_boundaries(base.tables["tblBaseViagens"].ref)[3]
        )
        for formula, column in (("CodigosDebito", 17), ('"Sim,Não"', 18)):
            intervals = sorted(
                (area.min_row, area.max_row)
                for validation in base.data_validations.dataValidation
                if validation.type == "list"
                and validation.formula1 in (formula, "=" + formula)
                for area in validation.sqref.ranges
                if area.min_col <= column <= area.max_col
            )
            covered = 1
            for first, last in intervals:
                if first > covered + 1:
                    break
                covered = max(covered, last)
            if covered < last_base_row:
                raise WorkbookValidationError(
                    "Validação de escolhas em BASE VIAGENS ausente."
                )

    financial_columns = (
        _financial_columns(base)
        if preserved
        else {header: index for index, header in enumerate(BASE_HEADERS[3:11], 4)}
    )
    pcdp_column = _pcdp_column(base) if preserved else 1
    total_column = financial_columns["Total da viagem (R$)"]
    seen_codes: set[str] = set()
    for row in range(2, support.max_row + 1):
        code = support.cell(row, 1).value
        if code in (None, ""):
            continue
        if (
            not isinstance(code, str)
            or code != code.strip()
            or code in seen_codes
            or any(c in code for c in "*?~")
        ):
            raise WorkbookValidationError("Código inválido ou duplicado em APOIO.")
        seen_codes.add(code)
        for col in (4, 5, 7):
            value = support.cell(row, col).value
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise WorkbookValidationError(
                    "Orçamento e transporte em APOIO precisam de valor numérico."
                )
        weights = [support.cell(row, col).value for col in (8, 9, 10)]
        if code == "PPGEL +":
            required = any(support.cell(row, col).value for col in (4, 5, 7)) or any(
                (
                    manual_values[
                        str(base.cell(r, pcdp_column).value).strip()
                    ].codigo_de_debito
                    == code
                    if preserved
                    and manual_values is not None
                    and str(base.cell(r, pcdp_column).value).strip() in manual_values
                    else base.cell(r, 17).value == code
                )
                and base.cell(r, total_column).value
                for r in range(2, base.max_row + 1)
            )
            if (required or any(w is not None for w in weights)) and (
                any(
                    isinstance(w, bool)
                    or not isinstance(w, (int, float))
                    or not math.isfinite(w)
                    or not 0 <= w <= 1
                    for w in weights
                )
                or not math.isclose(sum(weights), 1, rel_tol=0, abs_tol=1e-10)
            ):
                raise WorkbookValidationError(
                    "O rateio PPGEL + precisa totalizar 100%."
                )
    seen_pcdps: set[str] = set()
    for row in range(2, base.max_row + 1):
        pcdp = base.cell(row, pcdp_column).value
        if preserved and isinstance(pcdp, str):
            pcdp = pcdp.strip()
        if pcdp in (None, ""):
            continue
        if not isinstance(pcdp, str) or pcdp in seen_pcdps:
            raise WorkbookValidationError(
                "A base contém PCDPs inválidas ou duplicadas."
            )
        seen_pcdps.add(pcdp)
        manual = manual_values.get(pcdp) if manual_values is not None else None
        code = (
            manual.codigo_de_debito if manual is not None else base.cell(row, 17).value
        )
        if code not in (None, "") and code not in seen_codes:
            raise WorkbookValidationError(
                "A base contém código de débito desconhecido."
            )
        decision = (
            manual.descontar_do_curso
            if manual is not None
            else base.cell(row, 18).value
        )
        if decision not in (None, "", "Sim", "Não"):
            raise WorkbookValidationError("Decisão de desconto inválida na base.")
        for col in financial_columns.values():
            value = base.cell(row, col).value
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise WorkbookValidationError(
                    "A BASE precisa de valores financeiros numéricos e finitos."
                )


def _pcdp_column(base) -> int:
    columns = [
        cell.column
        for cell in base[1]
        if isinstance(cell.value, str)
        and cell.value.strip() in ("PCDP", "Número da Solicitação")
    ]
    if len(columns) != 1:
        raise WorkbookValidationError(
            "BASE VIAGENS: cabeçalho PCDP ausente ou ambíguo."
        )
    return columns[0]


def _financial_columns(base) -> dict[str, int]:
    columns: dict[str, int] = {}
    expected = set(BASE_HEADERS[3:11])
    for cell in base[1]:
        if cell.value in expected:
            if cell.value in columns:
                raise WorkbookValidationError(
                    f"BASE VIAGENS: campo financeiro ambíguo: {cell.value}."
                )
            columns[cell.value] = cell.column
        elif isinstance(cell.value, str) and "(R$)" in cell.value:
            raise WorkbookValidationError(
                f"BASE VIAGENS: campo financeiro sem mapeamento: {cell.value}."
            )
    missing = expected - columns.keys()
    if missing:
        raise WorkbookValidationError(
            "BASE VIAGENS: campos financeiros sem mapeamento: "
            + ", ".join(sorted(missing))
        )
    return columns


def _validate_manual_references(workbook: Workbook) -> None:
    from scdp_automation.xlsx_base import MANAGED_BASE_NAMES

    aliases = {
        "Número da Solicitação": "PCDP",
        "Codigo de debito": "Código de débito",
        "Descontar do curso": "Descontar do curso?",
    }
    original_headers = [cell.value for cell in workbook["BASE VIAGENS"][1]]
    old = [aliases.get(header, header) for header in original_headers]
    changed = {
        column
        for column in range(1, max(len(old), len(BASE_HEADERS)) + 1)
        if (old[column - 1] if column <= len(old) else None)
        != (BASE_HEADERS[column - 1] if column <= len(BASE_HEADERS) else None)
    }
    if not changed and tuple(original_headers) == BASE_HEADERS:
        return

    def check(value, context, *, base_scope=False):
        if not isinstance(value, str):
            return
        formula = value if value.startswith("=") else "=" + value
        for token in Tokenizer(formula).items:
            if token.type != "OPERAND" or token.subtype not in ("RANGE", "TEXT"):
                continue
            reference = token.value
            if token.subtype == "TEXT":
                reference = reference[1:-1].replace('""', '"')
                if "!" not in reference:
                    continue
            if reference.lower().startswith("tblbaseviagens"):
                headers = [
                    header.lstrip("@")
                    for header in re.findall(r"\[([^\[\]]+)\]", reference)
                    if not header.startswith("#")
                ]
                if (
                    not headers
                    or ":" in reference
                    or any(
                        header not in BASE_HEADERS or header not in original_headers
                        for header in headers
                    )
                ):
                    raise WorkbookValidationError(
                        f"{context}: referência estruturada à BASE VIAGENS incompatível; reconcilie a fórmula manual antes de reconstruir."
                    )
                continue
            if base_scope and re.fullmatch(
                r"(?:\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d+:\$?\d+)",
                reference,
            ):
                reference = "'BASE VIAGENS'!" + reference
            if "BASE VIAGENS" not in reference.upper():
                continue
            match = re.fullmatch(
                r"(?:'BASE VIAGENS'|BASE VIAGENS)!([\w$]+)(?::([\w$]+))?",
                reference,
                re.IGNORECASE,
            )
            if match:
                bounds = []
                for endpoint in (match[1], match[2] or match[1]):
                    column = re.match(r"\$?([A-Za-z]+)", endpoint)
                    bounds.append(
                        column_index_from_string(column[1]) if column else None
                    )
                first, last = bounds
                affected = (
                    any(first <= col <= last for col in changed)
                    if first is not None and last is not None
                    else True
                )
            else:
                affected = True
            if affected:
                raise WorkbookValidationError(
                    f"{context}: referência física à BASE VIAGENS mudou de significado; reconcilie a fórmula manual antes de reconstruir."
                )

    for name, definition in workbook.defined_names.items():
        if name not in MANAGED_BASE_NAMES | {"SCDPLayoutVersion"}:
            check(definition.attr_text, f"Intervalo {name}")
    for definition in workbook["BASE VIAGENS"].defined_names.values():
        if definition.name not in MANAGED_BASE_NAMES | {"SCDPLayoutVersion"}:
            check(
                definition.attr_text,
                f"BASE VIAGENS: intervalo {definition.name}",
                base_scope=True,
            )
    for sheet in (workbook["APOIO"], workbook["RESUMO GASTOS"]):
        for definition in sheet.defined_names.values():
            check(definition.attr_text, f"{sheet.title}: intervalo {definition.name}")
        for row in sheet:
            for cell in row:
                # R2 is an automation-owned height formula, checked exactly above;
                # its physical reference deliberately targets the rebuilt BASE.
                if sheet.title == "RESUMO GASTOS" and cell.coordinate == "R2":
                    continue
                if cell.data_type == "f":
                    check(cell.value, f"{sheet.title}!{cell.coordinate}")
        for validation in sheet.data_validations.dataValidation:
            for formula in (validation.formula1, validation.formula2):
                check(formula, f"{sheet.title}: validação {validation.sqref}")
        for area in sheet.conditional_formatting:
            for rule in sheet.conditional_formatting[area]:
                for formula in rule.formula or []:
                    check(formula, f"{sheet.title}: formatação {area.sqref}")


def validate_preserved_sheets(
    workbook: Workbook, manual_values: Mapping[str, ManualValues]
) -> None:
    """Validate the preserved contract independently of historical BASE layout."""
    _validate_contract(workbook, manual_values=manual_values)
    _validate_manual_references(workbook)
