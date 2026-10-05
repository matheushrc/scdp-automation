"""Original spending-summary presentation over simplified workbook inputs."""

from __future__ import annotations

import math
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from scdp_automation.xlsx_models import DEBIT_CATEGORIES
from scdp_automation.xlsx_reference import ReferenceData

LAYOUT_NAME = "SCDPLayoutVersion"
INPUT_HEADERS = (
    "Código de débito",
    "Nome por extenso",
    "Segmento",
    "Diárias e passagens distribuído (R$)",
    "Recurso total (R$)",
    "Transportes distribuído (R$)",
    "Transportes agendado (R$)",
    "Transportes pago (R$)",
    "Grupo no resumo",
    "Rateio para PPGE (%)",
    "Rateio para PPGEL (%)",
    "Rateio para PPGH (%)",
)
MAIN_ROWS = (*range(7, 22), *range(26, 36), 41, 43)
EXTRA_ROWS = (62, 63)
DETAIL_CODES = (
    "DIREÇÃO",
    "DIREÇÃO - AGAS",
    "DIREÇÃO - Banca Libras",
    "DIREÇÃO - CAAEX",
    "DIREÇÃO - Empr Junior",
    "DIREÇÃO - StartUp Summit",
    "DIREÇÃO - StartUp Weekend",
    "DIREÇÃO - Sunset",
)


def transport_formula(row: int) -> str:
    return f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),E{row}-D{row},"")'


def _copy_cell(source, target) -> None:
    target.value = source.value
    for attribute in ("font", "fill", "border", "alignment", "protection"):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.number_format = source.number_format
    if source.comment:
        target.comment = copy(source.comment)


def _copy_presentation(source: Worksheet, target: Worksheet) -> None:
    for row in source.iter_rows(min_row=1, max_row=58, max_col=13):
        for cell in row:
            if not isinstance(cell, MergedCell):
                _copy_cell(cell, target[cell.coordinate])
    for name, dimension in source.column_dimensions.items():
        if dimension.min is None or dimension.min <= 13:
            target.column_dimensions[name] = copy(dimension)
            target.column_dimensions[name].parent = target
    for row, dimension in source.row_dimensions.items():
        if row <= 58:
            target.row_dimensions[row] = copy(dimension)
            target.row_dimensions[row].parent = target
    for merged in source.merged_cells:
        if merged.max_col <= 13 and merged.max_row <= 58:
            target.merge_cells(str(merged))
    for attribute in (
        "sheet_format",
        "sheet_properties",
        "views",
        "page_setup",
        "page_margins",
        "print_options",
        "HeaderFooter",
    ):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.freeze_panes = source.freeze_panes
    target.print_title_rows = source.print_title_rows
    target.print_title_cols = source.print_title_cols
    for conditional in source.conditional_formatting:
        for area in conditional.sqref.ranges:
            if area.min_col > 13 or area.min_row > 58:
                continue
            clipped = CellRange(
                min_col=area.min_col,
                max_col=min(area.max_col, 13),
                min_row=area.min_row,
                max_row=min(area.max_row, 58),
            )
            for rule in source.conditional_formatting[conditional]:
                target.conditional_formatting.add(str(clipped), copy(rule))
    target.print_area = "B2:M68"


def _group(code: str) -> str:
    if code.startswith("DIREÇÃO - "):
        return "DIREÇÃO"
    return code


def update_display_names(workbook: Workbook) -> None:
    """Use official course names without changing debit keys or group membership."""
    support = workbook["APOIO"]
    summary = workbook["RESUMO GASTOS"]
    labels = {}
    categories = {category.code: category for category in DEBIT_CATEGORIES}
    for row in range(2, support.max_row + 1):
        code = support.cell(row, 1).value
        category = categories.get(code)
        if category is None:
            continue
        support.cell(row, 2).value = category.name
        if category.segment in ("SEG 1 GRADUAÇÃO", "SEG 2 MESTRADO"):
            labels[support.cell(row, 9).value] = category.name
            support.cell(row, 9).value = category.name
    for row in (*MAIN_ROWS, *EXTRA_ROWS):
        previous = summary.cell(row, 2).value
        if previous in labels:
            summary.cell(row, 2).value = labels[previous]


def install_layout(
    workbook: Workbook, reference_path: Path, data: ReferenceData
) -> None:
    """Replace the generic summary, retaining the original visual properties."""
    source = load_workbook(reference_path, data_only=False)
    try:
        workbook.loaded_theme = source.loaded_theme
        # Merged cells inherit font 0 when openpyxl reads them again.
        workbook._fonts = IndexedList([copy(source._fonts[0]), *workbook._fonts[1:]])
        del workbook["RESUMO GASTOS"]
        summary = workbook.create_sheet("RESUMO GASTOS")
        _copy_presentation(source["RESUMO GASTOS"], summary)
    finally:
        source.close()
    support = workbook["APOIO"]
    for col, header in enumerate(INPUT_HEADERS, 1):
        support.cell(1, col, header)
    for row, category in enumerate(DEBIT_CATEGORIES, 2):
        inputs = data.inputs.get(category.code)
        values = (
            (
                inputs.daily,
                inputs.total,
                transport_formula(row),
                inputs.scheduled,
                inputs.paid,
            )
            if inputs
            else (None, None, transport_formula(row), None, None)
        )
        for col, value in enumerate(values, 4):
            support.cell(row, col).value = value
            support.cell(row, col).number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
        support.cell(row, 9, data.labels.get(_group(category.code), category.name))
        if category.code == "PPGEL +":
            for col, value in enumerate(data.rateio or (None, None, None), 10):
                support.cell(row, col).value = value
                support.cell(row, col).number_format = "0.00%"
    for col in "DEFGH":
        support.column_dimensions[col].width = 25
    support.column_dimensions["I"].width = 44
    for col in "JKL":
        support.column_dimensions[col].width = 22
    support.tables["tblApoioDebito"].ref = f"A1:L{support.max_row}"
    # Recreate table columns when extending an existing table.
    support.tables["tblApoioDebito"].tableColumns = []
    for row, text in (
        (60, "OUTRAS CATEGORIAS E PENDÊNCIAS"),
        (62, data.labels.get("PPGDH", "PPGDH")),
        (63, "Afastamento no país"),
        (65, "PCDPs sem código de débito"),
        (66, "SEM CLASSIFICAÇÃO"),
        (68, "TOTAL CONSOLIDADO"),
    ):
        for col in range(2, 14):
            _copy_cell(summary.cell(7, col), summary.cell(row, col))
            summary.cell(row, col).value = None
        summary.cell(row, 2, text)
    for col in range(2, 14):
        _copy_cell(summary.cell(6, col), summary.cell(61, col))
    ppghd_row = next(
        row
        for row in range(2, support.max_row + 1)
        if support.cell(row, 1).value == "PPGDH"
    )
    support.cell(ppghd_row, 9).value = summary["B62"].value
    for col in (2, 3, 5, 6, 7, 9, 10, 11, 13):
        _copy_cell(summary.cell(45, col), summary.cell(68, col))
    summary["B68"] = "TOTAL CONSOLIDADO"
    for coordinate, formula in summary_formulas(workbook).items():
        summary[coordinate] = formula
    summary["C65"].number_format = "0"
    workbook.defined_names.add(DefinedName(LAYOUT_NAME, attr_text='"3"'))
    update_display_names(workbook)
    prepare_ranges(workbook)
    workbook.active = 2


def _sum_for_codes(
    rows: list[int], column: str = "K", *, canceled: bool = False
) -> str:
    terms = []
    for row in rows:
        term = f"SUMIFS('BASE VIAGENS'!${column}:${column},'BASE VIAGENS'!$M:$M,'APOIO'!$A${row},'BASE VIAGENS'!$L:$L,'APOIO'!$C${row}"
        if canceled:
            term += ",'BASE VIAGENS'!$C:$C,\"*Cancel*\""
        terms.append(term + ")")
    return "=" + "+".join(terms) if terms else "=0"


def legacy_summary_formulas(workbook: Workbook) -> dict[str, str]:
    support = workbook["APOIO"]
    summary = workbook["RESUMO GASTOS"]
    by_code = {support.cell(r, 1).value: r for r in range(2, support.max_row + 1)}
    formulas: dict[str, str] = {}
    for row in (*MAIN_ROWS, *EXTRA_ROWS):
        label = summary.cell(row, 2).value
        rows = [
            r
            for r in range(2, support.max_row + 1)
            if support.cell(r, 9).value == label
        ]
        if not rows:
            raise ValueError("Grupo do resumo sem categoria correspondente em APOIO.")
        for col, input_col in (("C", "E"), ("E", "D"), ("I", "F"), ("J", "G")):
            references = ",".join(f"'APOIO'!${input_col}${r}" for r in rows)
            if len(rows) == 1:
                formulas[f"{col}{row}"] = (
                    f'=IF(ISNUMBER({references}),{references},"Pendente")'
                )
            else:
                formulas[f"{col}{row}"] = (
                    f'=IF(COUNT({references})=0,"Pendente",SUM({references}))'
                )
            if row in (27, 28, 32):
                rate_column = {27: "J", 28: "K", 32: "L"}[row]
                split_row = by_code["PPGEL +"]
                extra = f"'APOIO'!${input_col}${split_row}*'APOIO'!${rate_column}${split_row}"
                formulas[f"{col}{row}"] = (
                    f"=IF(COUNT({references})=0,\"Pendente\",SUM({references})+IF(ISNUMBER('APOIO'!${input_col}${split_row}),{extra},0))"
                )
        formulas[f"F{row}"] = _sum_for_codes(rows)
        if row in (27, 28, 32):
            rate_column = {27: "J", 28: "K", 32: "L"}[row]
            split_row = by_code["PPGEL +"]
            formulas[f"F{row}"] += (
                "+"
                + _sum_for_codes([split_row])[1:]
                + f"*'APOIO'!${rate_column}${split_row}"
            )
        formulas[f"G{row}"] = f'=IF(ISNUMBER(E{row}),E{row}-F{row},"Pendente")'
        formulas[f"K{row}"] = (
            f'=IF(AND(ISNUMBER(I{row}),ISNUMBER(J{row})),I{row}-J{row},"Pendente")'
        )
        formulas[f"M{row}"] = (
            f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(K{row})),G{row}+K{row},"Pendente")'
        )
    for row, start, end in ((22, 7, 21), (36, 26, 35)):
        for col in ("C", "E", "F", "G", "I", "J", "K", "M"):
            formulas[f"{col}{row}"] = (
                f'=IF(COUNT({col}{start}:{col}{end})={end - start + 1},SUM({col}{start}:{col}{end}),"Pendente")'
            )
    for col in ("C", "E", "F", "G", "I", "J", "K", "M"):
        cells = f"{col}22,{col}36,{col}41,{col}43"
        formulas[f"{col}45"] = f'=IF(COUNT({cells})=4,SUM({cells}),"Pendente")'
    for row, code in enumerate(DETAIL_CODES, 50):
        formulas[f"C{row}"] = _sum_for_codes([by_code[code]])
    formulas["C58"] = "=SUM(C50:C57)"
    formulas["C65"] = "=COUNTIFS('BASE VIAGENS'!$M:$M,\"\",'BASE VIAGENS'!$A:$A,\"<>\")"
    formulas["F66"] = (
        "=SUMIFS('BASE VIAGENS'!$K:$K,'BASE VIAGENS'!$M:$M,\"\",'BASE VIAGENS'!$A:$A,\"<>\")"
    )
    for col in ("C", "E", "I", "J", "G", "K", "M"):
        formulas[f"{col}66"] = '=IF(C65=0,0,"Pendente")'
    formulas["F68"] = "=SUM(F45,F62:F63,F66)"
    for col in ("C", "E", "G", "I", "J", "K", "M"):
        formulas[f"{col}68"] = (
            f'=IF(COUNT({col}45,{col}62:{col}63,{col}66)=4,SUM({col}45,{col}62:{col}63,{col}66),"Pendente")'
        )
    return formulas


def legacy_validate_layout(workbook: Workbook) -> None:
    """Reject damaged formulas and invalid manual inputs before publication."""
    from scdp_automation.xlsx_output import WorkbookValidationError

    support = workbook["APOIO"]
    if (
        tuple(c.value for c in support[1]) != INPUT_HEADERS
        or support.max_row != len(DEBIT_CATEGORIES) + 1
    ):
        raise WorkbookValidationError(
            "As colunas ou categorias de APOIO foram alteradas."
        )
    for row, category in enumerate(DEBIT_CATEGORIES, 2):
        if tuple(support.cell(row, c).value for c in (1, 2, 3)) != (
            category.code,
            category.name,
            category.segment,
        ):
            raise WorkbookValidationError(
                "O catálogo de débitos em APOIO foi alterado."
            )
        for col in (4, 5, 7, 8):
            value = support.cell(row, col).value
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise WorkbookValidationError(
                    "Orçamento e transporte em APOIO precisam de valor numérico."
                )
        if support.cell(row, 6).value != transport_formula(row):
            raise WorkbookValidationError(
                "A fórmula de distribuição de transportes foi alterada."
            )
        weights = [support.cell(row, col).value for col in (10, 11, 12)]
        if category.code != "PPGEL +":
            if any(weight is not None for weight in weights):
                raise WorkbookValidationError(
                    "Rateio informado em categoria não habilitada."
                )
        elif any(weight is not None for weight in weights):
            if any(
                isinstance(weight, bool)
                or not isinstance(weight, (int, float))
                or not math.isfinite(weight)
                or not 0 <= weight <= 1
                for weight in weights
            ) or not math.isclose(sum(weights), 1, abs_tol=1e-10):
                raise WorkbookValidationError(
                    "O rateio PPGEL + precisa totalizar 100%."
                )
        elif any(support.cell(row, col).value for col in (4, 5, 7, 8)) or any(
            workbook["BASE VIAGENS"].cell(r, 13).value == "PPGEL +"
            and workbook["BASE VIAGENS"].cell(r, 11).value
            for r in range(2, workbook["BASE VIAGENS"].max_row + 1)
        ):
            raise WorkbookValidationError(
                "Informe o rateio PPGEL + antes de publicar despesas ou orçamento."
            )
    try:
        expected = legacy_summary_formulas(workbook)
    except ValueError as error:
        raise WorkbookValidationError(str(error)) from error
    for coordinate, formula in expected.items():
        if workbook["RESUMO GASTOS"][coordinate].value != formula:
            raise WorkbookValidationError(
                "Uma fórmula de RESUMO GASTOS foi alterada ou está ausente."
            )
    summary = workbook["RESUMO GASTOS"]
    for row in summary:
        for cell in row:
            if cell.data_type == "f" and cell.coordinate not in expected:
                raise WorkbookValidationError(
                    "Uma fórmula desconhecida foi inserida no resumo."
                )


def same_formula(actual: object, expected: str) -> bool:
    """Accept the optional APOIO sheet quotes removed by office applications."""
    if not isinstance(actual, str) or not actual.startswith("="):
        return False

    def tokens(formula: str) -> list[tuple[str, str, str]]:
        return [
            (
                token.type,
                token.subtype,
                token.value.replace("'APOIO'!", "APOIO!")
                if token.type == "OPERAND" and token.subtype == "RANGE"
                else token.value,
            )
            for token in Tokenizer(formula).items
        ]

    return tokens(actual) == tokens(expected)


def support_range(column: str) -> str:
    return f"OFFSET('APOIO'!${column}$1,1,0,MAX(1,'RESUMO GASTOS'!$R$1),1)"


def support_limit(workbook: Workbook) -> int:
    return workbook["APOIO"].max_row


def base_range(column: str) -> str:
    return f"OFFSET('BASE VIAGENS'!${column}$1,1,0,MAX(1,'RESUMO GASTOS'!$R$2),1)"


def support_expense_formula(row: int) -> str:
    return f'=IF(A{row}="","",SUMIFS({base_range("K")},{base_range("M")},A{row},{base_range("L")},C{row}))'


def range_segment_formula(row: int, limit: int | None = None) -> str:
    lookup = "OFFSET('APOIO'!$A$1,1,0,MAX(1,'RESUMO GASTOS'!$R$1),3)"
    return f"""=IF($M{row}="","",IFERROR(VLOOKUP($M{row},{lookup},3,FALSE),""))"""


def prepare_ranges(workbook: Workbook) -> None:
    """Use OFFSET ranges ending at the last populated debit or PCDP."""
    from scdp_automation.xlsx_output import CODE_LIST_NAME

    support = workbook["APOIO"]
    support.cell(1, 13, "Total utilizado por categoria (R$)")
    occupied = [
        row
        for row in range(2, support.max_row + 1)
        if any(
            support.cell(row, col).value is not None
            for col in (1, 2, 3, 4, 5, 7, 8, 9, 10, 11, 12)
        )
    ]
    last_row = max([2, *occupied]) + 1
    if support.max_row > last_row:
        support.delete_rows(last_row + 1, support.max_row - last_row)
    for row in range(2, last_row + 1):
        support.cell(row, 6).value = transport_formula(row)
        support.cell(row, 13).value = support_expense_formula(row)
        for col in (4, 5, 6, 7, 8, 13):
            support.cell(row, col).number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
    support.column_dimensions["M"].width = 30
    table = support.tables["tblApoioDebito"]
    table.ref = f"A1:M{last_row}"
    table.tableColumns = []
    workbook.defined_names.add(
        DefinedName(CODE_LIST_NAME, attr_text=support_range("A"))
    )
    summary = workbook["RESUMO GASTOS"]
    for col in range(2, 14):
        _copy_cell(summary.cell(62, col), summary.cell(64, col))
        summary.cell(64, col).value = None
    summary["B64"] = "Categorias sem grupo reconhecido"
    for row in range(1, 69):
        summary.cell(row, 15).value = (
            "Subtotal principal" if row in (22, 36, 41, 43) else None
        )
    summary.column_dimensions["O"].hidden = True
    summary.column_dimensions["P"].hidden = True
    summary.column_dimensions["R"].hidden = True
    for row in range(2, workbook["BASE VIAGENS"].max_row + 1):
        workbook["BASE VIAGENS"].cell(row, 12).value = range_segment_formula(row)
    for coordinate, formula in summary_formulas(workbook).items():
        summary[coordinate] = formula
    workbook.defined_names.add(DefinedName(LAYOUT_NAME, attr_text='"5"'))


def summary_formulas(workbook: Workbook) -> dict[str, str]:
    """Aggregate support ranges by group instead of enumerating their cells."""

    def apoio(column: str) -> str:
        return support_range(column)

    groups = apoio("I")
    codes = apoio("A")
    formulas = {}
    rate_code = '"PPGEL +"'
    formulas["R1"] = (
        """=MAX(1,IFERROR(LOOKUP(2,1/('APOIO'!$A:$A<>""),ROW('APOIO'!$A:$A))-1,1))"""
    )
    formulas["R2"] = (
        """=MAX(1,IFERROR(LOOKUP(2,1/('BASE VIAGENS'!$A:$A<>""),ROW('BASE VIAGENS'!$A:$A))-1,1))"""
    )
    for row in (*MAIN_ROWS, *EXTRA_ROWS):
        formulas[f"P{row}"] = f"=B{row}"
    for row in (*MAIN_ROWS, *EXTRA_ROWS):
        criterion = f"$B{row}"
        for output, source in (("C", "E"), ("E", "D"), ("I", "F"), ("J", "G")):
            amounts = apoio(source)
            count = f"SUMPRODUCT(--({groups}={criterion}),--ISNUMBER({amounts}))"
            summed = f"SUMIF({groups},{criterion},{amounts})"
            if row in (27, 28, 32):
                weight_col = {27: "J", 28: "K", 32: "L"}[row]
                weights = apoio(weight_col)
                summed += f"+SUMPRODUCT(--({codes}={rate_code}),{amounts},{weights})"
                count += f"+SUMPRODUCT(--({codes}={rate_code}),--ISNUMBER({amounts}),--({weights}>0))"
            formulas[f"{output}{row}"] = f'=IF({count}=0,"Pendente",{summed})'
        formulas[f"F{row}"] = f"=SUMIF({groups},{criterion},{apoio('M')})"
        if row in (27, 28, 32):
            weight_col = {27: "J", 28: "K", 32: "L"}[row]
            formulas[f"F{row}"] += (
                f"+SUMPRODUCT(--({codes}={rate_code}),{apoio('M')},{apoio(weight_col)})"
            )
        formulas[f"G{row}"] = f'=IF(ISNUMBER(E{row}),E{row}-F{row},"Pendente")'
        formulas[f"K{row}"] = (
            f'=IF(AND(ISNUMBER(I{row}),ISNUMBER(J{row})),I{row}-J{row},"Pendente")'
        )
        formulas[f"M{row}"] = (
            f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(K{row})),G{row}+K{row},"Pendente")'
        )
    for row, start, end in ((22, 7, 21), (36, 26, 35)):
        for col in ("C", "E", "F", "G", "I", "J", "K", "M"):
            formulas[f"{col}{row}"] = (
                f'=IF(COUNT({col}{start}:{col}{end})={end - start + 1},SUM({col}{start}:{col}{end}),"Pendente")'
            )
    # SUMIF over a row-kind range avoids adding individual subtotal cells.
    for col in ("C", "E", "F", "G", "I", "J", "K", "M"):
        amounts = f"{col}7:{col}43"
        formulas[f"{col}45"] = (
            f'=IF(COUNTIFS($O$7:$O$43,"Subtotal principal",{amounts},"Pendente")>0,"Pendente",SUMIF($O$7:$O$43,"Subtotal principal",{amounts}))'
        )
    for row, code in enumerate(DETAIL_CODES, 50):
        formulas[f"C{row}"] = f'=SUMIF({codes},"{code}",{apoio("M")})'
    formulas["C58"] = "=SUM(C50:C57)"
    formulas["C65"] = f'=COUNTIFS({base_range("M")},"",{base_range("A")},"<>")'
    formulas["F66"] = (
        f'=SUMIFS({base_range("K")},{base_range("M")},"",{base_range("A")},"<>")'
    )
    for col in ("C", "E", "I", "J", "G", "K", "M"):
        formulas[f"{col}66"] = '=IF(C65=0,0,"Pendente")'
    formulas["F68"] = f"=SUM({base_range('K')})"
    formulas["F64"] = "=F68-SUM(F45,F62:F63,F66)"
    # Unmapped groups remain visible, including budgets without expenses.
    unknown = f'SUMPRODUCT(--({codes}<>""),--({codes}<>{rate_code}),--ISNA(MATCH({groups},$P$7:$P$63,0)))'
    for col in ("C", "E", "I", "J", "G", "K", "M"):
        formulas[f"{col}64"] = f'=IF({unknown}=0,0,"Pendente")'
        formulas[f"{col}68"] = (
            f'=IF(COUNT({col}45,{col}62:{col}64,{col}66)=5,SUM({col}45,{col}62:{col}64,{col}66),"Pendente")'
        )
    return formulas


def validate_layout(workbook: Workbook) -> None:
    from scdp_automation.xlsx_output import WorkbookValidationError

    support = workbook["APOIO"]
    expected_headers = (*INPUT_HEADERS, "Total utilizado por categoria (R$)")
    if tuple(c.value for c in support[1]) != expected_headers:
        raise WorkbookValidationError("As colunas de APOIO foram alteradas.")
    expected = {category.code: category for category in DEBIT_CATEGORIES}
    seen = set()
    allowed_segments = {category.segment for category in DEBIT_CATEGORIES}
    for row in range(2, support.max_row + 1):
        code = support.cell(row, 1).value
        if not code:
            if any(
                support.cell(row, col).value is not None
                for col in (2, 3, 4, 5, 7, 8, 9, 10, 11, 12)
            ):
                raise WorkbookValidationError(
                    "Preencha o código da categoria em APOIO."
                )
            continue
        if (
            not isinstance(code, str)
            or code != code.strip()
            or code in seen
            or any(c in code for c in "*?~")
        ):
            raise WorkbookValidationError("Código inválido ou duplicado em APOIO.")
        seen.add(code)
        if (
            not isinstance(support.cell(row, 2).value, str)
            or support.cell(row, 3).value not in allowed_segments
        ):
            raise WorkbookValidationError("Nome ou segmento inválido em APOIO.")
        if code in expected and (
            support.cell(row, 2).value,
            support.cell(row, 3).value,
        ) != (expected[code].name, expected[code].segment):
            raise WorkbookValidationError(
                "O catálogo de débitos em APOIO foi alterado."
            )
        for col in (4, 5, 7, 8):
            value = support.cell(row, col).value
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise WorkbookValidationError(
                    "Orçamento e transporte em APOIO precisam de valor numérico."
                )
        if support.cell(row, 6).value != transport_formula(row) or support.cell(
            row, 13
        ).value != support_expense_formula(row):
            raise WorkbookValidationError("Uma fórmula de APOIO foi alterada.")
        weights = [support.cell(row, col).value for col in (10, 11, 12)]
        if code != "PPGEL +":
            if any(w is not None for w in weights):
                raise WorkbookValidationError(
                    "Rateio informado em categoria não habilitada."
                )
        else:
            required = any(support.cell(row, col).value for col in (4, 5, 7, 8)) or any(
                workbook["BASE VIAGENS"].cell(r, 13).value == code
                and workbook["BASE VIAGENS"].cell(r, 11).value
                for r in range(2, workbook["BASE VIAGENS"].max_row + 1)
            )
            if (required or any(w is not None for w in weights)) and (
                any(
                    isinstance(w, bool)
                    or not isinstance(w, (int, float))
                    or not math.isfinite(w)
                    or not 0 <= w <= 1
                    for w in weights
                )
                or not math.isclose(sum(weights), 1, abs_tol=1e-10)
            ):
                raise WorkbookValidationError(
                    "O rateio PPGEL + precisa totalizar 100%."
                )
    if not set(expected).issubset(seen):
        raise WorkbookValidationError("O catálogo de débitos em APOIO está incompleto.")
    formulas = summary_formulas(workbook)
    summary = workbook["RESUMO GASTOS"]
    if any(
        summary.cell(row, 15).value
        != ("Subtotal principal" if row in (22, 36, 41, 43) else None)
        for row in range(1, 69)
    ):
        raise WorkbookValidationError("A identificação dos subtotais foi alterada.")
    for coordinate, formula in formulas.items():
        if not same_formula(summary[coordinate].value, formula):
            raise WorkbookValidationError(
                "Uma fórmula de RESUMO GASTOS foi alterada ou está ausente."
            )
    for row in summary:
        for cell in row:
            if cell.data_type == "f" and cell.coordinate not in formulas:
                raise WorkbookValidationError(
                    "Uma fórmula desconhecida foi inserida no resumo."
                )
