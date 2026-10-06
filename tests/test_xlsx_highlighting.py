import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from scdp_automation.xlsx_output import (
    _load_workbook_for_refresh,
    import_reference_workbook,
    install_decision_highlighting,
)
from scdp_automation.xlsx_recalculate import recalculate_workbook
from tests.support.workbooks import reference_fixture


@unittest.skipUnless(shutil.which("libreoffice"), "LibreOffice necessário")
class HighlightingTests(unittest.TestCase):
    def test_red_cancellations_and_yellow_missing_segments_in_future_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            reference, output = root / "reference.xlsx", root / "output.xlsx"
            reference_fixture(reference)
            import_reference_workbook(reference, output)
            book = load_workbook(output)
            base = book["BASE VIAGENS"]
            base.conditional_formatting._cf_rules.clear()
            from openpyxl.formatting.rule import FormulaRule

            base.conditional_formatting.add(
                "A2:Q1048576", FormulaRule(formula=['AND($A2<>"",OR($P2="",$Q2=""))'])
            )
            book.save(output)
            book.close()
            book = _load_workbook_for_refresh(output)
            base = book["BASE VIAGENS"]
            install_decision_highlighting(base)
            rules = [
                (area, rule)
                for area in base.conditional_formatting
                for rule in base.conditional_formatting[area]
            ]
            self.assertEqual(len(rules), 2)
            rules.sort(key=lambda item: item[1].priority)
            red, yellow = rules[0][1], rules[1][1]
            self.assertTrue(red.stopIfTrue)
            self.assertEqual(red.dxf.fill.fgColor.rgb[-6:], "FFC7CE")
            self.assertEqual(yellow.dxf.fill.fgColor.rgb[-6:], "FFF2CC")
            self.assertTrue(all(str(area.sqref) == "A2:R1048576" for area, _ in rules))
            cases = [
                ("Cancelada", None, "SEG 1 GRADUAÇÃO", "111111/26", True, False),
                ("Concluída", None, "SEG 1 GRADUAÇÃO", "222222/26", False, False),
                ("Em Planejamento", "Sim", None, "333333/26", False, True),
                ("Em Prestação de Contas", None, None, "444444/26", False, True),
                ("Concluída", "Sim", "SEG 1 GRADUAÇÃO", "555555/26", False, False),
                ("Cancelada", "Não", "SEG 1 GRADUAÇÃO", "666666/26", False, False),
                ("Cancelada", None, None, "777777/26", True, True),
                ("Concluída", None, None, None, False, False),
            ]
            for row, (status, decision, segment, pcdp, _, _) in enumerate(cases, 100):
                base.cell(row, 1).value = pcdp
                base.cell(row, 3).value = status
                base.cell(row, 16).value = segment
                base.cell(row, 18).value = decision
                for col, rule in ((20, red), (21, yellow)):
                    base.cell(row, col).value = Translator(
                        "=" + rule.formula[0], origin="A2"
                    ).translate_formula(f"A{row}")
            book.save(output)
            book.close()
            recalculate_workbook(output)
            cached = load_workbook(output, data_only=True)
            for col, expectation in ((20, 4), (21, 5)):
                self.assertEqual(
                    [
                        cached["BASE VIAGENS"].cell(r, col).value
                        for r in range(100, 108)
                    ],
                    [case[expectation] for case in cases],
                )
            cached.close()
