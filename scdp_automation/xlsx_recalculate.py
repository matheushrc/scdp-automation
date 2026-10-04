"""Recalculate a disposable copy and transfer only its formula result caches."""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)
ET.register_namespace(
    "r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)


def recalculate_workbook(path: Path) -> None:
    """Make formula results visible while preserving the original XLSX design."""
    executable = shutil.which("libreoffice")
    if executable is None:
        raise RuntimeError("LibreOffice não encontrado para recalcular a cópia.")
    with tempfile.TemporaryDirectory(prefix="scdp-recalculate-") as directory:
        root = Path(directory)
        incoming, outgoing = root / "input", root / "calculated"
        incoming.mkdir()
        outgoing.mkdir()
        source = incoming / "workbook.xlsx"
        shutil.copy2(path, source)
        result = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={(root / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(outgoing),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
            env=os.environ
            | {"GSETTINGS_BACKEND": "memory", "SAL_USE_VCLPLUGIN": "svp"},
        )
        calculated = outgoing / source.name
        if result.returncode != 0 or not calculated.exists():
            raise RuntimeError(
                "LibreOffice não conseguiu recalcular a cópia descartável."
            )
        _transfer_caches(path, calculated)


def _transfer_caches(original: Path, calculated: Path) -> None:
    original_book = load_workbook(original, data_only=False)
    calculated_book = load_workbook(calculated, data_only=True)
    caches: dict[str, dict[str, object]] = {}
    try:
        for sheet in original_book:
            if sheet.title not in calculated_book.sheetnames:
                raise ValueError("O recálculo removeu uma aba necessária.")
            computed = calculated_book[sheet.title]
            values = {}
            for row in sheet:
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    value = computed[cell.coordinate]
                    if value.data_type == "e" or (
                        value.value is None and value.data_type != "str"
                    ):
                        raise ValueError(
                            f"Fórmula sem resultado válido em {sheet.title}!{cell.coordinate}."
                        )
                    values[cell.coordinate] = "" if value.value is None else value.value
            caches[sheet.title] = values
        sheetnames = original_book.sheetnames
    finally:
        original_book.close()
        calculated_book.close()
    temporary: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=original.parent, suffix=".xlsx", delete=False
        ) as stream:
            temporary = stream.name
        with (
            ZipFile(original) as source,
            ZipFile(temporary, "w", compression=ZIP_DEFLATED) as target,
        ):
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename.startswith(
                    "xl/worksheets/sheet"
                ) and entry.filename.endswith(".xml"):
                    index = (
                        int(
                            entry.filename.removeprefix(
                                "xl/worksheets/sheet"
                            ).removesuffix(".xml")
                        )
                        - 1
                    )
                    values = caches[sheetnames[index]]
                    document = ET.fromstring(content)
                    for cell in document.iter(f"{{{NS}}}c"):
                        coordinate = cell.attrib.get("r")
                        if coordinate not in values:
                            continue
                        old = cell.find(f"{{{NS}}}v")
                        if old is not None:
                            cell.remove(old)
                        value = values[coordinate]
                        cached = ET.SubElement(cell, f"{{{NS}}}v")
                        if isinstance(value, str):
                            cell.set("t", "str")
                            cached.text = value
                        elif isinstance(value, bool):
                            cell.set("t", "b")
                            cached.text = "1" if value else "0"
                        elif isinstance(value, (int, float)):
                            cell.attrib.pop("t", None)
                            cached.text = repr(value)
                        else:
                            raise TypeError(
                                "Tipo de resultado de fórmula não suportado."
                            )
                    content = ET.tostring(
                        document, encoding="utf-8", xml_declaration=True
                    )
                target.writestr(entry, content)
        os.replace(temporary, original)
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)
