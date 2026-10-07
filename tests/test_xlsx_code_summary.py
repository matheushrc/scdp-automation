import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from scdp_automation.xlsx_output import build_candidate
from tests.support.recalculation import recalculate_workbook
from tests.support.workbooks import (
    final_template_fixture,
    make_trip,
    write_existing_workbook,
)


@unittest.skipUnless(shutil.which("libreoffice"), "LibreOffice necessário")
class CodeSummaryTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template = self.root / "template.xlsx"
        self.current = self.root / "current.xlsx"
        self.candidate = self.root / "candidate.xlsx"
        final_template_fixture(self.template)

    def test_new_category_uses_code_without_group_or_hidden_metadata(self):
        trip = make_trip("111111/26")
        trip.total_da_viagem_r = 125
        write_existing_workbook(self.current, [trip])
        book = load_workbook(self.current)
        support, summary = book["APOIO"], book["RESUMO GASTOS"]
        self.assertNotIn("Grupo no resumo", [cell.value for cell in support[1]])
        support.append(
            ["TESTE", "Categoria teste", "SEG 1 GRADUAÇÃO", 10000, 0, None, 0]
        )
        summary["B25"] = "TESTE"
        for col in ("C", "E", "F", "G", "I", "J", "K", "M"):
            summary[f"{col}25"] = Translator(
                summary[f"{col}7"].value, origin=f"{col}7"
            ).translate_formula(f"{col}25")
        book["BASE VIAGENS"]["Q2"] = "TESTE"
        book.save(self.current)
        book.close()
        build_candidate(
            [trip], self.current, self.candidate, template_path=self.template
        )
        recalculate_workbook(self.candidate)
        cached = load_workbook(self.candidate, data_only=True)
        self.addCleanup(cached.close)
        self.assertEqual(cached["RESUMO GASTOS"]["E25"].value, 10000)
        self.assertEqual(cached["RESUMO GASTOS"]["F25"].value, 125)
        self.assertEqual(cached["RESUMO GASTOS"]["M25"].value, 9875)
        self.assertEqual(cached["RESUMO GASTOS"]["C20"].value, 19000)
        self.assertEqual(cached["RESUMO GASTOS"]["F20"].value, 125)

    def test_direction_subcodes_and_rateio_keep_their_totals(self):
        trips = [make_trip(pcdp) for pcdp in ("111111/26", "222222/26", "333333/26")]
        for trip, amount in zip(trips, (75, 25, 156), strict=True):
            trip.total_da_viagem_r = amount
        write_existing_workbook(
            self.current,
            trips,
            {
                "111111/26": "DIREÇÃO - AGAS",
                "222222/26": "DIREÇÃO",
                "333333/26": "PPGEL +",
            },
        )
        book = load_workbook(self.current)
        support = book["APOIO"]
        row = next(
            r
            for r in range(2, support.max_row + 1)
            if support.cell(r, 1).value == "DIREÇÃO - AGAS"
        )
        for col, value in ((4, 10), (5, 20), (7, 5)):
            support.cell(row, col).value = value
        book.save(self.current)
        book.close()
        build_candidate(
            trips, self.current, self.candidate, template_path=self.template
        )
        recalculate_workbook(self.candidate)
        cached = load_workbook(self.candidate, data_only=True)
        self.addCleanup(cached.close)
        summary = cached["RESUMO GASTOS"]
        self.assertEqual(summary["C12"].value, 1030)
        self.assertEqual(summary["F12"].value, 100)
        self.assertEqual(summary["J12"].value, 205)
        self.assertEqual([summary[f"F{row}"].value for row in (8, 9, 11)], [78, 39, 39])
        self.assertEqual(summary["C20"].value, 8030)
        self.assertEqual(summary["F20"].value, 256)

    def test_manual_formulas_and_support_references_survive_refresh(self):
        trip = make_trip()
        write_existing_workbook(
            self.current, [trip], {trip.numero_da_solicitacao: "AGRONOMIA"}
        )
        book = load_workbook(self.current)
        book["RESUMO GASTOS"]["E7"] = "=SUMIF(ApoioA,$B7,ApoioD)+10"
        book["RESUMO GASTOS"]["M25"] = "=COUNT(APOIO!L2:L3)"
        book.save(self.current)
        book.close()
        original = self.current.read_bytes()
        build_candidate(
            [trip], self.current, self.candidate, template_path=self.template
        )
        self.assertEqual(self.current.read_bytes(), original)
        book = load_workbook(self.candidate)
        self.addCleanup(book.close)
        self.assertEqual(
            book["RESUMO GASTOS"]["E7"].value, "=SUMIF(ApoioA,$B7,ApoioD)+10"
        )
        self.assertEqual(book["RESUMO GASTOS"]["M25"].value, "=COUNT(APOIO!L2:L3)")
        recalculate_workbook(self.candidate)
        cached = load_workbook(self.candidate, data_only=True)
        self.addCleanup(cached.close)
        self.assertEqual(cached["RESUMO GASTOS"]["E7"].value, 410)
