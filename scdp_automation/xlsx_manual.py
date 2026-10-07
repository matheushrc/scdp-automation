"""Read historical BASE manual inputs by explicit semantic field identity."""

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Literal

from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from scdp_automation.relatorio import PCDP_PATTERN, Viagem
from scdp_automation.xlsx_validation import WorkbookValidationError


@dataclass(frozen=True)
class ManualValues:
    codigo_de_debito: str | None
    descontar_do_curso: Literal["Sim", "Não"] | None


_HEADER_ALIASES = {
    "pcdp": ("PCDP", "Número da Solicitação"),
    "codigo_de_debito": ("Código de débito", "Codigo de debito"),
    "descontar_do_curso": ("Descontar do curso?", "Descontar do curso"),
}


def _pcdp(value: object, context: str) -> str:
    if not isinstance(value, str) or not re.fullmatch(PCDP_PATTERN, value.strip()):
        raise WorkbookValidationError(f"{context}: PCDP inválida: {value!r}.")
    return value.strip()


def _manual_values(code: object, decision: object, context: str) -> ManualValues:
    def text(value: object, field: str) -> str | None:
        if value is None:
            return None
        if not isinstance(value, str) or value.strip().startswith("="):
            raise WorkbookValidationError(f"{context}: {field} inválido: {value!r}.")
        return value.strip() or None

    normalized_code = text(code, "Código de débito")
    normalized_decision = text(decision, "Descontar do curso?")
    if normalized_decision not in (None, "Sim", "Não"):
        raise WorkbookValidationError(
            f"{context}: Descontar do curso? inválido: {decision!r}."
        )
    return ManualValues(normalized_code, normalized_decision)


def read_base_manual_values(base: Worksheet) -> dict[str, ManualValues]:
    """Import only recognized manual columns; callers load with data_only=False."""
    columns: dict[str, int] = {}
    for cell in base[1]:
        header = cell.value.strip() if isinstance(cell.value, str) else None
        for field, aliases in _HEADER_ALIASES.items():
            if header in aliases:
                if field in columns:
                    raise WorkbookValidationError(
                        f"{base.title}: cabeçalho ambíguo para {field}."
                    )
                columns[field] = cell.column
    for field in _HEADER_ALIASES:
        if field not in columns:
            raise WorkbookValidationError(
                f"{base.title}: cabeçalho ausente para {field}."
            )

    values: dict[str, ManualValues] = {}
    for row in range(2, base.max_row + 1):
        context = f"{base.title}, linha {row}"
        pcdp_cell = base.cell(row, columns["pcdp"]).value
        manual = _manual_values(
            base.cell(row, columns["codigo_de_debito"]).value,
            base.cell(row, columns["descontar_do_curso"]).value,
            context,
        )
        if pcdp_cell is None or (isinstance(pcdp_cell, str) and not pcdp_cell.strip()):
            if manual != ManualValues(None, None):
                raise WorkbookValidationError(f"{context}: valor manual sem PCDP.")
            continue
        pcdp = _pcdp(pcdp_cell, context)
        if pcdp in values:
            raise WorkbookValidationError(f"{context}: PCDP duplicada: {pcdp}.")
        values[pcdp] = manual
    return values


def apply_base_manual_values(
    trips: Sequence[Viagem], values: Mapping[str, ManualValues]
) -> list[Viagem]:
    """Return validated copies, retaining JSON-only trips and applying explicit clears."""
    normalized: dict[str, ManualValues] = {}
    for pcdp, manual in values.items():
        key = _pcdp(pcdp, "Entradas manuais")
        if key in normalized:
            raise WorkbookValidationError(f"Entradas manuais: PCDP duplicada: {key}.")
        normalized[key] = _manual_values(
            manual.codigo_de_debito, manual.descontar_do_curso, f"PCDP {key}"
        )
    seen: set[str] = set()
    result: list[Viagem] = []
    for trip in trips:
        pcdp = _pcdp(trip.numero_da_solicitacao, "JSON")
        if pcdp in seen:
            raise WorkbookValidationError(f"JSON: PCDP duplicada: {pcdp}.")
        seen.add(pcdp)
        data = trip.model_dump()
        manual = normalized.get(pcdp)
        if manual is None:
            _manual_values(
                trip.codigo_de_debito, trip.descontar_do_curso, f"JSON, PCDP {pcdp}"
            )
        if manual is not None:
            data.update(
                codigo_de_debito=manual.codigo_de_debito,
                descontar_do_curso=manual.descontar_do_curso,
            )
        try:
            result.append(Viagem.model_validate(data))
        except ValidationError as exc:
            raise WorkbookValidationError(f"JSON, PCDP {pcdp}: {exc}") from exc
    missing = tuple(sorted(normalized.keys() - seen))
    if missing:
        raise WorkbookValidationError(
            "Entradas manuais para PCDPs ausentes do JSON: " + ", ".join(missing),
            missing_pcdps=missing,
        )
    return result
