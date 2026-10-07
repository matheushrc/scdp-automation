"""Synchronize manual BASE fields and recover workbooks without a browser."""

from pathlib import Path

from openpyxl import load_workbook

from scdp_automation import xlsx_output
from scdp_automation.output_history import OutputHistory
from scdp_automation.relatorio import Viagem, load_trips, save_json
from scdp_automation.xlsx_manual import (
    apply_base_manual_values,
    read_base_manual_values,
)
from scdp_automation.xlsx_output import publish_workbook, validate_workbook_sources
from scdp_automation.xlsx_validation import WorkbookValidationError


def prepare_checkpoint_trips(
    checkpoint_path: Path,
    workbook_path: Path,
    *,
    template_path: Path | None = None,
    require_checkpoint: bool = False,
) -> list[Viagem]:
    """Validate sources and import manual inputs without writes or a session lock."""
    validate_workbook_sources(workbook_path, template_path=template_path)
    if not checkpoint_path.is_file() and (require_checkpoint or workbook_path.exists()):
        raise WorkbookValidationError(
            f"O checkpoint JSON em {checkpoint_path} está ausente; reconcilie os "
            f"arquivos antes de atualizar {workbook_path}. A planilha foi preservada."
        )
    trips = load_trips(checkpoint_path)
    existing = workbook_path.exists()
    source = (
        workbook_path if existing else template_path or xlsx_output.DEFAULT_TEMPLATE
    )
    workbook = load_workbook(source, data_only=False)
    try:
        values = read_base_manual_values(workbook["BASE VIAGENS"]) if existing else {}
        prepared = apply_base_manual_values(trips, values)
        catalog = {
            workbook["APOIO"].cell(row, 1).value
            for row in range(2, workbook["APOIO"].max_row + 1)
        }
        for trip in prepared:
            if (
                trip.codigo_de_debito is not None
                and trip.codigo_de_debito not in catalog
            ):
                raise WorkbookValidationError(
                    f"JSON, PCDP {trip.numero_da_solicitacao}: código de débito "
                    f"desconhecido em APOIO: {trip.codigo_de_debito}."
                )
        return prepared
    finally:
        workbook.close()


def recreate_workbook(
    checkpoint_path: Path,
    workbook_path: Path,
    *,
    template_path: Path | None = None,
) -> None:
    """Recover from the existing JSON and manual workbook under the output lock."""
    with OutputHistory(checkpoint_path, workbook_path) as history:
        trips = prepare_checkpoint_trips(
            checkpoint_path,
            workbook_path,
            template_path=template_path,
            require_checkpoint=True,
        )
        history.archive_previous()
        save_json(checkpoint_path, trips)
        publish_workbook(trips, workbook_path, template_path=template_path)
        history.prune()
