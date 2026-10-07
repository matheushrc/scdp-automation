import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from scdp_automation import xlsx_output
from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import final_template_fixture, make_trip


def assert_manual_sheets_preserved(case, actual_workbook, source):
    case.assertEqual(dict(actual_workbook.defined_names), dict(source.defined_names))
    case.assertEqual(actual_workbook.loaded_theme, source.loaded_theme)
    for name in ("APOIO", "RESUMO GASTOS"):
        actual, expected = actual_workbook[name], source[name]
        case.assertEqual(actual.max_row, expected.max_row)
        case.assertEqual(actual.max_column, expected.max_column)
        case.assertEqual(
            set(map(str, actual.merged_cells)), set(map(str, expected.merged_cells))
        )
        case.assertEqual(actual.data_validations, expected.data_validations)
        case.assertEqual(dict(actual.tables), dict(expected.tables))
        case.assertEqual(
            {
                str(area.sqref): actual.conditional_formatting[area]
                for area in actual.conditional_formatting
            },
            {
                str(area.sqref): expected.conditional_formatting[area]
                for area in expected.conditional_formatting
            },
        )
        for property_name in (
            "print_area",
            "print_options",
            "page_setup",
            "page_margins",
            "sheet_properties",
            "sheet_format",
            "sheet_view",
            "protection",
        ):
            case.assertEqual(
                getattr(actual, property_name), getattr(expected, property_name)
            )
        case.assertEqual(
            {
                key: (dict(value), value._style)
                for key, value in actual.column_dimensions.items()
            },
            {
                key: (dict(value), value._style)
                for key, value in expected.column_dimensions.items()
            },
        )
        case.assertEqual(
            {
                key: (dict(value), value._style)
                for key, value in actual.row_dimensions.items()
            },
            {
                key: (dict(value), value._style)
                for key, value in expected.row_dimensions.items()
            },
        )
        case.assertEqual(dict(actual.defined_names), dict(expected.defined_names))
        for row in expected:
            for cell in row:
                copied = actual[cell.coordinate]
                if not (name == "RESUMO GASTOS" and cell.coordinate == "M2"):
                    case.assertEqual(
                        copied.value, cell.value, f"{name}!{cell.coordinate}"
                    )
                case.assertEqual(
                    copied._style, cell._style, f"{name}!{cell.coordinate}"
                )


class FidelityTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template = self.root / "template.xlsx"
        self.current = self.root / "current.xlsx"
        self.candidate = self.root / "candidate.xlsx"
        final_template_fixture(self.template)
        final_template_fixture(self.current)
        default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.template)
        default.start()
        self.addCleanup(default.stop)

    def test_missing_template_preserves_published_files(self):
        original = self.current.read_bytes()
        checkpoint = self.root / "checkpoint.json"
        checkpoint.write_text('{"private": "synthetic"}')
        checkpoint_bytes = checkpoint.read_bytes()
        with (
            patch.object(xlsx_output, "DEFAULT_TEMPLATE", self.root / "missing.xlsx"),
            self.assertRaises(WorkbookValidationError),
        ):
            xlsx_output.build_candidate([], self.current, self.candidate)
        self.assertEqual(self.current.read_bytes(), original)
        self.assertEqual(checkpoint.read_bytes(), checkpoint_bytes)
        self.assertFalse(self.candidate.exists())

    def test_old_template_is_rejected(self):
        book = load_workbook(self.template)
        book.defined_names["SCDPLayoutVersion"].attr_text = '"8"'
        book.save(self.template)
        book.close()
        original = self.current.read_bytes()
        with self.assertRaises(WorkbookValidationError):
            xlsx_output.build_candidate([], self.current, self.candidate)
        self.assertEqual(self.current.read_bytes(), original)
        self.assertFalse(self.candidate.exists())

    def test_old_output_is_rejected(self):
        book = load_workbook(self.current)
        book.defined_names["SCDPLayoutVersion"].attr_text = '"8"'
        book.save(self.current)
        book.close()
        original = self.current.read_bytes()
        with self.assertRaises(WorkbookValidationError):
            xlsx_output.build_candidate([], self.current, self.candidate)
        self.assertEqual(self.current.read_bytes(), original)
        self.assertFalse(self.candidate.exists())

    def test_fresh_base_does_not_import_template_trips(self):
        book = load_workbook(self.template)
        base = book["BASE VIAGENS"]
        for coordinate, value in (
            ("A2", "999001/26"),
            ("Q2", "AGRONOMIA"),
            ("R2", "Sim"),
            ("O2", "Antiga"),
            ("N2", date(2026, 1, 1)),
        ):
            base[coordinate] = value
        book.save(self.template)
        book.close()
        self.current.unlink()
        original = self.template.read_bytes()
        xlsx_output.build_candidate(
            [make_trip("999001/26")], self.current, self.candidate
        )
        actual = load_workbook(self.candidate)
        source = load_workbook(self.template)
        self.addCleanup(actual.close)
        self.addCleanup(source.close)
        self.assertEqual(self.template.read_bytes(), original)
        self.assertEqual(actual["BASE VIAGENS"]["A2"].value, "999001/26")
        for column in ("N", "O", "Q", "R"):
            self.assertIsNone(actual["BASE VIAGENS"][f"{column}2"].value)
        assert_manual_sheets_preserved(self, actual, source)

    def test_refresh_preserves_manual_sheets_and_full_pcdp_choices(self):
        book = load_workbook(self.current)
        base = book["BASE VIAGENS"]
        for row, (pcdp, code, decision) in enumerate(
            (("999001/26", "AGRONOMIA", "Não"), ("999001/26-1C", "PPGE", "Sim")), 2
        ):
            base.cell(row, 1, pcdp)
            base.cell(row, 17, code)
            base.cell(row, 18, decision)
        base.tables["tblBaseViagens"].ref = "A1:R3"
        book["APOIO"]["D2"] = 12345
        book["RESUMO GASTOS"]["F7"] = "=SUM(1,2)"
        book.save(self.current)
        book.close()
        original = self.current.read_bytes()
        trips = [
            make_trip(pcdp) for pcdp in ("999001/26-1C", "999001/26", "999001/26-2C")
        ]
        xlsx_output.build_candidate(trips, self.current, self.candidate)
        source, actual = load_workbook(self.current), load_workbook(self.candidate)
        self.addCleanup(source.close)
        self.addCleanup(actual.close)
        self.assertEqual(self.current.read_bytes(), original)
        assert_manual_sheets_preserved(self, actual, source)
        self.assertEqual(
            [actual["BASE VIAGENS"].cell(r, 17).value for r in (2, 3, 4)],
            ["PPGE", "AGRONOMIA", None],
        )
        self.assertEqual(
            [actual["BASE VIAGENS"].cell(r, 18).value for r in (2, 3, 4)],
            ["Sim", "Não", None],
        )

    def test_empty_description_does_not_restore_old_value(self):
        book = load_workbook(self.current)
        book["BASE VIAGENS"]["A2"] = "999001/26"
        book["BASE VIAGENS"]["O2"] = "Descrição antiga"
        book.save(self.current)
        book.close()
        trip = make_trip("999001/26")
        trip.descricao_do_motivo_da_viagem = ""
        xlsx_output.build_candidate([trip], self.current, self.candidate)
        updated = load_workbook(self.candidate)
        self.addCleanup(updated.close)
        self.assertIn(updated["BASE VIAGENS"]["O2"].value, (None, ""))
        self.assertEqual(
            trip.model_dump(mode="json")["descricao_do_motivo_da_viagem"], ""
        )

    def test_publication_records_the_value_update_date_in_m2(self):
        updated_at = datetime(2026, 10, 7, 1, 30, tzinfo=ZoneInfo("America/Sao_Paulo"))
        with patch("scdp_automation.xlsx_output.datetime") as clock:
            clock.now.return_value = updated_at
            xlsx_output.build_candidate([], self.current, self.candidate)
        workbook = load_workbook(self.candidate)
        self.addCleanup(workbook.close)
        cell = workbook["RESUMO GASTOS"]["M2"]
        self.assertEqual(cell.value.date(), date(2026, 10, 7))
        self.assertEqual(cell.number_format, "dd/mm/yyyy")

    def test_new_expense_requires_rateio(self):
        book = load_workbook(self.current)
        base, support = book["BASE VIAGENS"], book["APOIO"]
        base["A2"], base["Q2"], base["K2"] = "999001/26", "PPGEL +", 0
        row = next(
            r
            for r in range(2, support.max_row + 1)
            if support.cell(r, 1).value == "PPGEL +"
        )
        for column in (4, 5, 7, 8, 9, 10, 11):
            support.cell(row, column).value = None
        book.save(self.current)
        book.close()
        original = self.current.read_bytes()
        trip = make_trip("999001/26")
        trip.total_da_viagem_r = 100
        with self.assertRaisesRegex(WorkbookValidationError, "100%"):
            xlsx_output.build_candidate([trip], self.current, self.candidate)
        self.assertEqual(self.current.read_bytes(), original)
        self.assertFalse(self.candidate.exists())

    def test_unverified_trip_does_not_restore_workbook_verification_date(self):
        book = load_workbook(self.current)
        book["BASE VIAGENS"]["A2"] = "999001/26"
        for column in ("L", "M", "N"):
            book["BASE VIAGENS"][f"{column}2"] = date(2020, 1, 1)
        book.save(self.current)
        book.close()
        trip = make_trip("999001/26")
        xlsx_output.build_candidate([trip], self.current, self.candidate)
        updated = load_workbook(self.candidate)
        self.addCleanup(updated.close)
        self.assertIsNone(updated["BASE VIAGENS"]["N2"].value)
