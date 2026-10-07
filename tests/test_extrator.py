import asyncio
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import AsyncMock, MagicMock, Mock, patch

from scdp_automation import xlsx_output
from scdp_automation.config import REPO_ROOT, current_year
from scdp_automation.extrator import (
    collect_pending_descriptions,
    preflight_extraction,
    resolve_report_url,
    run,
    save_checkpoint_and_publish,
)
from scdp_automation.output_history import OutputHistory
from scdp_automation.relatorio import Viagem, load_trips, save_json

OUTPUT_PATH = REPO_ROOT / "output" / f"viagens_scdp_{current_year()}.json"
WORKBOOK_PATH = REPO_ROOT / "output" / f"gastos_scdp_{current_year()}.xlsx"


def make_valid_trip(
    pcdp: str = "999999/26",
    name: str = "Pessoa de Teste",
    description: str | None = None,
) -> Viagem:
    return Viagem.model_validate(
        {
            "numero_da_solicitacao": pcdp,
            "nome_do_proposto": name,
            "orgao_solicitante": "ORG-TESTE",
            "orgao_superior": "ORG-SUPERIOR-TESTE",
            "tipo_da_viagem": "NACIONAL",
            "situacao_da_viagem": "Autorizada",
            "motivo_viagem": "Nacional - A Serviço",
            "trechos": [
                {
                    "inicio": "01/03/2026",
                    "termino": "02/03/2026",
                    "origem": "Cidade Alfa (AA)",
                    "destino": "Cidade Beta (BB)",
                    "meio_de_transporte": "Aéreo",
                    "quantidade_diarias": 1.0,
                    "diarias_r": 100.0,
                    "passagens_e_taxas_iniciais_r": 20.0,
                    "total_r": 120.0,
                }
            ],
            "custo_com_bilhetes_remarcados_nao_utilizados_cancelados_r": {
                "passagens_e_taxas_iniciais_r": 0.0,
                "total_r": 0.0,
            },
            "sub_total": {
                "quantidade_diarias": 1.0,
                "diarias_r": 100.0,
                "passagens_e_taxas_iniciais_r": 20.0,
                "total_r": 120.0,
            },
            "total_adicional_r": 0.0,
            "descontos_r": 0.0,
            "restituicao_r": 0.0,
            "reembolso_r": 0.0,
            "total_da_viagem_r": 120.0,
            "descricao_do_motivo_da_viagem": description,
        }
    )


class ExtratorTests(unittest.TestCase):
    def test_report_link_is_resolved_against_authenticated_scdp_page(self) -> None:
        report_url = resolve_report_url(
            "https://www2.scdp.gov.br/novoscdp/home.xhtml",
            "pages/relatorio/relatorio_viagem.xhtml?faces-redirect=true",
        )

        self.assertEqual(
            report_url,
            "https://www2.scdp.gov.br/novoscdp/pages/relatorio/"
            "relatorio_viagem.xhtml?faces-redirect=true",
        )


class LoginRoutingTests(unittest.IsolatedAsyncioTestCase):
    async def test_existing_govbr_login_does_not_search_for_scdp_button(self) -> None:
        from unittest.mock import AsyncMock, MagicMock, patch

        from playwright.async_api import Page

        from scdp_automation.extrator import start_login_if_needed

        page = MagicMock(spec=Page)
        page.url = "https://sso.acesso.gov.br/login"
        page.locator.side_effect = AssertionError("Botão SCDP não existe no gov.br")
        with (
            patch("scdp_automation.extrator.load_credentials"),
            patch(
                "scdp_automation.extrator.authenticate_gov_br", new=AsyncMock()
            ) as authenticate,
        ):
            await start_login_if_needed(page)
        authenticate.assert_awaited_once()

    async def test_authenticated_menu_can_finish_loading_before_login_check(
        self,
    ) -> None:
        from unittest.mock import AsyncMock, MagicMock

        from playwright.async_api import Page

        from scdp_automation.extrator import start_login_if_needed

        menu = MagicMock()
        menu.wait_for = AsyncMock()
        page = MagicMock(spec=Page)
        page.url = "https://www2.scdp.gov.br/novoscdp/home.xhtml"
        page.get_by_role.return_value = menu
        await start_login_if_needed(page)
        menu.wait_for.assert_awaited_once_with(state="visible", timeout=2_000)
        page.locator.assert_not_called()

    async def test_authenticated_menu_loading_while_login_link_appears_skips_login(
        self,
    ) -> None:
        from unittest.mock import AsyncMock, MagicMock

        from playwright.async_api import Page, TimeoutError

        from scdp_automation.extrator import start_login_if_needed

        menu = MagicMock()
        menu.wait_for = AsyncMock(side_effect=TimeoutError("slow home"))
        menu.is_visible = AsyncMock(return_value=True)
        login_link = MagicMock()
        login_link.wait_for = AsyncMock()
        login_link.click = AsyncMock()
        page = MagicMock(spec=Page)
        page.url = "https://www2.scdp.gov.br/novoscdp/home.xhtml"
        page.get_by_role.return_value = menu
        page.locator.return_value = login_link

        with (
            patch("scdp_automation.extrator.authenticate_gov_br", new=AsyncMock()),
            patch("scdp_automation.extrator.load_credentials"),
        ):
            await start_login_if_needed(page)

        menu.is_visible.assert_awaited_once()
        login_link.click.assert_not_awaited()

    async def test_click_returning_to_scdp_home_uses_existing_session(self) -> None:
        from unittest.mock import AsyncMock, MagicMock

        from playwright.async_api import Page, TimeoutError

        from scdp_automation.extrator import start_login_if_needed

        menu = MagicMock()
        menu.wait_for = AsyncMock(side_effect=[TimeoutError("home slow"), None])
        menu.is_visible = AsyncMock(return_value=False)
        login_link = MagicMock()
        login_link.wait_for = AsyncMock()
        login_link.click = AsyncMock()
        page = MagicMock(spec=Page)
        page.url = "https://www2.scdp.gov.br/novoscdp/home.xhtml"
        page.get_by_role.return_value = menu
        page.locator.return_value = login_link

        async def return_to_scdp(url_pattern: object, **_: object) -> None:
            page.url = "https://www2.scdp.gov.br/novoscdp/pages/main.xhtml"

        page.wait_for_url = AsyncMock(side_effect=return_to_scdp)
        with patch(
            "scdp_automation.extrator.authenticate_gov_br", new=AsyncMock()
        ) as authenticate:
            await start_login_if_needed(page)

        wait_call = page.wait_for_url.await_args
        assert wait_call is not None
        self.assertEqual(
            wait_call.args[0].pattern,
            r"^https://(?:sso\.acesso\.gov\.br/login|www2\.scdp\.gov\.br/novoscdp/pages/main\.xhtml)",
        )
        menu.wait_for.assert_awaited_with(state="visible", timeout=30_000)
        authenticate.assert_not_awaited()


class DescriptionCollectionTests(unittest.IsolatedAsyncioTestCase):
    async def test_single_tab_serializes_and_checkpoints_each_description(self) -> None:
        from tempfile import TemporaryDirectory
        from unittest.mock import AsyncMock, MagicMock

        from playwright.async_api import Page

        page = MagicMock(spec=Page)
        page.bring_to_front = AsyncMock()
        trips = [
            Viagem.model_construct(
                numero_da_solicitacao=f"00000{i}/26",
                descricao_do_motivo_da_viagem=None,
            )
            for i in range(1, 3)
        ]
        active_queries = 0

        async def consult(*_: object) -> str:
            nonlocal active_queries
            self.assertEqual(active_queries, 0)
            active_queries += 1
            await asyncio.sleep(0)
            active_queries -= 1
            return "Descrição consultada"

        with TemporaryDirectory() as directory:
            output = Path(directory) / "viagens.json"
            with (
                patch(
                    "scdp_automation.extrator._consult_with_safe_failure",
                    new=AsyncMock(side_effect=consult),
                ),
                patch("scdp_automation.extrator.save_json") as save_json,
            ):
                await collect_pending_descriptions(page, trips, trips, output)

        self.assertEqual(save_json.call_count, 2)
        self.assertEqual(page.bring_to_front.await_count, 2)
        self.assertEqual(
            [trip.descricao_do_motivo_da_viagem for trip in trips],
            ["Descrição consultada", "Descrição consultada"],
        )


class QueryRetryTests(unittest.IsolatedAsyncioTestCase):
    async def test_timeout_retries_read_only_query(self) -> None:
        from unittest.mock import MagicMock, patch

        from playwright.async_api import Page, TimeoutError

        from scdp_automation.extrator import consult_trip_reason

        with patch(
            "scdp_automation.extrator._consult_trip_reason",
            side_effect=[TimeoutError("timeout"), "Motivo completo"],
        ):
            reason = await consult_trip_reason(MagicMock(spec=Page), "000548/26")
        self.assertEqual(reason, "Motivo completo")

    async def test_repeated_timeout_is_not_saved_as_empty_description(self) -> None:
        from unittest.mock import MagicMock, patch

        from playwright.async_api import Page, TimeoutError

        from scdp_automation.extrator import consult_trip_reason

        with (
            patch(
                "scdp_automation.extrator._consult_trip_reason",
                side_effect=TimeoutError("timeout"),
            ),
            self.assertRaises(TimeoutError),
        ):
            await consult_trip_reason(MagicMock(spec=Page), "000548/26")


class WorkbookPublishIntegrationTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        from tests.support.workbooks import final_template_fixture

        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        template = Path(directory.name) / "template.xlsx"
        final_template_fixture(template)
        default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", template)
        default.start()
        self.addCleanup(default.stop)

    def test_checkpoint_is_saved_before_workbook_publication(self) -> None:
        trips = [make_valid_trip("000001/26")]
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        checkpoint = Path(directory.name) / "checkpoint.json"
        workbook = Path(directory.name) / "test-workbook.xlsx"
        calls: list[str] = []

        with (
            patch(
                "scdp_automation.extrator.save_json",
                side_effect=lambda *_: calls.append("checkpoint"),
            ) as save,
            patch(
                "scdp_automation.extrator.publish_workbook",
                side_effect=lambda *_: calls.append("workbook"),
            ) as publish,
        ):
            result = save_checkpoint_and_publish(
                trips, checkpoint, workbook, history=OutputHistory(checkpoint, workbook)
            )

        self.assertEqual(calls, ["checkpoint", "workbook"])
        save.assert_called_once_with(checkpoint, trips)
        publish.assert_called_once_with(trips, workbook)
        self.assertIsNone(result)

    def test_workbook_failure_keeps_complete_checkpoint(self) -> None:
        trips = [make_valid_trip("000001/26"), make_valid_trip("000002/26")]

        with TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "viagens.json"
            workbook = Path(directory) / "gastos.xlsx"
            with (
                patch("scdp_automation.extrator.save_json", wraps=save_json) as save,
                patch(
                    "scdp_automation.extrator.publish_workbook",
                    side_effect=OSError("workbook unavailable"),
                ) as publish,
                self.assertRaisesRegex(OSError, "workbook unavailable"),
            ):
                save_checkpoint_and_publish(
                    trips,
                    checkpoint,
                    workbook,
                    history=OutputHistory(checkpoint, workbook),
                )

            save.assert_called_once_with(checkpoint, trips)
            publish.assert_called_once_with(trips, workbook)
            self.assertEqual(
                [trip.numero_da_solicitacao for trip in load_trips(checkpoint)],
                ["000001/26", "000002/26"],
            )

    def test_rollover_publication_failure_keeps_history_until_later_success(self):
        import os

        from openpyxl import load_workbook

        from tests.support.workbooks import final_template_fixture

        with TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.xlsx"
            final_template_fixture(template)
            output = root / "output"
            output.mkdir()
            old_json = output / "viagens_scdp_2026.json"
            old_book = output / "gastos_scdp_2026.xlsx"
            old_trips = [make_valid_trip("000001/26")]
            save_json(old_json, old_trips)
            xlsx_output.publish_workbook(old_trips, old_book, template_path=template)
            old = load_workbook(old_book)
            old["APOIO"]["D2"] = 5000
            old.save(old_book)
            old.close()
            old_contents = (old_json.read_bytes(), old_book.read_bytes())
            backup = output / "backup"
            backup.mkdir()
            prior_contents = {}
            for index in range(5):
                identifier = f"20260101T00000000000{index}Zabcdefgh".replace(
                    "abcdefgh", "abcdef01"
                )
                for stem, extension in (
                    ("viagens_scdp_2026", "json"),
                    ("gastos_scdp_2026", "xlsx"),
                ):
                    path = backup / f"{stem}.history.{identifier}.{extension}"
                    path.write_text(f"previous-{index}-{extension}")
                    prior_contents[path] = path.read_bytes()
            manual_recovery = backup / "recovery.manual.xlsx"
            manual_recovery.write_text("manual")
            new_json = output / "viagens_scdp_2027.json"
            new_book = output / "gastos_scdp_2027.xlsx"
            trips = [make_valid_trip("000002/27")]
            replace = os.replace

            def block_workbook(source, destination):
                if str(source).endswith(".candidate.xlsx"):
                    raise PermissionError("Excel holds workbook")
                return replace(source, destination)

            with patch.object(xlsx_output, "DEFAULT_TEMPLATE", template):
                with (
                    OutputHistory(new_json, new_book) as history,
                    patch(
                        "scdp_automation.xlsx_output.os.replace",
                        side_effect=block_workbook,
                    ),
                    self.assertRaises(xlsx_output.WorkbookPublishError),
                ):
                    save_checkpoint_and_publish(
                        trips, new_json, new_book, history=history
                    )
                self.assertEqual(
                    (old_json.read_bytes(), old_book.read_bytes()), old_contents
                )
                self.assertTrue(
                    all(
                        path.read_bytes() == content
                        for path, content in prior_contents.items()
                    )
                )
                recoveries = list(backup.glob("recovery.*.xlsx"))
                self.assertEqual(len(recoveries), 2)
                self.assertFalse(list(output.glob("*.candidate.xlsx")))
                self.assertFalse(new_book.exists())
                with OutputHistory(new_json, new_book) as history:
                    save_checkpoint_and_publish(
                        trips, new_json, new_book, history=history
                    )
                    save_checkpoint_and_publish(
                        trips, new_json, new_book, history=history
                    )
            self.assertFalse(old_json.exists())
            self.assertFalse(old_book.exists())
            self.assertEqual(list(backup.glob("recovery.*.xlsx")), [manual_recovery])
            self.assertEqual(len(list(backup.glob("*.history.*"))), 8)
            new = load_workbook(new_book)
            self.addCleanup(new.close)
            self.assertEqual(new["BASE VIAGENS"]["A2"].value, "000002/27")
            self.assertNotEqual(new["APOIO"]["D2"].value, 5000)

    async def run_with_mocks(
        self,
        trips: list[Viagem],
        args: list[str],
        backup_path: Path | None,
        verify_session: bool = False,
    ) -> tuple[MagicMock, Mock, AsyncMock, Mock]:
        page = MagicMock()
        page.url = "https://www2.scdp.gov.br/novoscdp/home.xhtml"
        page.bring_to_front = AsyncMock()
        context = MagicMock()
        context.pages = [page]
        browser = MagicMock()
        browser.contexts = [context]
        browser.close = AsyncMock()
        playwright_manager = MagicMock()
        playwright_manager.__aenter__ = AsyncMock(return_value=object())
        playwright_manager.__aexit__ = AsyncMock(return_value=False)
        session = OutputHistory(OUTPUT_PATH, WORKBOOK_PATH)
        observed = []

        def check_session(stage):
            if verify_session:
                from filelock import FileLock, Timeout

                contender = FileLock(session.lock.lock_file, timeout=0)
                with (
                    self.assertRaises(Timeout, msg=f"session unlocked during {stage}"),
                    contender,
                ):
                    pass
                observed.append(stage)

        def preflight(*paths):
            check_session("preflight")
            preflight_extraction(*paths)

        def read_previous(*_):
            check_session("load")
            return []

        async def connect(*_):
            check_session("browser")
            return browser

        def publish(*_):
            check_session("publish")

        def checkpoint(*_):
            check_session("checkpoint")

        async def consult(*_):
            check_session("pending")

        workbook_publisher = Mock(side_effect=publish)
        pending_descriptions = AsyncMock(side_effect=consult)
        logger_info = Mock()

        with (
            patch("scdp_automation.extrator.configure_logging"),
            patch(
                "scdp_automation.extrator.preflight_extraction", side_effect=preflight
            ),
            patch("scdp_automation.extrator.load_trips", side_effect=read_previous),
            patch(
                "scdp_automation.extrator.async_playwright",
                return_value=playwright_manager,
            ),
            patch(
                "scdp_automation.extrator.selected_profile_directory",
                return_value="Default",
            ),
            patch("scdp_automation.extrator.Path.is_dir", return_value=True),
            patch("scdp_automation.extrator.Path.is_file", return_value=True),
            patch(
                "scdp_automation.extrator.connect_visible_chrome",
                new=AsyncMock(side_effect=connect),
            ),
            patch("scdp_automation.extrator.start_login_if_needed", new=AsyncMock()),
            patch("scdp_automation.extrator.wait_for_login", new=AsyncMock()),
            patch("scdp_automation.extrator.open_annual_cch_report", new=AsyncMock()),
            patch(
                "scdp_automation.extrator.collect_listing",
                new=AsyncMock(return_value=trips),
            ),
            patch("scdp_automation.extrator.save_json", side_effect=checkpoint),
            patch("scdp_automation.extrator.OutputHistory") as history_factory,
            patch("scdp_automation.extrator.publish_workbook", new=workbook_publisher),
            patch(
                "scdp_automation.extrator.collect_pending_descriptions",
                new=pending_descriptions,
            ),
            patch("scdp_automation.extrator.logger.info", new=logger_info),
        ):
            if verify_session:
                history_factory.return_value = session
                with (
                    patch.object(session, "archive_previous", return_value=backup_path),
                    patch.object(session, "prune"),
                ):
                    await run(args)
                self.assertEqual(
                    observed,
                    [
                        "preflight",
                        "load",
                        "browser",
                        "preflight",
                        "checkpoint",
                        "publish",
                        "pending",
                        "preflight",
                        "checkpoint",
                        "publish",
                    ],
                )
                self.assertFalse(session.lock.is_locked)
                return page, workbook_publisher, pending_descriptions, logger_info
            history_factory.return_value.__enter__.return_value = (
                history_factory.return_value
            )
            history_factory.return_value.archive_previous.return_value = backup_path
            await run(args)

        return page, workbook_publisher, pending_descriptions, logger_info

    async def test_output_session_covers_json_browser_checkpoints_and_publications(
        self,
    ):
        await self.run_with_mocks([make_valid_trip()], [], None, verify_session=True)

    async def test_limit_does_not_skip_workbook_publication(self) -> None:
        trips = [
            make_valid_trip("000001/26"),
            make_valid_trip("000002/26"),
            make_valid_trip("000003/26"),
        ]
        backup = Path("output/gastos.backup.xlsx")

        page, workbook_publisher, pending, _ = await self.run_with_mocks(
            trips, ["--limite", "1"], backup
        )

        self.assertEqual(workbook_publisher.call_count, 2)
        workbook_publisher.assert_called_with(trips, WORKBOOK_PATH)
        pending.assert_awaited_once_with(page, trips, [trips[0]], OUTPUT_PATH)

    async def test_classification_reminder_has_no_trip_data(self) -> None:
        secrets = ("000999/26", "NOME_PRIVADO_TESTE", "DESCRICAO_PRIVADA_TESTE")
        trip = make_valid_trip(secrets[0], secrets[1], secrets[2])
        backup = Path("output/gastos.backup.xlsx")

        _, workbook_publisher, _, logger_info = await self.run_with_mocks(
            [trip], [], backup
        )

        self.assertEqual(workbook_publisher.call_count, 2)
        workbook_publisher.assert_called_with([trip], WORKBOOK_PATH)
        log_calls = repr(logger_info.call_args_list)
        for secret in secrets:
            self.assertNotIn(secret, log_calls)
        self.assertIn("BASE VIAGENS", log_calls)
        self.assertIn("APOIO", log_calls)
        self.assertIn(str(WORKBOOK_PATH.resolve()), log_calls)
        self.assertIn(str(backup.resolve()), log_calls)


class BrowserModuleTests(unittest.TestCase):
    def test_entry_points_use_shared_visible_chrome_connector(self) -> None:
        from scdp_automation import cli, extrator, navegador_chrome

        self.assertIs(
            cli.connect_visible_chrome, navegador_chrome.connect_visible_chrome
        )
        self.assertIs(
            extrator.connect_visible_chrome, navegador_chrome.connect_visible_chrome
        )


class CliLoginTests(unittest.TestCase):
    def test_help_does_not_trigger_profile_selection(self) -> None:
        from unittest.mock import patch

        from scdp_automation import cli

        with (
            patch("sys.argv", ["scdp-extrair", "--help"]),
            redirect_stdout(StringIO()),
            patch("scdp_automation.cli.prepare_chrome_profile") as prepare_profile,
            self.assertRaises(SystemExit) as raised,
        ):
            cli.main()

        self.assertEqual(raised.exception.code, 0)
        prepare_profile.assert_not_called()

    def test_login_option_runs_authentication_without_extractor(self) -> None:
        from unittest.mock import MagicMock, patch

        from scdp_automation import cli

        login_coroutine = object()
        with (
            patch("sys.argv", ["scdp-extrair", "--login"]),
            patch(
                "scdp_automation.cli.login_only",
                new=MagicMock(return_value=login_coroutine),
            ) as login,
            patch("scdp_automation.cli.prepare_chrome_profile") as prepare_profile,
            patch(
                "scdp_automation.cli.run", return_value="extract-coroutine"
            ) as extract,
            patch("scdp_automation.cli.asyncio.run") as run_async,
        ):
            cli.main()

        login.assert_called_once_with()
        prepare_profile.assert_called_once()
        run_async.assert_called_once_with(login_coroutine)
        extract.assert_not_called()

    def test_profile_selection_runs_before_login_action(self) -> None:
        from unittest.mock import MagicMock, patch

        from scdp_automation import cli

        login_coroutine = object()
        with (
            patch(
                "sys.argv",
                [
                    "scdp-extrair",
                    "--selecionar-perfil-chrome",
                    "--login",
                ],
            ),
            patch("scdp_automation.cli.prepare_chrome_profile") as prepare_profile,
            patch(
                "scdp_automation.cli.login_only",
                new=MagicMock(return_value=login_coroutine),
            ) as login,
            patch("scdp_automation.cli.asyncio.run") as run_async,
        ):
            cli.main()

        prepare_profile.assert_called_once()
        self.assertTrue(prepare_profile.call_args.kwargs["force_reselect"])
        login.assert_called_once_with()
        run_async.assert_called_once_with(login_coroutine)

    def test_profile_selector_flag_is_removed_before_extractor_arguments(self) -> None:
        from unittest.mock import MagicMock, patch

        from scdp_automation import cli

        extraction_coroutine = object()
        with (
            patch(
                "sys.argv",
                [
                    "scdp-extrair",
                    "--selecionar-perfil-chrome",
                    "--limite",
                    "2",
                ],
            ),
            patch("scdp_automation.cli.prepare_chrome_profile"),
            patch(
                "scdp_automation.cli.run",
                new=MagicMock(return_value=extraction_coroutine),
            ) as extract,
            patch("scdp_automation.cli.asyncio.run") as run_async,
        ):
            cli.main()

        extract.assert_called_once_with(["--limite", "2"])
        run_async.assert_called_once_with(extraction_coroutine)


class SafeFailureTests(unittest.IsolatedAsyncioTestCase):
    async def test_query_failure_does_not_expose_original_exception_text(self) -> None:
        from unittest.mock import AsyncMock, MagicMock, patch

        from playwright.async_api import Error as PlaywrightError
        from playwright.async_api import Page

        from scdp_automation.extrator import _consult_with_safe_failure

        private_detail = "https://example.invalid/path?token=private-value"
        with (
            patch(
                "scdp_automation.extrator.consult_trip_reason",
                new=AsyncMock(side_effect=PlaywrightError(private_detail)),
            ),
            self.assertRaises(RuntimeError) as raised,
        ):
            await _consult_with_safe_failure(MagicMock(spec=Page), "999999/26", 1, 3)

        self.assertEqual(
            str(raised.exception),
            "Falha ao consultar descrição. O progresso foi preservado; execute novamente para retomar.",
        )
        self.assertTrue(raised.exception.__suppress_context__)
        self.assertNotIn(private_detail, str(raised.exception))
