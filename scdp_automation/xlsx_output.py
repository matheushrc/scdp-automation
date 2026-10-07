"""Refresh BASE from complete trips, preserving the final template's manual sheets."""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from collections.abc import Sequence
from copy import copy
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from scdp_automation.config import current_year
from scdp_automation.relatorio import Viagem
from scdp_automation.xlsx_models import TripSummary, summarize_trips
from scdp_automation.xlsx_validation import (
    WorkbookValidationError,
    same_formula,
    segment_formula,
    validate_workbook,
)

DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[1]
    / "output"
    / f"gastos_scdp_{current_year()}.xlsx"
)
_CHECKOUT = Path(__file__).resolve().parents[1]
_TEMPLATE_ROOT = (
    _CHECKOUT.parent.parent if _CHECKOUT.parent.name == ".worktrees" else _CHECKOUT
)
DEFAULT_TEMPLATE = _TEMPLATE_ROOT / "input" / "gastos_scdp_template.xlsx"
BASE_TABLE_NAME = "tblBaseViagens"
FILELOCK_TIMEOUT = 30


class WorkbookLockedError(TimeoutError):
    """The output lock could not be acquired before its timeout."""


class WorkbookBackupError(OSError):
    """The existing workbook could not be backed up before replacement."""


class WorkbookPublishError(RuntimeError):
    """A validated workbook could not be promoted to the output path."""


def install_decision_highlighting(base: Worksheet) -> None:
    """Distinguish undecided cancellations from missing travel classifications."""
    from openpyxl.formatting.rule import FormulaRule

    decision, segment, last = "R", "P", "R"
    red_formula = f'AND($C2="Cancelada",${decision}2="",$A2<>"")'
    yellow_formula = f'AND($A2<>"",${segment}2="")'
    managed_formulas = {red_formula, yellow_formula}
    for area in list(base.conditional_formatting):
        rules = base.conditional_formatting[area]
        rules[:] = [
            rule
            for rule in rules
            if not (
                rule.formula
                and len(rule.formula) == 1
                and rule.formula[0] in managed_formulas
            )
        ]
        if not rules:
            del base.conditional_formatting[str(area.sqref)]
        for rule in rules:
            if rule.priority < 3:
                rule.priority += 2
    for formula, color, priority in (
        (red_formula, "FFC7CE", 1),
        (yellow_formula, "FFF2CC", 2),
    ):
        rule = FormulaRule(
            formula=[formula], fill=PatternFill("solid", fgColor=color), stopIfTrue=True
        )
        rule.priority = priority
        base.conditional_formatting.add(f"A2:{last}1048576", rule)


def _load_workbook_for_refresh(path: Path) -> Workbook:
    try:
        workbook = load_workbook(path, data_only=False)
    except (OSError, InvalidFileException, BadZipFile, KeyError, ValueError) as error:
        raise WorkbookValidationError(
            f"Não foi possível abrir o workbook em {path}."
        ) from error
    try:
        validate_workbook(workbook)
    except Exception:
        workbook.close()
        raise
    return workbook


def _write_base_rows(
    workbook: Workbook,
    summaries: Sequence[TripSummary],
    manual_codes: dict[str, str],
    manual_decisions: dict[str, str | None] | None = None,
) -> None:
    base = workbook["BASE VIAGENS"]
    old_last_row = max(base.max_row, 2)
    target_last_row = max(2, len(summaries) + 1)
    style_source = [copy(cell._style) for cell in base[2]]
    height_source = base.row_dimensions[2].height

    if old_last_row > target_last_row:
        base.delete_rows(target_last_row + 1, old_last_row - target_last_row)

    for index, row in enumerate(range(2, target_last_row + 1)):
        if row > old_last_row:
            for column, style in enumerate(style_source, start=1):
                base.cell(row, column)._style = copy(style)
            if height_source is not None:
                base.row_dimensions[row].height = height_source

        if index >= len(summaries):
            values: tuple[object, ...] = (None,) * 11
            pcdp = ""
        else:
            summary = summaries[index]
            pcdp = summary.pcdp
            values = (
                summary.pcdp,
                summary.proposed,
                summary.status,
                summary.daily_count,
                summary.daily_amount,
                summary.ticket_amount,
                summary.additional_amount,
                summary.discount_amount,
                summary.restitution_amount,
                summary.reimbursement_amount,
                summary.trip_total,
            )

        for column, value in enumerate(values, start=1):
            base.cell(row, column).value = value
        base.cell(row, 16).value = segment_formula(row)
        base.cell(row, 15).value = (
            summary.description if index < len(summaries) else None
        )
        base.cell(row, 17).value = manual_codes.get(pcdp)
        base.cell(row, 18).value = (manual_decisions or {}).get(pcdp)
        dates = (
            (summary.start_date, summary.end_date, summary.verified_date)
            if index < len(summaries)
            else (None, None, None)
        )
        for column, value in enumerate(dates, 12):
            base.cell(row, column).value = value
            base.cell(row, column).number_format = "dd/mm/yyyy"
        base.cell(row, 4).number_format = "0.0"
        for column in range(5, 12):
            base.cell(row, column).number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'

    base.tables[BASE_TABLE_NAME].ref = f"A1:R{target_last_row}"
    base.tables[BASE_TABLE_NAME].tableColumns = []
    from scdp_automation.xlsx_presentation import format_base_sheet

    format_base_sheet(base)


def _validate_candidate(
    candidate_path: Path, summaries: Sequence[TripSummary] | None = None
) -> None:
    workbook = _load_workbook_for_refresh(candidate_path)
    try:
        base = workbook["BASE VIAGENS"]
        actual_rows = [
            row for row in range(2, base.max_row + 1) if base.cell(row, 1).value
        ]
        for row in actual_rows:
            if not same_formula(base.cell(row, 16).value, segment_formula(row)):
                raise WorkbookValidationError(
                    "A fórmula Segmento está ausente ou foi alterada na base."
                )
        if summaries is not None:
            if [base.cell(row, 1).value for row in actual_rows] != [
                summary.pcdp for summary in summaries
            ]:
                raise WorkbookValidationError(
                    "As PCDPs do candidato não correspondem à listagem completa."
                )
            for row, summary in zip(actual_rows, summaries, strict=True):
                expected = (
                    summary.pcdp,
                    summary.proposed,
                    summary.status,
                    summary.daily_count,
                    summary.daily_amount,
                    summary.ticket_amount,
                    summary.additional_amount,
                    summary.discount_amount,
                    summary.restitution_amount,
                    summary.reimbursement_amount,
                    summary.trip_total,
                )
                if tuple(base.cell(row, col).value for col in range(1, 12)) != expected:
                    raise WorkbookValidationError(
                        "Os campos agregados do candidato estão inconsistentes."
                    )
    finally:
        workbook.close()


def build_candidate(
    trips: Sequence[Viagem],
    current_path: Path,
    candidate_path: Path,
    *,
    template_path: Path | None = None,
) -> None:
    """Validate the required final template, refresh a separate candidate, then validate it."""
    source_path = template_path or DEFAULT_TEMPLATE
    resolved = (source_path.resolve(), current_path.resolve(), candidate_path.resolve())
    if len(set(resolved)) != 3:
        raise WorkbookValidationError(
            "Template, candidato e workbook publicado devem ter caminhos separados."
        )
    if candidate_path.exists():
        raise WorkbookValidationError(
            f"O caminho de candidato já existe: {candidate_path}"
        )
    template = _load_workbook_for_refresh(source_path)
    workbook = template
    try:
        summaries = summarize_trips(trips)
        manual_codes: dict[str, str] = {}
        manual_decisions: dict[str, str | None] = {}
        if current_path.exists():
            workbook = _load_workbook_for_refresh(current_path)
            base = workbook["BASE VIAGENS"]
            existing_pcdps = {
                base.cell(row, 1).value
                for row in range(2, base.max_row + 1)
                if base.cell(row, 1).value
            }
            missing = tuple(
                sorted(existing_pcdps - {summary.pcdp for summary in summaries})
            )
            if missing:
                raise WorkbookValidationError(
                    f"A listagem completa está sem {len(missing)} PCDP(s) já publicadas; reconcilie o relatório antes de atualizar.",
                    missing_pcdps=missing,
                )
            for row in range(2, base.max_row + 1):
                pcdp = base.cell(row, 1).value
                if pcdp:
                    if base.cell(row, 17).value:
                        manual_codes[pcdp] = base.cell(row, 17).value
                    manual_decisions[pcdp] = base.cell(row, 18).value or None
        else:
            # Template travel data is never an extraction input, even for matching keys.
            base = workbook["BASE VIAGENS"]
            for row in base.iter_rows(min_row=2):
                for cell in row:
                    cell.value = None
        _write_base_rows(workbook, summaries, manual_codes, manual_decisions)
        install_decision_highlighting(workbook["BASE VIAGENS"])
        updated = workbook["RESUMO GASTOS"]["M2"]
        updated.value = datetime.now(ZoneInfo("America/Sao_Paulo")).date()
        updated.number_format = "dd/mm/yyyy"
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        validate_workbook(workbook)
        candidate_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(candidate_path)
        _validate_candidate(candidate_path, summaries)
    except Exception:
        candidate_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
        if workbook is not template:
            template.close()


def _candidate_path_for(workbook_path: Path) -> Path:
    return workbook_path.with_name(
        f"{workbook_path.stem}.{uuid4().hex}.candidate{workbook_path.suffix}"
    )


def _backup_path_for(workbook_path: Path) -> Path:
    timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ")
    backup = workbook_path.with_name(
        f"{workbook_path.stem}.backup.{timestamp}{workbook_path.suffix}"
    )
    if backup.exists():
        backup = backup.with_name(f"{backup.stem}.{uuid4().hex[:8]}{backup.suffix}")
    return backup


def _workbook_lock_path(workbook_path: Path) -> str:
    identifier = hashlib.sha256(str(workbook_path.resolve()).encode()).hexdigest()
    return str(Path(tempfile.gettempdir()) / f"scdp-workbook-{identifier}.lock")


def publish_workbook(
    trips: Sequence[Viagem],
    workbook_path: Path = DEFAULT_WORKBOOK,
    *,
    template_path: Path | None = None,
    create_backup: bool = True,
) -> Path | None:
    """Serialize refreshes, validate a neighboring candidate, then replace atomically."""
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = _workbook_lock_path(workbook_path)
    lock = FileLock(lock_path, timeout=FILELOCK_TIMEOUT)
    try:
        with lock:
            candidate = _candidate_path_for(workbook_path)
            candidate_retained = False
            try:
                if template_path is None:
                    build_candidate(trips, workbook_path, candidate)
                else:
                    build_candidate(
                        trips, workbook_path, candidate, template_path=template_path
                    )
                _validate_candidate(candidate)
                backup: Path | None = None
                if create_backup and workbook_path.exists():
                    backup = _backup_path_for(workbook_path)
                    try:
                        shutil.copy2(workbook_path, backup)
                    except OSError as error:
                        backup.unlink(missing_ok=True)
                        raise WorkbookBackupError(
                            f"Não foi possível criar o backup em {backup}; "
                            f"o workbook publicado permanece intacto: {error}"
                        ) from error
                try:
                    os.replace(candidate, workbook_path)
                except OSError as error:
                    candidate_retained = True
                    raise WorkbookPublishError(
                        "Não foi possível substituir o workbook; feche o arquivo no Excel e "
                        f"tente novamente. O candidato validado permanece em {candidate}."
                    ) from error
                return backup
            except Exception:
                if not candidate_retained:
                    candidate.unlink(missing_ok=True)
                raise
    except Timeout as error:
        raise WorkbookLockedError(
            f"O workbook está em uso por outra atualização; feche-o e tente novamente: {workbook_path}"
        ) from error
