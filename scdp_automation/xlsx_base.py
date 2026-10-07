"""Replace only BASE with the current template inside the existing workbook."""

from copy import copy, deepcopy

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from scdp_automation.xlsx_validation import WorkbookValidationError

MANAGED_BASE_NAMES = frozenset("Viagens" + column for column in "ABCDEFGHIJKLMNOPQ")


def _copy_style(source, target) -> None:
    # Style indices belong to their workbook's registries. Assigning objects lets
    # openpyxl register each component in the destination registry.
    for attribute in ("font", "fill", "border", "alignment", "protection"):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.number_format = source.number_format


def rebuild_base_from_template(workbook: Workbook, template: Workbook) -> Worksheet:
    """Rebuild BASE, retaining the destination container and all manual sheets."""
    source = template["BASE VIAGENS"]
    if source._charts or source._images or source._pivots or source.legacy_drawing:
        raise WorkbookValidationError(
            "BASE VIAGENS do template contém desenho ou elemento não suportado."
        )
    old = workbook["BASE VIAGENS"]
    position = workbook.index(old)
    workbook.remove(old)
    target = workbook.create_sheet("BASE VIAGENS", position)
    for row in source:
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            copied = target.cell(cell.row, cell.column, cell.value)
            _copy_style(cell, copied)
            if cell.hyperlink:
                copied.hyperlink = copy(cell.hyperlink)
            if cell.comment:
                copied.comment = copy(cell.comment)
    for attribute in ("row_dimensions", "column_dimensions"):
        for key, dimension in getattr(source, attribute).items():
            copied = copy(dimension)
            copied.parent = target
            copied._style = None
            _copy_style(dimension, copied)
            getattr(target, attribute)[key] = copied
    for attribute in (
        "sheet_format",
        "sheet_properties",
        "views",
        "sheet_state",
        "protection",
        "auto_filter",
        "data_validations",
        "conditional_formatting",
        "print_options",
        "page_margins",
        "page_setup",
        "row_breaks",
        "col_breaks",
        "HeaderFooter",
    ):
        setattr(target, attribute, deepcopy(getattr(source, attribute)))
    target.freeze_panes = source.freeze_panes
    target.print_area = source.print_area
    target.print_title_rows = source.print_title_rows
    target.print_title_cols = source.print_title_cols
    for merged in source.merged_cells:
        target.merge_cells(str(merged))
    for table in source.tables.values():
        target.add_table(deepcopy(table))
    for name in old.defined_names.values():
        if (
            name.name not in source.defined_names
            and name.name not in MANAGED_BASE_NAMES | {"SCDPLayoutVersion"}
        ):
            preserved_name = deepcopy(name)
            preserved_name.localSheetId = position
            target.defined_names.add(preserved_name)
    for name in source.defined_names.values():
        copied_name = deepcopy(name)
        copied_name.localSheetId = position
        target.defined_names.add(copied_name)
    for name in MANAGED_BASE_NAMES | {"SCDPLayoutVersion"}:
        workbook.defined_names.pop(name, None)
        definition = template.defined_names.get(name)
        if definition is not None:
            workbook.defined_names.add(deepcopy(definition))
    return target
