"""Compact presentation for the editable spending inputs."""

from __future__ import annotations

import math
from textwrap import wrap

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

BASE_WIDTHS = (15, 26, 18, 12, 14, 14, 14, 14, 14, 14, 16, 15, 15, 15, 48, 17, 20, 16)
DATE_HEADERS = {
    "Data de início da viagem",
    "Data de término da viagem",
    "Data da última verificação",
}


def _compact_sheet(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    widths = widths[: sheet.max_column]
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
        cell.font = Font(color="FFFFFFFF", bold=True)
        cell.alignment = alignment
    sheet.sheet_format.defaultRowHeight = 32
    sheet.freeze_panes = "A2"
    for column, width in enumerate(widths, 1):
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.width = width
        dimension.hidden = sheet.cell(1, column).value in DATE_HEADERS
        dimension.alignment = alignment
    for row in sheet.iter_rows(max_col=len(widths)):
        lines = 1
        for cell, width in zip(row, widths, strict=True):
            cell.alignment = alignment
            if cell.value is None or cell.data_type == "f":
                continue
            if sheet.column_dimensions[cell.column_letter].hidden:
                continue
            for paragraph in str(cell.value).splitlines():
                lines = max(lines, len(wrap(paragraph, width=max(1, width - 2))))
        height = max(32, math.ceil(14 * lines + 8))
        sheet.row_dimensions[row[0].row].height = min(
            409, max(64, height) if row[0].row == 1 else height
        )


def format_base_sheet(sheet: Worksheet) -> None:
    _compact_sheet(sheet, BASE_WIDTHS)
