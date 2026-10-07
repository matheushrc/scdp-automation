import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from scdp_automation import xlsx_output
from scdp_automation.xlsx_models import TripSummary, summarize_trips
from scdp_automation.xlsx_output import build_candidate, publish_workbook
from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import (
    final_template_fixture,
    make_trip,
    read_base_rows,
    write_existing_workbook,
)


def workbook_set_up(case):
    directory = TemporaryDirectory()
    case.addCleanup(directory.cleanup)
    template = Path(directory.name) / "template.xlsx"
    final_template_fixture(template)
    default = patch.object(xlsx_output, "DEFAULT_TEMPLATE", template)
    default.start()
    case.addCleanup(default.stop)


class TripSummaryTests(unittest.TestCase):
    def test_summarize_trip_uses_trip_subtotal_once(self) -> None:
        trip = make_trip()

        summary = summarize_trips([trip])[0]

        self.assertIsInstance(summary, TripSummary)
        self.assertEqual(
            summary,
            TripSummary(
                pcdp="123456/26-2B",
                proposed="Pessoa Exemplo",
                status="Autorizada",
                daily_count=8.5,
                daily_amount=1234.56,
                ticket_amount=789.01,
                additional_amount=33.05,
                discount_amount=4.01,
                restitution_amount=5.02,
                reimbursement_amount=6.03,
                trip_total=2063.66,
                start_date=date(2026, 3, 1),
                end_date=date(2026, 3, 5),
            ),
        )

    def test_summarize_trips_rejects_duplicate_full_pcdp(self) -> None:
        with self.assertRaisesRegex(ValueError, "PCDP duplicada na listagem") as error:
            summarize_trips([make_trip(), make_trip()])
        self.assertNotIn("123456/26-2B", str(error.exception))

        summaries = summarize_trips([make_trip("123456/26"), make_trip("123456/26-2B")])

        self.assertEqual(
            [summary.pcdp for summary in summaries], ["123456/26", "123456/26-2B"]
        )
        complementary_summaries = summarize_trips(
            [make_trip("123456/26-1C"), make_trip("123456/26-2C")]
        )
        self.assertEqual(
            [summary.pcdp for summary in complementary_summaries],
            ["123456/26-1C", "123456/26-2C"],
        )


class WorkbookRefreshTests(unittest.TestCase):
    setUp = workbook_set_up

    def test_refresh_exports_request_description_before_classification(self) -> None:
        from openpyxl import load_workbook

        trip = make_trip()
        trip.descricao_do_motivo_da_viagem = "Participação em evento acadêmico."
        with TemporaryDirectory() as directory:
            current = Path(directory) / "missing.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            build_candidate([trip], current, candidate)
            workbook = load_workbook(candidate)
            headers = [cell.value for cell in workbook["BASE VIAGENS"][1]]
            self.assertIn("Descrição do pedido", headers)
            self.assertEqual(
                headers[-3:], ["Segmento", "Código de débito", "Descontar do curso?"]
            )
            column = headers.index("Descrição do pedido") + 1
            self.assertEqual(
                workbook["BASE VIAGENS"].cell(2, column).value,
                trip.descricao_do_motivo_da_viagem,
            )
            workbook.close()

    def test_refresh_preserves_manual_codes_by_full_pcdp_after_reordering(self) -> None:
        first = make_trip("123456/26-1C")
        second = make_trip("123456/26-2C")
        new_trip = make_trip("234567/26-1C")

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(
                current,
                [first, second],
                {
                    first.numero_da_solicitacao: "AGRONOMIA",
                    second.numero_da_solicitacao: "PPGEL +",
                },
            )
            original_bytes = current.read_bytes()

            build_candidate([second, new_trip, first], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            rows = read_base_rows(candidate)
            self.assertEqual(
                list(rows), ["123456/26-2C", "234567/26-1C", "123456/26-1C"]
            )
            self.assertEqual(rows["123456/26-2C"][16], "PPGEL +")
            self.assertEqual(rows["123456/26-1C"][16], "AGRONOMIA")

    def test_refresh_leaves_new_pcdp_code_blank(self) -> None:
        existing = make_trip("123456/26")
        added = make_trip("234567/26")

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(
                current, [existing], {existing.numero_da_solicitacao: "AGRONOMIA"}
            )

            build_candidate([existing, added], current, candidate)

            rows = read_base_rows(candidate)
            self.assertEqual(rows[existing.numero_da_solicitacao][16], "AGRONOMIA")
            self.assertIsNone(rows[added.numero_da_solicitacao][16])

    def test_refresh_rejects_unknown_code(self) -> None:
        trip = make_trip()

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(
                current, [trip], {trip.numero_da_solicitacao: "CODIGO DESCONHECIDO"}
            )
            original_bytes = current.read_bytes()

            with self.assertRaisesRegex(ValueError, "código de débito desconhecido"):
                build_candidate([trip], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())

    def test_refresh_rejects_duplicate_pcdp(self) -> None:
        trip = make_trip()

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(current, [trip])
            original_bytes = current.read_bytes()

            with self.assertRaisesRegex(ValueError, "PCDP duplicada"):
                build_candidate([trip, trip], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())

    def test_refresh_rejects_missing_previous_pcdp_without_changing_source(
        self,
    ) -> None:
        first = make_trip("123456/26")
        second = make_trip("234567/26")

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(current, [first, second])
            original_bytes = current.read_bytes()

            with self.assertRaises(WorkbookValidationError) as context:
                build_candidate([first], current, candidate)

            self.assertEqual(context.exception.missing_pcdps, ("234567/26",))
            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())

    def test_refresh_rejects_corrupt_workbook_without_changing_source(self) -> None:
        with TemporaryDirectory() as directory:
            current = Path(directory) / "corrupt.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            current.write_bytes(b"not an Excel workbook")
            original_bytes = current.read_bytes()

            with self.assertRaisesRegex(ValueError, "abrir o workbook"):
                build_candidate([make_trip()], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())

    def test_refresh_rejects_missing_sheet_workbook_without_changing_source(
        self,
    ) -> None:
        from openpyxl import Workbook

        with TemporaryDirectory() as directory:
            current = Path(directory) / "missing-sheet.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            Workbook().save(current)
            original_bytes = current.read_bytes()

            with self.assertRaisesRegex(WorkbookValidationError, "três worksheets"):
                build_candidate([make_trip()], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())

    def test_refresh_rejects_nonnumeric_allocation_without_changing_source(
        self,
    ) -> None:
        trip = make_trip()

        with TemporaryDirectory() as directory:
            current = Path(directory) / "current.xlsx"
            candidate = Path(directory) / "candidate.xlsx"
            write_existing_workbook(
                current, [trip], allocations={"AGRONOMIA": "dez mil"}
            )
            original_bytes = current.read_bytes()

            with self.assertRaisesRegex(WorkbookValidationError, "numérico"):
                build_candidate([trip], current, candidate)

            self.assertEqual(current.read_bytes(), original_bytes)
            self.assertFalse(candidate.exists())


class WorkbookPublicationTests(unittest.TestCase):
    setUp = workbook_set_up

    def test_first_publish_creates_workbook_without_backup(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "gastos.xlsx"

            backup = publish_workbook([make_trip()], path)

            self.assertIsNone(backup)
            self.assertTrue(path.exists())
            self.assertEqual(list(Path(directory).glob("*.backup.*.xlsx")), [])

    def test_refresh_preserves_manual_data_without_publisher_backup(self) -> None:
        trip = make_trip()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "gastos.xlsx"
            publish_workbook([trip], path)
            from openpyxl import load_workbook

            original = load_workbook(path)
            original["APOIO"]["D2"] = 5000.0
            original.save(path)
            original.close()
            self.assertIsNone(publish_workbook([trip], path))
            self.assertEqual(
                {p.name for p in Path(directory).iterdir()}, {"gastos.xlsx"}
            )
            refreshed = load_workbook(path)
            self.addCleanup(refreshed.close)
            self.assertEqual(refreshed["APOIO"]["D2"].value, 5000.0)

    def test_invalid_candidate_keeps_published_workbook_unchanged(self) -> None:
        trip = make_trip()

        with TemporaryDirectory() as directory:
            path = Path(directory) / "gastos.xlsx"
            publish_workbook([trip], path)
            original_bytes = path.read_bytes()

            def build_corrupt_candidate(trips, current, candidate):
                from openpyxl import load_workbook

                build_candidate(trips, current, candidate)
                book = load_workbook(candidate)
                book["BASE VIAGENS"]["Q2"] = "DESCONHECIDO"
                book.save(candidate)
                book.close()

            with (
                patch(
                    "scdp_automation.xlsx_output.build_candidate",
                    new=build_corrupt_candidate,
                ),
                self.assertRaisesRegex(WorkbookValidationError, "desconhecido"),
            ):
                publish_workbook([trip], path)

            self.assertEqual(path.read_bytes(), original_bytes)
            self.assertEqual(list(Path(directory).glob("*.backup.*.xlsx")), [])
            self.assertEqual(list(Path(directory).glob("*.candidate.xlsx")), [])

    def test_replace_failure_retains_validated_candidate(self) -> None:
        trip = make_trip()

        with TemporaryDirectory() as directory:
            path = Path(directory) / "gastos.xlsx"
            publish_workbook([trip], path)
            original_bytes = path.read_bytes()

            with (
                patch(
                    "scdp_automation.xlsx_output.os.replace",
                    side_effect=OSError("falha de promoção"),
                ),
                self.assertRaisesRegex(RuntimeError, "candidato validado"),
            ):
                publish_workbook([trip], path)

            self.assertEqual(path.read_bytes(), original_bytes)
            self.assertEqual(list(Path(directory).glob("*.candidate.xlsx")), [])
            recoveries = list((Path(directory) / "backup").glob("recovery.*.xlsx"))
            self.assertEqual(len(recoveries), 1)
            self.assertTrue(recoveries[0].is_file())
            from openpyxl import load_workbook

            recovered = load_workbook(recoveries[0])
            self.addCleanup(recovered.close)
            self.assertEqual(
                recovered["BASE VIAGENS"]["A2"].value, trip.numero_da_solicitacao
            )

    def test_recovery_transfer_failure_retains_validated_candidate(self) -> None:
        from openpyxl import load_workbook

        from scdp_automation.xlsx_output import WorkbookPublishError

        mkdir = Path.mkdir
        for failure_stage in ("mkdir", "rename"):
            with (
                self.subTest(failure_stage=failure_stage),
                TemporaryDirectory() as directory,
            ):
                path = Path(directory) / "gastos.xlsx"
                publish_workbook([make_trip()], path)
                original_bytes = path.read_bytes()

                def fail_recovery_mkdir(target, *args, **kwargs):
                    if target.name == "backup":
                        raise PermissionError("recovery unavailable")
                    return mkdir(target, *args, **kwargs)

                transfer_failure = (
                    patch.object(Path, "mkdir", new=fail_recovery_mkdir)
                    if failure_stage == "mkdir"
                    else patch.object(
                        Path,
                        "rename",
                        side_effect=PermissionError("recovery unavailable"),
                    )
                )
                with (
                    patch(
                        "scdp_automation.xlsx_output.os.replace",
                        side_effect=PermissionError("Excel holds workbook"),
                    ),
                    transfer_failure,
                    self.assertRaises(WorkbookPublishError) as raised,
                ):
                    publish_workbook([make_trip(), make_trip(pcdp="000002/26")], path)
                self.assertEqual(path.read_bytes(), original_bytes)
                candidates = list(Path(directory).glob("*.candidate.xlsx"))
                self.assertEqual(len(candidates), 1)
                self.assertIsInstance(raised.exception, WorkbookPublishError)
                self.assertIn(str(candidates[0]), str(raised.exception))
                self.assertIn("recovery unavailable", str(raised.exception))
                self.assertIsInstance(raised.exception.__cause__, PermissionError)
                self.assertEqual(
                    str(raised.exception.__cause__), "Excel holds workbook"
                )
                candidate = load_workbook(candidates[0])
                try:
                    self.assertEqual(
                        {
                            candidate["BASE VIAGENS"]["A2"].value,
                            candidate["BASE VIAGENS"]["A3"].value,
                        },
                        {"123456/26-2B", "000002/26"},
                    )
                finally:
                    candidate.close()
                self.assertFalse(
                    list((Path(directory) / "backup").glob("recovery.*.xlsx"))
                )


class WorkbookValidationTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "template.xlsx"
        final_template_fixture(path)
        from openpyxl import load_workbook

        self.book = load_workbook(path)
        self.addCleanup(self.book.close)

    def test_header_only_base_table_is_rejected(self):
        from scdp_automation.xlsx_validation import validate_workbook

        self.book["BASE VIAGENS"].tables["tblBaseViagens"].ref = "A1:R1"
        with self.assertRaises(WorkbookValidationError):
            validate_workbook(self.book)

    def test_manual_formulas_and_missing_recalculation_flags_are_not_repaired(self):
        from scdp_automation.xlsx_validation import validate_workbook

        self.book["APOIO"]["F2"] = "=1+2"
        self.book["RESUMO GASTOS"]["F7"] = 99
        self.book.calculation.fullCalcOnLoad = False
        before = [
            (sheet.title, cell.coordinate, cell.value, cell.style_id)
            for sheet in self.book
            for row in sheet
            for cell in row
        ]
        validate_workbook(self.book)
        self.assertEqual(
            [
                (sheet.title, cell.coordinate, cell.value, cell.style_id)
                for sheet in self.book
                for row in sheet
                for cell in row
            ],
            before,
        )
        self.assertFalse(self.book.calculation.fullCalcOnLoad)

    def test_final_named_ranges_cannot_point_to_previous_classification_columns(self):
        from scdp_automation.xlsx_validation import validate_workbook

        name = self.book.defined_names["ViagensM"]
        name.attr_text = name.attr_text.replace("$Q$1", "$P$1")
        before = name.attr_text
        with self.assertRaises(WorkbookValidationError):
            validate_workbook(self.book)
        self.assertEqual(name.attr_text, before)

    def test_invalid_manual_codes_decisions_and_financial_values_are_rejected(self):
        from scdp_automation.xlsx_validation import validate_workbook

        cases = (
            ("APOIO", "A3", "AGRONOMIA"),
            ("APOIO", "A3", " INVÁLIDO"),
            ("APOIO", "A3", "INVÁLIDO*"),
            ("APOIO", "D2", float("inf")),
            ("APOIO", "G2", True),
            ("BASE VIAGENS", "R2", "Talvez"),
            ("BASE VIAGENS", "Q2", "DESCONHECIDO"),
            ("BASE VIAGENS", "K2", float("nan")),
            ("BASE VIAGENS", "D2", True),
        )
        self.book["BASE VIAGENS"]["A2"] = "111111/26"
        for sheet, coordinate, value in cases:
            with self.subTest(sheet=sheet, coordinate=coordinate, value=value):
                cell = self.book[sheet][coordinate]
                original = cell.value
                cell.value = value
                with self.assertRaises(WorkbookValidationError):
                    validate_workbook(self.book)
                cell.value = original

    def test_summary_named_ranges_cannot_point_to_wrong_totals(self):
        from scdp_automation.xlsx_validation import validate_workbook

        self.book.defined_names["ResumoF"].attr_text = self.book.defined_names[
            "ResumoF"
        ].attr_text.replace("$F$1", "$E$1")
        with self.assertRaises(WorkbookValidationError):
            validate_workbook(self.book)

    def test_essential_range_heights_are_validated_without_repair(self):
        from scdp_automation.xlsx_validation import validate_workbook

        summary = self.book["RESUMO GASTOS"]
        for coordinate in ("R1", "R2", "R3"):
            with self.subTest(coordinate=coordinate):
                original = summary[coordinate].value
                summary[coordinate] = "=1"
                with self.assertRaises(WorkbookValidationError):
                    validate_workbook(self.book)
                self.assertEqual(summary[coordinate].value, "=1")
                summary[coordinate] = original


class LimitedChoiceRangeTests(unittest.TestCase):
    def test_limited_source_choices_cover_base_and_candidate_extends_for_new_rows(self):
        from openpyxl import load_workbook

        with TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.xlsx"
            final_template_fixture(template)
            book = load_workbook(template)
            base = book["BASE VIAGENS"]
            for validation in base.data_validations.dataValidation:
                column = "Q" if "CodigosDebito" in validation.formula1 else "R"
                validation.sqref = f"{column}2:{column}3"
                if column == "Q":
                    validation.formula1 = "=CodigosDebito"
            book.save(template)
            book.close()
            before = template.read_bytes()
            current = root / "current.xlsx"
            try:
                publish_workbook(
                    [make_trip(f"{i:06d}/26") for i in range(5)],
                    current,
                    template_path=template,
                )
            except WorkbookValidationError as error:
                self.fail(f"Sufficient limited source choices rejected: {error}")
            result = load_workbook(current)
            self.addCleanup(result.close)
            for coordinate in ("Q6", "R6", "Q1048576", "R1048576"):
                self.assertTrue(
                    any(
                        coordinate in validation
                        for validation in result[
                            "BASE VIAGENS"
                        ].data_validations.dataValidation
                    )
                )
            self.assertEqual(template.read_bytes(), before)

    def test_source_choices_must_cover_all_table_and_populated_rows(self):
        from openpyxl import load_workbook

        from scdp_automation.xlsx_validation import validate_workbook

        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.xlsx"
            final_template_fixture(template)
            book = load_workbook(template)
            self.addCleanup(book.close)
            base = book["BASE VIAGENS"]
            for validation in base.data_validations.dataValidation:
                column = "Q" if "CodigosDebito" in validation.formula1 else "R"
                validation.sqref = f"{column}2:{column}3"
            base.tables["tblBaseViagens"].ref = "A1:R4"
            with self.assertRaises(WorkbookValidationError):
                validate_workbook(book)
            base.tables["tblBaseViagens"].ref = "A1:R2"
            base["A4"] = "000004/26"
            with self.assertRaises(WorkbookValidationError):
                validate_workbook(book)
