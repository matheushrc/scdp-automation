import shutil
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from scdp_automation.xlsx_output import build_candidate
from scdp_automation.xlsx_validation import validate_workbook
from tests.support.fidelity import assert_manual_sheets_preserved
from tests.support.recalculation import recalculate_workbook
from tests.support.workbooks import (
    final_template_fixture,
    make_trip,
    write_existing_workbook,
)


@unittest.skipUnless(shutil.which("libreoffice"), "LibreOffice unavailable")
class RecalculationTests(unittest.TestCase):
    def test_cancelled_trip_decisions_survive_layout_10_refresh(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, current, candidate = (
                root / name
                for name in ("template.xlsx", "current.xlsx", "candidate.xlsx")
            )
            final_template_fixture(template)
            cancelled, other = make_trip("111111/26"), make_trip("222222/26")
            cancelled.situacao_da_viagem = "Cancelada"
            cancelled.total_da_viagem_r = 100
            other.total_da_viagem_r = 25
            for decision, expected in ((None, "Pendente"), ("Sim", 100), ("Não", 0)):
                with self.subTest(decision=decision):
                    candidate.unlink(missing_ok=True)
                    write_existing_workbook(
                        current,
                        [cancelled, other],
                        {"111111/26": "AGRONOMIA", "222222/26": "PPGE"},
                    )
                    source = load_workbook(current)
                    source["BASE VIAGENS"]["R2"] = decision
                    source["BASE VIAGENS"]["R3"] = "Não"
                    source.save(current)
                    source.close()
                    original = current.read_bytes()
                    build_candidate(
                        [other, cancelled], current, candidate, template_path=template
                    )
                    source, actual = load_workbook(current), load_workbook(candidate)
                    try:
                        self.assertEqual(
                            actual.defined_names["SCDPLayoutVersion"].attr_text, '"10"'
                        )
                        self.assertEqual(
                            [
                                (
                                    actual["BASE VIAGENS"].cell(row, 1).value,
                                    actual["BASE VIAGENS"].cell(row, 17).value,
                                    actual["BASE VIAGENS"].cell(row, 18).value,
                                )
                                for row in (2, 3)
                            ],
                            [
                                ("222222/26", "PPGE", "Não"),
                                ("111111/26", "AGRONOMIA", decision),
                            ],
                        )
                        assert_manual_sheets_preserved(self, actual, source)
                    finally:
                        source.close()
                        actual.close()
                    recalculate_workbook(candidate)
                    cached = load_workbook(candidate, data_only=True)
                    try:
                        self.assertEqual(cached["RESUMO GASTOS"]["F7"].value, expected)
                    finally:
                        cached.close()
                    self.assertEqual(current.read_bytes(), original)

    def test_visible_results_recalculate_while_original_formulas_and_styles_survive(
        self,
    ):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, current, candidate = (
                root / name
                for name in ("template.xlsx", "current.xlsx", "candidate.xlsx")
            )
            final_template_fixture(template)
            trip = make_trip()
            trip.total_da_viagem_r = 156
            write_existing_workbook(
                current, [trip], {trip.numero_da_solicitacao: "AGRONOMIA"}
            )
            build_candidate([trip], current, candidate, template_path=template)
            workbook = load_workbook(candidate)
            font = copy(workbook["RESUMO GASTOS"]["B2"].font)
            formula = workbook["RESUMO GASTOS"]["F7"].value
            workbook.close()
            recalculate_workbook(candidate)
            cached = load_workbook(candidate, data_only=True)
            self.assertEqual(
                [
                    cached["RESUMO GASTOS"][c].value
                    for c in ("F7", "G7", "J7", "K7", "M7", "F20")
                ],
                [156, 244, 200, 400, 644, 156],
            )
            self.assertEqual(cached["APOIO"]["F2"].value, 1000)
            cached.close()
            workbook = load_workbook(candidate)
            validate_workbook(workbook)
            self.assertEqual(workbook["RESUMO GASTOS"]["F7"].value, formula)
            self.assertEqual(copy(workbook["RESUMO GASTOS"]["B2"].font), font)
            workbook["APOIO"]["D2"] = 500
            workbook.save(candidate)
            workbook.close()
            recalculate_workbook(candidate)
            cached = load_workbook(candidate, data_only=True)
            self.assertEqual(cached["RESUMO GASTOS"]["G7"].value, 344)
            self.assertEqual(cached["APOIO"]["F2"].value, 1100)
            cached.close()
