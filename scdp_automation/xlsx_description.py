"""Add request descriptions while preserving existing workbook references."""

from __future__ import annotations

from copy import copy

from openpyxl.workbook.workbook import Workbook

from scdp_automation.xlsx_adjustments import translate_base_references


def shift_description_references(formula: str, sheet: str) -> str:
    return translate_base_references(formula, sheet, {15: 16, 16: 17, 17: 18})


def migrate_description(workbook: Workbook) -> None:
    """Insert a description before the final three classification columns once."""
    from scdp_automation.xlsx_output import (
        BASE_HEADERS,
        BASE_TABLE_NAME,
        LEGACY_BASE_HEADERS,
        install_base_controls,
    )
    from scdp_automation.xlsx_presentation import format_base_sheet

    base = workbook["BASE VIAGENS"]
    headers = tuple(cell.value for cell in base[1])
    if headers == BASE_HEADERS:
        return
    if headers != LEGACY_BASE_HEADERS:
        raise ValueError("As colunas de BASE VIAGENS foram alteradas.")

    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.value = shift_description_references(cell.value, sheet.title)
    for name in workbook.defined_names.values():
        name.attr_text = shift_description_references("=" + name.attr_text, "")[1:]

    base.insert_cols(15)
    base.cell(1, 15)._style = copy(base.cell(1, 16)._style)
    for row in range(2, base.max_row + 1):
        base.cell(row, 15)._style = copy(base.cell(row, 2)._style)
    for column, header in enumerate(BASE_HEADERS, 1):
        base.cell(1, column).value = header
    for validation in base.data_validations.dataValidation:
        validation.sqref = " ".join(
            shift_description_references("=" + str(area), "BASE VIAGENS")[1:]
            for area in validation.sqref.ranges
        )
        if validation.formula1 and validation.formula1.startswith("="):
            validation.formula1 = shift_description_references(
                validation.formula1, "BASE VIAGENS"
            )
    for area in base.conditional_formatting:
        for rule in base.conditional_formatting[area]:
            if rule.formula:
                rule.formula = [
                    shift_description_references("=" + formula, "BASE VIAGENS")[1:]
                    for formula in rule.formula
                ]
    table = base.tables[BASE_TABLE_NAME]
    table.ref = f"A1:R{max(2, base.max_row)}"
    table.tableColumns = []
    if table.autoFilter:
        table.autoFilter.ref = table.ref
    install_base_controls(base)
    format_base_sheet(base)
