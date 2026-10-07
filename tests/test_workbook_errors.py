"""Expected workbook failures identify the file and preserve user data."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scdp_automation import cli
from scdp_automation.xlsx_output import (
    WorkbookPublishError,
    _load_workbook_for_refresh,
    validate_workbook_sources,
)
from scdp_automation.xlsx_validation import WorkbookValidationError
from tests.support.workbooks import final_template_fixture


class WorkbookErrorTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.template = self.root / "input" / "gastos_scdp_template.xlsx"
        self.output = self.root / "output" / "gastos.xlsx"
        self.template.parent.mkdir()
        self.output.parent.mkdir()
        final_template_fixture(self.template)

    def test_version_9_sources_are_rejected_without_changes(self):
        from scdp_automation.extrator import preflight_extraction

        checkpoint = self.output.with_suffix(".json")
        backup = self.output.parent / "backup" / "synthetic.xlsx"
        backup.parent.mkdir()
        checkpoint.write_text("[]")
        backup.write_bytes(b"synthetic backup")
        for role, source in (
            ("obrigatório", self.template),
            ("publicado", self.output),
        ):
            with self.subTest(role=role):
                final_template_fixture(self.template)
                final_template_fixture(self.output)
                book = load_workbook(source)
                book.defined_names["SCDPLayoutVersion"].attr_text = '"9"'
                if role == "publicado":
                    book["APOIO"].insert_cols(8)
                    book["APOIO"]["H1"] = "Transportes pago (R$)"
                book.save(source)
                book.close()
                paths = (self.template, self.output, checkpoint, backup)
                before = [path.read_bytes() for path in paths]
                with (
                    patch(
                        "scdp_automation.xlsx_output.DEFAULT_TEMPLATE", self.template
                    ),
                    self.assertRaises(WorkbookValidationError) as caught,
                ):
                    preflight_extraction(checkpoint, self.output)
                for text in (
                    str(source),
                    role,
                    "esperada 10" if role == "obrigatório" else "APOIO",
                ):
                    self.assertIn(text, str(caught.exception))
                self.assertEqual(before, [path.read_bytes() for path in paths])

    def test_incompatible_output_identifies_path_and_non_destructive_recovery(self):
        final_template_fixture(self.output)
        workbook = load_workbook(self.output)
        workbook.defined_names["SCDPLayoutVersion"].attr_text = '"8"'
        workbook["APOIO"].insert_cols(8)
        workbook["APOIO"]["H1"] = "Transportes pago (R$)"
        workbook.save(self.output)
        workbook.close()
        before = self.output.read_bytes(), self.template.read_bytes()
        with self.assertRaises(WorkbookValidationError) as caught:
            validate_workbook_sources(self.output, template_path=self.template)
        message = str(caught.exception)
        for text in (
            str(self.output),
            "publicado",
            "APOIO",
            str(self.template),
            "JSON",
            "--recriar-planilha",
            "reconciliação explícita",
            "JSON sozinho não recupera",
            "git pull",
        ):
            self.assertIn(text, message)
        self.assertNotIn("começar novamente", message)
        self.assertNotIn("fora da saída ativa", message)
        self.assertIsInstance(caught.exception.__cause__, WorkbookValidationError)
        self.assertEqual(before, (self.output.read_bytes(), self.template.read_bytes()))

    def test_missing_marker_in_template_requests_template_restore_only(self):
        workbook = load_workbook(self.template)
        del workbook.defined_names["SCDPLayoutVersion"]
        workbook.save(self.template)
        workbook.close()
        with self.assertRaises(WorkbookValidationError) as caught:
            validate_workbook_sources(self.output, template_path=self.template)
        message = str(caught.exception)
        for text in (str(self.template), "obrigatório", "ausente", "10", "Restaure"):
            self.assertIn(text, message)
        self.assertNotIn("JSON", message)

    def test_missing_template_and_corrupt_template_explain_restore(self):
        for contents in (None, b"not an xlsx"):
            with self.subTest(contents=contents):
                self.template.unlink(missing_ok=True)
                if contents is not None:
                    self.template.write_bytes(contents)
                with self.assertRaises(WorkbookValidationError) as caught:
                    validate_workbook_sources(self.output, template_path=self.template)
                self.assertIn(str(self.template), str(caught.exception))
                self.assertIn("Restaure", str(caught.exception))
                self.assertNotIn("Excel", str(caught.exception))

    def test_orphan_output_identifies_both_paths_and_preserves_files(self):
        from scdp_automation.extrator import preflight_extraction

        final_template_fixture(self.output)
        checkpoint = self.output.with_suffix(".json")
        before = self.output.read_bytes()
        with (
            patch("scdp_automation.xlsx_output.DEFAULT_TEMPLATE", self.template),
            self.assertRaises(WorkbookValidationError) as caught,
        ):
            preflight_extraction(checkpoint, self.output)
        self.assertIn(str(self.output), str(caught.exception))
        self.assertIn(str(checkpoint), str(caught.exception))
        self.assertIn("reconcilie", str(caught.exception))
        self.assertEqual(before, self.output.read_bytes())
        self.assertFalse(checkpoint.exists())

    def test_permission_failure_identifies_access_action(self):
        with (
            patch(
                "scdp_automation.xlsx_output.load_workbook",
                side_effect=PermissionError(),
            ),
            self.assertRaises(WorkbookValidationError) as caught,
        ):
            validate_workbook_sources(self.output, template_path=self.template)
        self.assertIn(str(self.template), str(caught.exception))
        self.assertIn("Excel", str(caught.exception))
        self.assertIn("acesso", str(caught.exception))

    def test_file_context_preserves_missing_pcdps_and_cause(self):
        original = WorkbookValidationError("PCDP ausente", missing_pcdps=("123/26",))
        with (
            patch(
                "scdp_automation.xlsx_output.validate_workbook", side_effect=original
            ),
            self.assertRaises(WorkbookValidationError) as caught,
        ):
            _load_workbook_for_refresh(self.template)
        self.assertIn(str(self.template), str(caught.exception))
        self.assertEqual(caught.exception.missing_pcdps, ("123/26",))
        self.assertIs(caught.exception.__cause__, original)

    def test_candidate_failure_identifies_candidate_role_and_path(self):
        from scdp_automation.xlsx_output import _validate_candidate

        candidate = self.root / "candidate.xlsx"
        final_template_fixture(candidate)
        workbook = load_workbook(candidate)
        workbook["APOIO"]["D2"] = "bad budget"
        workbook.save(candidate)
        workbook.close()
        with self.assertRaises(WorkbookValidationError) as caught:
            _validate_candidate(candidate)
        self.assertIn("Candidato", str(caught.exception))
        self.assertIn(str(candidate), str(caught.exception))
        self.assertIn("APOIO", str(caught.exception))

    def test_cli_reports_known_errors_and_leaves_unexpected_bugs_visible(self):
        # The browser/profile setup is external; the CLI presentation remains real.
        for error in (
            WorkbookValidationError(f"Arquivo {self.output}: reconcilie"),
            WorkbookPublishError(f"Recuperação em {self.output}"),
        ):
            with self.subTest(error=error), patch.object(cli, "prepare_chrome_profile"):

                async def fail(args, failure=error):
                    raise failure

                stderr = io.StringIO()
                with (
                    patch.object(cli, "run", fail),
                    contextlib.redirect_stderr(stderr),
                    self.assertRaises(SystemExit) as caught,
                ):
                    cli.main([])
                self.assertNotEqual(caught.exception.code, 0)
                self.assertIn(str(self.output), stderr.getvalue())
                self.assertNotIn("Traceback", stderr.getvalue())

        async def bug(args):
            raise RuntimeError("programming bug")

        with (
            patch.object(cli, "prepare_chrome_profile"),
            patch.object(cli, "run", bug),
            self.assertRaisesRegex(RuntimeError, "programming bug"),
        ):
            cli.main([])
