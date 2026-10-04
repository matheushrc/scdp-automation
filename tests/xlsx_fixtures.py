"""Small public workbook fixtures with the same source coordinates."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.writer.theme import theme_xml


def reference_fixture(path: Path) -> None:
    workbook = Workbook()
    workbook._fonts[0] = Font(name="Arial", size=10)
    workbook.loaded_theme = theme_xml.replace("4F81BD", "123456").encode()
    base = workbook.active
    base.title = "BD D&P"
    support = workbook.create_sheet("APOIO")
    summary = workbook.create_sheet("RESUMO GASTOS")
    workbook.create_sheet("CONSULTA ANALÍTICA")
    labels = (
        "ADMINISTRAÇÃO",
        "AGRONOMIA",
        "C COMPUTAÇÃO",
        "C ECONÔMICAS (NOVO)",
        "CIÊNCIAS SOCIAIS",
        "ENFERMAGEM",
        "ENG AMBIENTAL",
        "ENG CIVIL (NOVO)",
        "FILOSOFIA",
        "GEOGRAFIA",
        "HISTÓRIA",
        "LETRAS",
        "MATEMÁTICA",
        "MEDICINA",
        "PEDAGOGIA",
        "PPGCB - C Biomédicas (a partir de 2021)",
        "PPGE - Educação",
        "PPGEL - Estudos Linguísticos",
        "PPGEnf - Enfermagem",
        "PPGFil - Filosofia (a partir de 2019)",
        "PPGGeo - Geografia (a partir de 2019)",
        "PPGH - História (a partir de 2017)",
        "PROFMAT - Matemática",
        "PROFIAP - Prof em Adm Pública",
        "PPGDH - Dr História",
        "LS Enf em Oncologia",
        "OUTROS",
        "RESOLUÇÃO 049/2022-CONSUNI/CPPGEC",
    )
    codes = (
        "ADMINISTRAÇÃO",
        "AGRONOMIA",
        "C COMPUTAÇÃO",
        "C ECONÔMICAS",
        "CIÊNCIAS SOCIAIS",
        "ENFERMAGEM",
        "ENG AMBIENTAL",
        "ENGENHARIA CIVIL",
        "FILOSOFIA",
        "GEOGRAFIA",
        "HISTÓRIA",
        "LETRAS",
        "MATEMÁTICA",
        "MEDICINA",
        "PEDAGOGIA",
        "PPGCB",
        "PPGE",
        "PPGEL",
        "PPGEnf",
        "PPGFil",
        "PPGGeo",
        "PPGH",
        "PROFMAT",
        "PROFIAP",
        "PPGDH",
        "LS Enf em Oncologia",
        "DIREÇÃO",
        "CAPPG - Res 49",
    )
    for row, (label, code) in enumerate(zip(labels, codes, strict=True), 7):
        support.cell(row, 2, label)
        support.cell(row, 3, code)
        support.cell(row, 22, 200.0)
        support.cell(row, 25, 150.0)
        support.cell(row, 13, f"=V{row}")
        support.cell(row, 16, f"=Y{row}")
        support.cell(row, 11, 9999.0)  # Not the summary's allocation.
    support["I57"] = "PPGEL +"
    support["J57"] = 156
    for row, label in ((58, "PPGE"), (59, "PPGEL"), (60, "PPGH")):
        support.cell(row, 9, label)
        support.cell(row, 10, 52)
    for row, label in list(zip(range(7, 22), labels[:15], strict=True)) + [
        (26, labels[15]),
        (27, labels[16]),
        (28, labels[17]),
        (29, labels[18]),
        (30, labels[19]),
        (31, labels[20]),
        (32, labels[21]),
        (33, labels[23]),
        (34, labels[22]),
        (35, labels[25]),
        (41, labels[26]),
        (43, labels[27]),
    ]:
        summary.cell(row, 2, label)
        summary.cell(row, 3, 1000.0)
        summary.cell(row, 5, 400.0)
        summary.cell(row, 9, f"=C{row}-E{row}")
        summary.cell(row, 6, f"=VLOOKUP(B{row},APOIO!$B$7:$F$42,3,FALSE)")
    summary["B2"] = "RESUMO DE GASTOS - Diárias - Passagens - Transportes"
    summary["B3"] = "ASSESSORIA - Campus Exemplo"
    for row, title in (
        (5, "CURSOS DE GRADUAÇÃO"),
        (24, "PROGRAMAS DE PÓS"),
        (38, "OUTROS (Administrativo)"),
    ):
        summary.cell(row, 2, title)
        summary.cell(row, 3, "Recurso Total")
        summary.cell(row, 5, "Diárias & Passagens")
        summary.cell(row, 9, "Transportes")
        summary.cell(row, 13, "SALDO GERAL")
        for column, text in (
            (5, "Distribuído"),
            (6, "Utilizado"),
            (7, "Saldo"),
            (9, "Distribuído"),
            (10, "Utilizado"),
            (11, "Saldo"),
        ):
            summary.cell(row + 1, column, text)
    for merged in (
        "B2:I2",
        "B3:I3",
        "B5:B6",
        "C5:C6",
        "E5:G5",
        "I5:K5",
        "M5:M6",
        "B24:B25",
        "C24:C25",
        "E24:G24",
        "I24:K24",
        "M24:M25",
        "B38:B39",
        "C38:C39",
        "E38:G38",
        "I38:K38",
        "M38:M39",
    ):
        summary.merge_cells(merged)
    for row, text in (
        (22, "TOTAL CURSOS GRADUAÇÃO"),
        (36, "TOTAL PROGRAMAS PÓS-GRADUAÇÃO"),
        (45, "TOTAL (Diárias, Passagens, Transportes)"),
        (48, "OBS.: DIÁRIAS - OUTROS (ADMINISTRATIVO)"),
        (58, "TOTAL"),
    ):
        summary.cell(row, 2, text)
    for row, text in enumerate(
        (
            "Direção",
            "AGAS",
            "Banca de Libra",
            "CAAEX",
            "Empresa Junior",
            "StartUP Summit",
            "StartUp Weekend",
            "Sunset",
        ),
        50,
    ):
        summary.cell(row, 2, text)
    for row in summary.iter_rows(min_row=1, max_row=58, max_col=13):
        for cell in row:
            cell.font = Font(name="Liberation Sans", size=10, color="123456")
            if cell.__class__.__name__ == "MergedCell":
                continue
            cell.fill = PatternFill("solid", fgColor="DDEEFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="thin", color="112233"))
            cell.number_format = "#,##0.00"
    summary.column_dimensions["B"].width = 44
    summary.column_dimensions["C"].width = 17
    summary.row_dimensions[5].height = 32
    summary.print_area = "B2:M58"
    summary.sheet_view.showGridLines = False
    summary.page_setup.orientation = "landscape"
    for row, code in ((5, "AGRONOMIA"), (11, None), (17, "PPGEL +")):
        pcdp = "999001/26" if row != 17 else "999002/26-1C"
        base.cell(row, 4, pcdp)
        base.cell(row, 5, "Pessoa Exemplo")
        base.cell(row, 9, "Autorizada")
        if code:
            base.cell(
                row, 2, "SEG 1 GRADUAÇÃO" if code == "AGRONOMIA" else "SEG 2 MESTRADO"
            )
            base.cell(row, 3, code)
        base.cell(row + 3, 15, "Sub-Total")
        for col, value in ((16, 2.5), (17, 100), (18, 50)):
            base.cell(row + 3, col, value)
        base.cell(row + 4, 16, "Total da Viagem (R$)")
        for col, value in ((6, 10), (9, 5), (12, -2), (15, 3), (19, 156)):
            base.cell(row + 4, col, value)
    workbook.save(path)
    workbook.close()
