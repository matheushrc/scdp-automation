"""Spending rules and migration for the editable annual workbook."""

from __future__ import annotations

import math
import re
from copy import copy
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from scdp_automation.xlsx_reference import ReferenceData

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook

HISTORY_GROUP = "Mestrado e Doutorado em História"
HISTORY_CODE = "PPGH/PPGDH"
BASE_COLUMN_MAP = {12: 15, 13: 16, 14: 17, 15: 12, 16: 13, 17: 14}


def total_formula(row: int) -> str:
    return f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),D{row}+E{row},"")'


def dynamic_range(sheet: str, column: str, width: int = 1) -> str:
    height_row = {"APOIO": 1, "BASE VIAGENS": 2, "RESUMO GASTOS": 3}[sheet]
    height = f"MAX(1,'RESUMO GASTOS'!$R${height_row})"
    return f"OFFSET('{sheet}'!${column}$1,1,0,{height},{width})"


def install_names(workbook: Workbook) -> None:
    summary = workbook["RESUMO GASTOS"]
    for row, sheet, key in (
        (1, "APOIO", "A"),
        (2, "BASE VIAGENS", "A"),
        (3, "RESUMO GASTOS", "B"),
    ):
        summary.cell(
            row, 18
        ).value = f"""=MAX(1,IFERROR(LOOKUP(2,1/('{sheet}'!${key}:${key}<>""),ROW('{sheet}'!${key}:${key}))-1,1))"""
    summary.column_dimensions["R"].hidden = True
    for sheet, prefix, columns in (
        ("APOIO", "Apoio", "ABCDEFGHIJKLM"),
        ("BASE VIAGENS", "Viagens", "ABCDEFGHIJKLMNOPQ"),
        ("RESUMO GASTOS", "Resumo", "BCDEFGHIJKLMNOP"),
    ):
        for column in columns:
            workbook.defined_names.add(
                DefinedName(
                    prefix + column,
                    attr_text=dynamic_range(
                        sheet,
                        chr(
                            64 + BASE_COLUMN_MAP.get(ord(column) - 64, ord(column) - 64)
                        )
                        if sheet == "BASE VIAGENS"
                        else column,
                    ),
                )
            )
    workbook.defined_names.add(
        DefinedName("ApoioCatalogo", attr_text=dynamic_range("APOIO", "A", 3))
    )
    workbook.defined_names.add(
        DefinedName("CodigosDebito", attr_text=dynamic_range("APOIO", "A"))
    )


def charged_for_codes(codes: str) -> str:
    """Array-compatible subtotals by debit, including only decided cancellations."""
    return (
        f'SUMIFS(ViagensK,ViagensM,{codes},ViagensC,"<>Cancelada")'
        f'+SUMIFS(ViagensK,ViagensM,{codes},ViagensC,"Cancelada",ViagensN,"Sim")'
    )


def pending_for_codes(codes: str) -> str:
    return f'COUNTIFS(ViagensM,{codes},ViagensC,"Cancelada",ViagensN,"")'


def expense_formula(row: int) -> str:
    return f'=IF(A{row}="","",IF({pending_for_codes(f"A{row}")}>0,"Pendente",{charged_for_codes(f"A{row}")}))'


def group_formulas(row: int, *, weight_column: str | None = None) -> dict[str, str]:
    selected = f"--(ApoioI=$B{row})"
    split = f'--(ApoioA="PPGEL +")*Apoio{weight_column}' if weight_column else None
    result = {}
    for output, source in (("E", "D"), ("I", "E"), ("J", "G")):
        amounts = f"Apoio{source}"
        count = f"SUMPRODUCT({selected},--ISNUMBER({amounts}))"
        summed = f"SUMIF(ApoioI,$B{row},{amounts})"
        if split:
            summed += f"+SUMPRODUCT({split},{amounts})"
            count += f"+SUMPRODUCT({split},--ISNUMBER({amounts}))"
        result[f"{output}{row}"] = f'=IF({count}=0,"Pendente",{summed})'
    result[f"C{row}"] = (
        f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(I{row})),E{row}+I{row},"Pendente")'
    )
    pending = f"SUMPRODUCT({selected},{pending_for_codes('ApoioA')})"
    charged = f"SUMPRODUCT({selected},({charged_for_codes('ApoioA')}))"
    if split:
        pending += f"+SUMPRODUCT({split},{pending_for_codes('ApoioA')})"
        charged += f"+SUMPRODUCT({split},({charged_for_codes('ApoioA')}))"
    result[f"F{row}"] = f'=IF({pending}>0,"Pendente",{charged})'
    result[f"G{row}"] = (
        f'=IF(AND(ISNUMBER(E{row}),ISNUMBER(F{row})),E{row}-F{row},"Pendente")'
    )
    result[f"K{row}"] = (
        f'=IF(AND(ISNUMBER(I{row}),ISNUMBER(J{row})),I{row}-J{row},"Pendente")'
    )
    result[f"M{row}"] = (
        f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(K{row})),G{row}+K{row},"Pendente")'
    )
    result[f"P{row}"] = f"=B{row}"
    return result


def merge_history(workbook: Workbook) -> None:
    support = workbook["APOIO"]
    rows = [
        r
        for r in range(2, support.max_row + 1)
        if support.cell(r, 1).value in ("PPGH", "PPGDH", HISTORY_CODE)
    ]
    if not rows:
        raise ValueError("Categoria de História ausente na migração.")
    target = rows[0]
    for col in (4, 5, 7, 8):
        values = [support.cell(r, col).value for r in rows]
        numbers = [
            v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)
        ]
        support.cell(target, col).value = sum(numbers) if numbers else None
    support.cell(target, 1).value = HISTORY_CODE
    support.cell(target, 2).value = HISTORY_GROUP
    support.cell(target, 9).value = HISTORY_GROUP
    support.cell(target, 2).comment = None
    for row in rows[1:]:
        for cell in support[row]:
            cell.value = None
            cell.comment = None
    support.cell(target, 6).value = total_formula(target)
    support.cell(target, 13).value = expense_formula(target)
    if "tblApoioDebito" in support.tables:
        support.tables["tblApoioDebito"].ref = f"A1:M{support.max_row}"
        support.tables["tblApoioDebito"].tableColumns = []
    base = workbook["BASE VIAGENS"]
    header = [c.value for c in base[1]]
    debit_col = header.index("Código de débito") + 1
    for row in range(2, base.max_row + 1):
        if base.cell(row, debit_col).value in ("PPGH", "PPGDH"):
            base.cell(row, debit_col).value = HISTORY_CODE
    summary = workbook["RESUMO GASTOS"]
    summary["B32"] = HISTORY_GROUP
    for col in ("C", "E", "F", "G", "I", "J", "K", "M", "P"):
        summary[f"{col}62"] = None
    summary["B62"] = None


def translate_base_references(formula: str, sheet: str) -> str:
    tokens = Tokenizer(formula).items
    mapping = {chr(64 + old): chr(64 + new) for old, new in BASE_COLUMN_MAP.items()}
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        prefix, separator, address = token.value.rpartition("!")
        if separator and prefix.strip("'") != "BASE VIAGENS":
            continue
        if not separator:
            if sheet != "BASE VIAGENS":
                continue
            address = token.value
        if not re.fullmatch(
            r"\$?[A-Z]{1,3}(?:\$?\d+)?(?::\$?[A-Z]{1,3}(?:\$?\d+)?)?", address
        ):
            continue
        moved = re.sub(
            r"[A-Z]+", lambda match: mapping.get(match[0], match[0]), address
        )
        token.value = prefix + separator + moved
    return "=" + "".join(token.value for token in tokens)


def reorder_base_columns(workbook: Workbook) -> None:
    """Keep stable named ranges while moving classification behind dates."""
    from scdp_automation.xlsx_layout import range_segment_formula
    from scdp_automation.xlsx_output import BASE_HEADERS, install_base_controls

    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.value = translate_base_references(cell.value, sheet.title)
    for name in workbook.defined_names.values():
        name.attr_text = translate_base_references("=" + name.attr_text, "")[1:]
    base = workbook["BASE VIAGENS"]
    for row in range(2, base.max_row + 1):
        saved = {
            old: (
                base.cell(row, old).value,
                copy(base.cell(row, old)._style),
                copy(base.cell(row, old).comment),
            )
            for old in BASE_COLUMN_MAP
        }
        for old, new in BASE_COLUMN_MAP.items():
            value, style, comment = saved[old]
            cell = base.cell(row, new)
            cell.value, cell._style, cell.comment = value, style, comment
        base.cell(row, 15).value = range_segment_formula(row)
    dimensions = {
        old: copy(base.column_dimensions[chr(64 + old)]) for old in BASE_COLUMN_MAP
    }
    for old, new in BASE_COLUMN_MAP.items():
        dimension = dimensions[old]
        dimension.index = chr(64 + new)
        dimension.min = dimension.max = new
        base.column_dimensions[chr(64 + new)] = dimension
    for validation in base.data_validations.dataValidation:
        if validation.formula1 in ("=CodigosDebito", "CodigosDebito"):
            validation.sqref = "P2:P1048576"
        elif validation.formula1 == '"Sim,Não"':
            validation.sqref = "Q2:Q1048576"
    for area in base.conditional_formatting:
        for rule in base.conditional_formatting[area]:
            if rule.formula:
                rule.formula = [f.replace("$N2=", "$Q2=") for f in rule.formula]
    for col, header in enumerate(BASE_HEADERS, 1):
        base.cell(1, col).value = header
    base.tables["tblBaseViagens"].ref = f"A1:Q{max(2, base.max_row)}"
    base.tables["tblBaseViagens"].tableColumns = []
    install_base_controls(base)


def migrate_adjustments(workbook: Workbook, data: ReferenceData | None = None) -> None:
    """One-time migration; subsequent refreshes never rebuild manual sheets."""
    from scdp_automation.xlsx_layout import INPUT_HEADERS, LAYOUT_NAME, prepare_ranges
    from scdp_automation.xlsx_output import install_base_controls

    support = workbook["APOIO"]
    for row in range(2, support.max_row + 1):
        daily, total, transport = (support.cell(row, col).value for col in (4, 5, 6))
        if isinstance(transport, str) and transport.startswith("="):
            transport = (
                total - daily
                if all(
                    isinstance(v, (int, float)) and not isinstance(v, bool)
                    for v in (daily, total)
                )
                else None
            )
        support.cell(row, 5).value = transport
        support.cell(row, 6).value = total_formula(row)
    for col, header in enumerate(INPUT_HEADERS, 1):
        support.cell(1, col).value = header
    base = workbook["BASE VIAGENS"]
    install_base_controls(base)
    if data is not None:
        historical = {t.pcdp: t for t in data.trips}
        for row in range(2, base.max_row + 1):
            pcdp = base.cell(row, 1).value
            if base.cell(row, 14).value in (None, "") and pcdp in data.decisions:
                base.cell(row, 14).value = data.decisions[pcdp]
            trip = historical.get(pcdp)
            if trip:
                for col, value in ((15, trip.start_date), (16, trip.end_date)):
                    if base.cell(row, col).value is None:
                        base.cell(row, col).value = value
                    base.cell(row, col).number_format = "dd/mm/yyyy"
    reorder_base_columns(workbook)
    merge_history(workbook)
    prepare_ranges(workbook)
    workbook.defined_names.add(DefinedName(LAYOUT_NAME, attr_text='"8"'))


def validate_editable_layout(workbook: Workbook) -> None:
    """Validate manual inputs without imposing a fixed row order or formulas."""
    from scdp_automation.xlsx_layout import INPUT_HEADERS
    from scdp_automation.xlsx_output import WorkbookValidationError

    support = workbook["APOIO"]
    if tuple(c.value for c in support[1]) != (
        *INPUT_HEADERS,
        "Total utilizado por categoria (R$)",
    ):
        raise WorkbookValidationError("As colunas de APOIO foram alteradas.")
    seen = set()
    for row in range(2, support.max_row + 1):
        code = support.cell(row, 1).value
        if not code:
            continue
        if (
            not isinstance(code, str)
            or code != code.strip()
            or code in seen
            or any(c in code for c in "*?~")
        ):
            raise WorkbookValidationError("Código inválido ou duplicado em APOIO.")
        seen.add(code)
        for col in (4, 5, 7, 8):
            value = support.cell(row, col).value
            if value is not None and (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                raise WorkbookValidationError(
                    "Orçamento e transporte em APOIO precisam de valor numérico."
                )
        weights = [support.cell(row, col).value for col in (10, 11, 12)]
        if code == "PPGEL +":
            required = any(support.cell(row, col).value for col in (4, 5, 7, 8)) or any(
                workbook["BASE VIAGENS"].cell(r, 16).value == code
                and workbook["BASE VIAGENS"].cell(r, 11).value
                for r in range(2, workbook["BASE VIAGENS"].max_row + 1)
            )
            if (required or any(w is not None for w in weights)) and (
                any(
                    isinstance(w, bool)
                    or not isinstance(w, (int, float))
                    or not math.isfinite(w)
                    or not 0 <= w <= 1
                    for w in weights
                )
                or not math.isclose(sum(weights), 1, abs_tol=1e-10)
            ):
                raise WorkbookValidationError(
                    "O rateio PPGEL + precisa totalizar 100%."
                )
